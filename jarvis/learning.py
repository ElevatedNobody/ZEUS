"""
Project-learning cards: the organisation's own completed work as a planning
reference.

History comes from two places, each labelled:
  * DERIVED - activities ZEUS itself saw complete, with evidenced actual
    dates, their quantities, blockers met and evidence count;
  * IMPORTED - a register of past projects (CSV / XLSX) with its file and
    origin (company records, or SAMPLE / DEMO data - always shown as such).

A card finds comparable completed activities (same kind of work on the same
kind of element, compatible units, similar names) and reports, separately:

    FACT            each comparable's actual duration, quantity, productivity
    INFERENCE       how the current plan compares with them
    PREDICTION      the duration the current quantity implies at their
                    productivity (a range, never one invented number)
    RECOMMENDATION  what to check before the work starts

With no evidence there is no card: the answer says nothing comparable exists.
"""

from __future__ import annotations

import csv
import math
import re
import statistics
from pathlib import Path
from typing import Any, Dict, List, Optional

from jarvis import dates, forecast, origin_label, store
from jarvis.extraction import _work_types
from jarvis.linking import _ELEMENT_WORDS, _tokens

_COLUMNS = {
    "project": ("project", "project name", "source project"), "code": ("code", "activity id", "activity code"),
    "name": ("activity", "name", "activity name", "description"), "category": ("category", "work category", "work type"),
    "unit": ("unit", "uom"), "quantity": ("quantity", "qty", "actual quantity"),
    "planned_duration": ("planned duration", "planned_duration", "planned days"),
    "actual_start": ("actual start", "actual_start", "start"), "actual_finish": ("actual finish", "actual_finish", "finish"),
    "actual_duration": ("actual duration", "actual_duration", "actual days", "duration"),
    "contractor": ("contractor", "agency"), "location": ("location", "site", "state"),
    "conditions": ("conditions", "notes", "remarks"), "blockers": ("blockers", "hindrances", "delays"),
    "evidence": ("evidence", "source", "reference", "evidence count"),
}


def _season(day) -> str:
    d = dates.from_iso(day)
    if not d:
        return ""
    return "monsoon" if d.month in (6, 7, 8, 9) else "post-monsoon" if d.month in (10, 11) else \
        "winter" if d.month in (12, 1, 2) else "summer"


def _elements(text: str) -> set:
    low = (text or "").lower()
    return {w for w in _ELEMENT_WORDS if re.search(rf"\b{w}s?\b", low)} | (
        {"pier"} if re.search(r"\bp-?\d{1,3}\b", low) else set())


def derive_history(project_id: Optional[str] = None) -> int:
    """Completed activities with evidenced actual dates -> history rows."""
    count = 0
    projects = store.query("SELECT project_id, name, origin FROM jv_projects" +
                           (" WHERE project_id = ?" if project_id else ""), (project_id,) if project_id else ())
    for p in projects:
        with store.connection() as conn:
            try:
                data = forecast.compute(conn, p["project_id"])
            except forecast.NoSchedule:
                continue
            meta = {r["activity_id"]: dict(r) for r in conn.execute(
                "SELECT activity_id, code, name, unit, planned_qty, contractor, location FROM jv_activities WHERE "
                "project_id = ?", (p["project_id"],))}
            for aid, a in data["activities"].items():
                if a.get("status") != "complete" or a.get("type") != "task":
                    continue
                sources = a.get("status_sources") or {}
                if "actual_finish" not in sources:
                    continue
                m = meta.get(aid, {})
                start_evidenced = "actual_start" in sources
                duration = (dates.days_between(a["es"], a["ef"]) + 1) if start_evidenced else None
                quantity = m.get("planned_qty")
                blockers = [dict(r) for r in conn.execute(
                    "SELECT category, description, raised_on, expected_clear FROM jv_blockers WHERE activity_id = ? AND "
                    "status IN ('open', 'cleared')", (aid,))]
                evidence = conn.execute("SELECT COUNT(DISTINCT from_id) FROM jv_links WHERE to_kind = 'activity' AND "
                                        "to_id = ? AND from_kind = 'evidence'", (aid,)).fetchone()[0]
                conn.execute("DELETE FROM jv_historical_activities WHERE project_id = ? AND code = ? AND source = 'derived'",
                             (p["project_id"], m.get("code")))
                store.insert(conn, "jv_historical_activities", {
                    "hist_id": store.new_id("HIST"), "project_id": p["project_id"], "project_name": p["name"],
                    "activity_id": aid, "code": m.get("code"), "name": m.get("name") or a["name"],
                    "work_category": ",".join(_work_types(m.get("name") or "")), "unit": m.get("unit"),
                    "quantity": quantity, "planned_duration_days": a.get("duration"),
                    "actual_start": a["es"] if start_evidenced else None, "actual_finish": a["ef"],
                    "actual_duration_days": duration,
                    "productivity_per_day": round(quantity / duration, 3) if quantity and duration else None,
                    "contractor": m.get("contractor"), "location": m.get("location"),
                    "conditions": store.dumps({"season": _season(a["ef"]), "month": (a["ef"] or "")[:7],
                                               "start_evidenced": start_evidenced}),
                    "blockers": store.dumps(blockers), "evidence_count": evidence, "origin": p["origin"],
                    "source": "derived", "source_file": None, "created_at": store.now()})
                count += 1
    return count


def import_history(path: str, origin: str = "company") -> Dict[str, Any]:
    source = Path(path)
    rows = _read_table(source)
    if not rows:
        raise ValueError(f"{source.name}: no rows")
    headers = {h.strip().lower(): h for h in rows[0] if h}
    mapping = {k: next((headers[s] for s in syn if s in headers), None) for k, syn in _COLUMNS.items()}
    if not mapping["name"]:
        raise ValueError(f"{source.name}: no activity / name column")
    added, rejected = 0, []
    with store.connection() as conn:
        for index, row in enumerate(rows, start=2):
            get = lambda k: row.get(mapping[k]) if mapping[k] else None  # noqa: E731
            name = str(get("name") or "").strip()
            start, finish = dates.iso(get("actual_start")), dates.iso(get("actual_finish"))
            duration = _num(get("actual_duration"))
            if duration is None and start and finish:
                duration = dates.days_between(start, finish) + 1
            if not name or duration is None:
                rejected.append(f"row {index}: no actual duration or actual dates - not evidence of an outcome")
                continue
            quantity = _num(get("quantity"))
            project_name = str(get("project") or source.stem)
            code = str(get("code") or f"{source.stem}-{index}")
            conn.execute("DELETE FROM jv_historical_activities WHERE project_id IS NULL AND project_name = ? AND code = ? "
                         "AND source = 'imported'", (project_name, code))
            store.insert(conn, "jv_historical_activities", {
                "hist_id": store.new_id("HIST"), "project_id": None, "project_name": project_name, "activity_id": None,
                "code": code, "name": name, "work_category": str(get("category") or "") or ",".join(_work_types(name)),
                "unit": str(get("unit") or "") or None, "quantity": quantity,
                "planned_duration_days": int(_num(get("planned_duration"))) if _num(get("planned_duration")) else None,
                "actual_start": start, "actual_finish": finish, "actual_duration_days": int(duration),
                "productivity_per_day": round(quantity / duration, 3) if quantity and duration else None,
                "contractor": str(get("contractor") or "") or None, "location": str(get("location") or "") or None,
                "conditions": store.dumps({"notes": str(get("conditions") or ""), "season": _season(finish)}),
                "blockers": store.dumps([b.strip() for b in re.split(r"[;|]", str(get("blockers") or "")) if b.strip()]),
                "evidence_count": int(_num(get("evidence")) or 0) if _num(get("evidence")) is not None else 0,
                "origin": origin, "source": "imported", "source_file": source.name, "created_at": store.now()})
            added += 1
        store.audit("ingest_historical_activities", "history", source.name, "", {"added": added,
                                                                                 "rejected": len(rejected)}, conn=conn)
    return {"ok": True, "added": added, "rejected": rejected, "origin": origin}


def _read_table(path: Path) -> List[Dict[str, Any]]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(str(path), read_only=True, data_only=True)
        sheet = wb.worksheets[0]
        values = [list(r) for r in sheet.iter_rows(values_only=True)]
        wb.close()
        header = [str(c) if c is not None else "" for c in values[0]]
        return [dict(zip(header, r)) for r in values[1:] if any(v is not None for v in r)]
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(line for line in handle if not line.lstrip().startswith("#")))


def _num(value) -> Optional[float]:
    try:
        return float(str(value).replace(",", "").strip()) if value not in (None, "") else None
    except ValueError:
        return None


def similarity(target: Dict[str, Any], hist: Dict[str, Any]) -> Dict[str, Any]:
    t_work, h_work = set(target["work"]), set(_work_types(f"{hist['name']} {hist.get('work_category') or ''}"))
    t_el, h_el = target["elements"], _elements(hist["name"])
    reasons, total = [], 0.0
    if t_work and h_work:
        shared = t_work & h_work
        if shared:
            total += 0.45
            reasons.append(f"same work: {', '.join(sorted(shared))}")
        else:
            total -= 0.3
    if t_el and h_el and t_el & h_el:
        total += 0.2
        reasons.append(f"same element: {', '.join(sorted(t_el & h_el))}")
    if target.get("unit") and hist.get("unit"):
        if str(target["unit"]).lower() == str(hist["unit"]).lower():
            total += 0.1
            reasons.append(f"same unit ({hist['unit']})")
    tt, ht = _tokens(target["name"]), _tokens(hist["name"])
    jaccard = len(tt & ht) / max(1, len(tt | ht))
    total += 0.25 * jaccard
    return {"score": round(max(0.0, min(1.0, total)), 3), "reasons": reasons}


def cards(activity: str, project: Optional[str] = None, limit: int = 6, min_score: float = 0.45) -> Dict[str, Any]:
    with store.connection() as conn:
        row = conn.execute("SELECT * FROM jv_activities WHERE activity_id = ? OR (code = ? AND project_id = COALESCE(?, "
                           "project_id))", (activity, activity, project)).fetchone()
        if not row:
            raise ValueError(f"unknown activity '{activity}'")
        a = dict(row)
        try:
            now = forecast.compute(conn, a["project_id"])["activities"].get(a["activity_id"], {})
        except forecast.NoSchedule:
            now = {}
        history = [dict(r) for r in conn.execute("SELECT * FROM jv_historical_activities WHERE activity_id IS NULL OR "
                                                 "activity_id != ?", (a["activity_id"],))]
    target = {"name": a["name"], "code": a["code"], "unit": a.get("unit"), "quantity": a.get("planned_qty"),
              "work": _work_types(f"{a['name']} {a.get('work_category') or ''}"), "elements": _elements(a["name"]),
              "planned_duration": now.get("duration"), "forecast_start": now.get("es"), "forecast_finish": now.get("ef"),
              "status": now.get("status")}
    scored = []
    for h in history:
        sim = similarity(target, h)
        if sim["score"] >= min_score and h.get("actual_duration_days"):
            scored.append(dict(h, similarity=sim["score"], match_reasons=sim["reasons"],
                               blockers=store.loads(h["blockers"], []), conditions=store.loads(h["conditions"], {}),
                               origin_label=origin_label(h["origin"])))
    scored.sort(key=lambda h: (-h["similarity"], h["project_name"] or ""))
    comparables = scored[:limit]
    if not comparables:
        return {"ok": True, "target": target, "comparables": [], "stats": None,
                "message": "No completed comparable activity with recorded actual dates exists in this organisation's "
                           "history or imported records - nothing is estimated."}
    durations = [c["actual_duration_days"] for c in comparables]
    productivities = [c["productivity_per_day"] for c in comparables if c.get("productivity_per_day") and
                      (not target["unit"] or not c.get("unit") or str(c["unit"]).lower() == str(target["unit"]).lower())]
    stats = {"n": len(comparables), "median_duration": statistics.median(durations), "min_duration": min(durations),
             "max_duration": max(durations),
             "median_productivity": round(statistics.median(productivities), 2) if productivities else None,
             "productivity_unit": f"{target['unit'] or comparables[0].get('unit') or ''}/day" if productivities else None,
             "origins": sorted({c["origin_label"] for c in comparables})}
    facts = [f"{c['project_name']} {c['code']} ({c['name']}): actual {c['actual_duration_days']} d"
             + (f" vs planned {c['planned_duration_days']} d" if c.get("planned_duration_days") else "")
             + (f"; {c['quantity']:g} {c['unit'] or ''} at {c['productivity_per_day']:g}/day" if c.get("productivity_per_day") else "")
             + (f"; blockers: {', '.join(_blocker_label(b) for b in c['blockers'])}" if c["blockers"] else "")
             + f" [{c['origin_label']}, {c['source']}]" for c in comparables]
    inference, prediction, recommendation = [], [], []
    if target["planned_duration"]:
        gap = target["planned_duration"] - stats["median_duration"]
        inference.append(f"The current plan allows {target['planned_duration']} d; the comparable median is "
                         f"{stats['median_duration']:g} d ({'+' if gap >= 0 else ''}{gap:g} d).")
        if gap < 0:
            recommendation.append(f"Review {a['code']}'s duration: the actual durations of {stats['n']} comparable "
                                  f"activities put the median at {stats['median_duration']:g} d.")
    blocked = [c for c in comparables if c["blockers"]]
    if blocked:
        categories = sorted({_blocker_label(b) for c in blocked for b in c["blockers"]})
        inference.append(f"{len(blocked)} of {stats['n']} comparables met blockers ({', '.join(categories)}).")
        recommendation.append(f"Before {a['code']} starts, check the recurring causes: {', '.join(categories)}.")
    if target["quantity"] and productivities:
        low, high = min(productivities), max(productivities)
        fast, slow = math.ceil(target["quantity"] / high), math.ceil(target["quantity"] / low)
        prediction.append(f"{target['quantity']:g} {target['unit'] or ''} at the comparables' {low:g}-{high:g} "
                          f"{target['unit'] or ''}/day would take {fast}-{slow} d (median productivity "
                          f"{stats['median_productivity']:g}/day -> {math.ceil(target['quantity'] / stats['median_productivity'])} d).")
    elif not target["quantity"]:
        prediction.append("No quantity is recorded for this activity, so no productivity-based duration is given.")
    card = {"card_id": store.new_id("CARD"), "project_id": a["project_id"], "activity_id": a["activity_id"],
            "generated_at": store.now(), "generated_by": store.current_user(),
            "comparables": store.dumps([{k: c.get(k) for k in ("hist_id", "project_name", "code", "similarity")}
                                        for c in comparables]),
            "stats": store.dumps(stats), "implication": " ".join(inference + recommendation),
            "basis": "comparable completed activities with recorded actual dates"}
    with store.connection() as conn:
        store.insert(conn, "jv_learning_cards", card)
    return {"ok": True, "card_id": card["card_id"], "target": target, "comparables": comparables, "stats": stats,
            "FACT": facts, "INFERENCE": inference, "PREDICTION": prediction, "RECOMMENDATION": recommendation}


def _blocker_label(b) -> str:
    if isinstance(b, dict):
        return b.get("category") or (b.get("description") or "")[:40]
    return str(b)
