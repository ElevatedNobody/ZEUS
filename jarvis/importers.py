"""
Schedule providers - ZEUS does not assume one scheduling system.

Each provider reads one format into the same neutral `ImportedSchedule`:
activities (codes, names, WBS, types, durations, planned / baseline / actual
dates, progress, quantities), dependencies (FS/SS/FF/SF + lag), WBS, the
calendar and the data date. Adding a format means adding a provider class.

    CsvProvider     .csv - one row per activity, a predecessors column
    XlsxProvider    .xlsx / .xlsm - the same columns on the first sheet with them
    MspdiProvider   .xml - Microsoft Project XML (MSPDI)
    XerProvider     .xer - Primavera P6 export (TASK / TASKPRED / PROJWBS / PROJECT)

Nothing is guessed silently: a row that cannot be read becomes a warning.
"""

from __future__ import annotations

import csv
import datetime as _dt
import io
import re
import xml.etree.ElementTree as ET
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Dict, List, Optional

from jarvis import dates

_TYPES = {"task": "task", "activity": "task", "milestone": "finish_milestone", "finish milestone": "finish_milestone",
          "finish_milestone": "finish_milestone", "start milestone": "start_milestone",
          "start_milestone": "start_milestone", "tt_task": "task", "tt_mile": "start_milestone",
          "tt_finmile": "finish_milestone", "tt_rsrc": "task", "tt_loe": "task"}

_COLUMNS = {
    "code": ("code", "activity id", "activity_id", "task code", "task_code", "id", "act id", "activity code"),
    "name": ("name", "activity name", "activity_name", "task name", "task_name", "description", "activity"),
    "wbs_code": ("wbs", "wbs code", "wbs_code", "wbs id"),
    "wbs_name": ("wbs name", "wbs_name"),
    "type": ("type", "activity type", "task type", "task_type"),
    "duration": ("duration", "original duration", "duration (days)", "duration_days", "planned duration", "od"),
    "planned_start": ("start", "planned start", "planned_start", "early start", "forecast start"),
    "planned_finish": ("finish", "planned finish", "planned_finish", "early finish", "forecast finish"),
    "baseline_start": ("baseline start", "baseline_start", "bl start", "bl project start"),
    "baseline_finish": ("baseline finish", "baseline_finish", "bl finish", "bl project finish"),
    "actual_start": ("actual start", "actual_start"),
    "actual_finish": ("actual finish", "actual_finish"),
    "percent_complete": ("% complete", "percent complete", "percent_complete", "physical % complete", "progress"),
    "remaining": ("remaining duration", "remaining", "remaining_days", "rd"),
    "predecessors": ("predecessors", "predecessor", "preds", "predecessor list"),
    "contractor": ("contractor", "agency", "subcontractor", "responsible"),
    "location": ("location", "chainage", "area", "zone"),
    "unit": ("unit", "uom", "units"),
    "quantity": ("quantity", "qty", "planned quantity", "planned_qty", "budgeted quantity"),
    "constraint_date": ("constraint date", "start no earlier than", "snet"),
    "aliases": ("aliases", "also known as", "field names", "site names"),
}


@dataclass
class ImportedActivity:
    code: str
    name: str
    wbs_code: str = ""
    wbs_name: str = ""
    type: str = "task"
    duration: Optional[int] = None
    planned_start: Optional[str] = None
    planned_finish: Optional[str] = None
    baseline_start: Optional[str] = None
    baseline_finish: Optional[str] = None
    actual_start: Optional[str] = None
    actual_finish: Optional[str] = None
    percent_complete: Optional[float] = None
    remaining: Optional[int] = None
    contractor: str = ""
    location: str = ""
    unit: str = ""
    quantity: Optional[float] = None
    constraint_date: Optional[str] = None
    aliases: List[str] = field(default_factory=list)


@dataclass
class ImportedDependency:
    pred_code: str
    succ_code: str
    type: str = "FS"
    lag: int = 0


@dataclass
class ImportedSchedule:
    provider: str
    source: str
    data_date: Optional[str] = None
    project_name: str = ""
    activities: List[ImportedActivity] = field(default_factory=list)
    dependencies: List[ImportedDependency] = field(default_factory=list)
    wbs: List[Dict[str, str]] = field(default_factory=list)
    calendar: Dict = field(default_factory=lambda: {"workdays": "0123456", "holidays": []})
    warnings: List[str] = field(default_factory=list)

    def as_dict(self) -> Dict:
        return asdict(self)


def _num(value) -> Optional[float]:
    if value is None or str(value).strip() in ("", "-", "—"):
        return None
    text = str(value).strip().rstrip("%").replace(",", "")
    text = re.sub(r"\s*(d|days?|wd)$", "", text, flags=re.IGNORECASE)
    try:
        return float(text)
    except ValueError:
        return None


def _date(value) -> Optional[str]:
    if value is None or value == "":
        return None
    if isinstance(value, (_dt.date, _dt.datetime)):
        return dates.iso(value)
    return dates.iso(str(value).strip())


# "CODE", "CODE FS", "CODE: SS+2", "CODE FF -1d", "CODE +3" - the type and lag need a
# space (or colon) before them, so a code such as BR-17 is never read as BR with lag -17.
# MS Project's "12FS+2 days" form is accepted for numeric ids.
_PRED = re.compile(r"^(?P<code>\S+?)(?:(?:\s+|\s*:\s*)(?P<type>FS|SS|FF|SF)(?:\s*(?P<lag>[+-]\s*\d+)\s*(?:d|days?)?)?"
                   r"|\s+(?P<lag2>[+-]\s*\d+)\s*(?:d|days?)?)?$", re.IGNORECASE)
_PRED_NUMERIC = re.compile(r"^(?P<code>\d+)(?P<type>FS|SS|FF|SF)(?P<lag>[+-]\d+)?\s*(?:d|days?)?$", re.IGNORECASE)


def parse_predecessors(text: str, succ_code: str) -> List[ImportedDependency]:
    """'BR-P17 FS+2; RD-DRN SS, X' -> dependencies."""
    found = []
    for token in re.split(r"[;,\n]+", text or ""):
        token = token.strip()
        if not token:
            continue
        m = _PRED_NUMERIC.match(token) or _PRED.match(token)
        if not m:
            found.append(ImportedDependency(token, succ_code))
            continue
        groups = m.groupdict()
        lag_text = groups.get("lag") or groups.get("lag2")
        lag = int(re.sub(r"\s", "", lag_text)) if lag_text else 0
        found.append(ImportedDependency(m.group("code"), succ_code, (groups.get("type") or "FS").upper(), lag))
    return found


class Provider:
    name = "base"
    extensions: tuple = ()

    @classmethod
    def can_read(cls, path: Path) -> bool:
        return path.suffix.lower() in cls.extensions

    def read(self, path: Path) -> ImportedSchedule:
        raise NotImplementedError


class _TabularProvider(Provider):
    def _from_rows(self, rows: List[Dict[str, object]], path: Path, provider: str) -> ImportedSchedule:
        out = ImportedSchedule(provider=provider, source=str(path))
        if not rows:
            out.warnings.append("no rows")
            return out
        headers = {str(h).strip().lower(): h for h in rows[0].keys() if h is not None}
        mapping = {}
        for field_name, synonyms in _COLUMNS.items():
            for synonym in synonyms:
                if synonym in headers:
                    mapping[field_name] = headers[synonym]
                    break
        if "code" not in mapping or "name" not in mapping:
            out.warnings.append(f"no activity code / name column (found: {', '.join(headers)})")
            return out
        wbs_seen = {}
        for index, row in enumerate(rows, start=2):
            get = lambda key: row.get(mapping[key]) if key in mapping else None  # noqa: E731
            code = str(get("code") or "").strip()
            name = str(get("name") or "").strip()
            if not code:
                continue
            if code.startswith("#") or code.lower() in ("data date", "data_date"):
                if code.lower().lstrip("#").strip() in ("data date", "data_date"):
                    out.data_date = _date(name)
                continue
            kind = _TYPES.get(str(get("type") or "task").strip().lower(), "task")
            duration = _num(get("duration"))
            act = ImportedActivity(
                code=code, name=name or code, wbs_code=str(get("wbs_code") or "").strip(),
                wbs_name=str(get("wbs_name") or "").strip(), type=kind,
                duration=int(duration) if duration is not None else None,
                planned_start=_date(get("planned_start")), planned_finish=_date(get("planned_finish")),
                baseline_start=_date(get("baseline_start")), baseline_finish=_date(get("baseline_finish")),
                actual_start=_date(get("actual_start")), actual_finish=_date(get("actual_finish")),
                percent_complete=_num(get("percent_complete")),
                remaining=int(_num(get("remaining"))) if _num(get("remaining")) is not None else None,
                contractor=str(get("contractor") or "").strip(), location=str(get("location") or "").strip(),
                unit=str(get("unit") or "").strip(), quantity=_num(get("quantity")),
                constraint_date=_date(get("constraint_date")),
                aliases=[a.strip() for a in re.split(r"[;|]", str(get("aliases") or "")) if a.strip()])
            if act.duration is None and act.planned_start and act.planned_finish and kind == "task":
                act.duration = max(1, (dates.from_iso(act.planned_finish) - dates.from_iso(act.planned_start)).days + 1)
            if act.duration is None and kind == "task":
                out.warnings.append(f"row {index} ({code}): no duration - 1 day assumed")
                act.duration = 1
            out.activities.append(act)
            if act.wbs_code and act.wbs_code not in wbs_seen:
                wbs_seen[act.wbs_code] = act.wbs_name
            out.dependencies.extend(parse_predecessors(str(get("predecessors") or ""), code))
        out.wbs = [{"code": c, "name": n} for c, n in wbs_seen.items()]
        known = {a.code for a in out.activities}
        for dep in list(out.dependencies):
            if dep.pred_code not in known:
                out.warnings.append(f"{dep.succ_code}: predecessor '{dep.pred_code}' is not in the schedule - ignored")
                out.dependencies.remove(dep)
        return out


class CsvProvider(_TabularProvider):
    name = "csv"
    extensions = (".csv",)

    def read(self, path: Path) -> ImportedSchedule:
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        lines = [line for line in text.splitlines() if line.strip()]
        data_date = None
        body = []
        for line in lines:
            m = re.match(r"^\s*#\s*data[_ ]date\s*[:=,]\s*(.+)$", line, re.IGNORECASE)
            if m:
                data_date = _date(m.group(1).strip().strip(","))
            elif not line.lstrip().startswith("#"):
                body.append(line)
        rows = list(csv.DictReader(io.StringIO("\n".join(body))))
        out = self._from_rows(rows, path, self.name)
        out.data_date = out.data_date or data_date
        return out


class XlsxProvider(_TabularProvider):
    name = "xlsx"
    extensions = (".xlsx", ".xlsm")

    def read(self, path: Path) -> ImportedSchedule:
        from openpyxl import load_workbook
        workbook = load_workbook(str(path), data_only=True, read_only=True)
        best = None
        data_date = None
        for sheet in workbook.worksheets:
            rows = [list(r) for r in sheet.iter_rows(values_only=True)]
            for r in rows[:6]:
                if r and isinstance(r[0], str) and re.match(r"^\s*#?\s*data[_ ]date", r[0], re.IGNORECASE):
                    data_date = _date(r[1] if len(r) > 1 else None)
            for start, header in enumerate(rows[:10]):
                labels = [str(c).strip().lower() for c in header if c is not None]
                if any(s in labels for s in _COLUMNS["code"]) and any(s in labels for s in _COLUMNS["name"]):
                    records = [dict(zip([str(c) if c is not None else "" for c in header], r))
                               for r in rows[start + 1:] if any(v is not None for v in r)]
                    best = best or records
                    break
        workbook.close()
        out = self._from_rows(best or [], path, self.name)
        out.data_date = out.data_date or data_date
        return out


class MspdiProvider(Provider):
    name = "ms_project_xml"
    extensions = (".xml",)
    _TYPE = {"0": "FF", "1": "FS", "2": "SF", "3": "SS"}

    @classmethod
    def can_read(cls, path: Path) -> bool:
        if path.suffix.lower() != ".xml":
            return False
        head = path.read_bytes()[:4000].decode("utf-8", "replace")
        return "schemas.microsoft.com/project" in head

    @staticmethod
    def _hours(duration: str) -> float:
        m = re.match(r"PT(\d+)H(\d+)M", duration or "")
        return (int(m.group(1)) + int(m.group(2)) / 60.0) if m else 0.0

    def read(self, path: Path) -> ImportedSchedule:
        root = ET.parse(str(path)).getroot()
        ns = {"p": root.tag.split("}")[0].strip("{")} if root.tag.startswith("{") else {}
        q = (lambda tag: f"p:{tag}") if ns else (lambda tag: tag)
        text = lambda node, tag: (node.findtext(q(tag), default="", namespaces=ns) or "").strip()  # noqa: E731
        out = ImportedSchedule(provider=self.name, source=str(path), project_name=text(root, "Name"))
        out.data_date = _date(text(root, "StatusDate")) or _date(text(root, "CurrentDate"))
        hours_per_day = float(text(root, "MinutesPerDay") or 480) / 60.0
        by_uid: Dict[str, str] = {}
        tasks = root.findall(f"{q('Tasks')}/{q('Task')}", ns)
        for task in tasks:
            if text(task, "Summary") == "1" or text(task, "UID") == "0":
                wbs = text(task, "WBS")
                if wbs:
                    out.wbs.append({"code": wbs, "name": text(task, "Name")})
                continue
            uid = text(task, "UID")
            code = text(task, "WBS") or text(task, "ID") or uid
            by_uid[uid] = code
            milestone = text(task, "Milestone") == "1"
            baseline = next((b for b in task.findall(q("Baseline"), ns) if text(b, "Number") in ("0", "")), None)
            out.activities.append(ImportedActivity(
                code=code, name=text(task, "Name") or code,
                wbs_code=".".join(code.split(".")[:-1]) if "." in code else "",
                type="finish_milestone" if milestone else "task",
                duration=0 if milestone else max(1, round(self._hours(text(task, "Duration")) / hours_per_day)),
                planned_start=_date(text(task, "Start")), planned_finish=_date(text(task, "Finish")),
                baseline_start=_date(text(baseline, "Start")) if baseline is not None else None,
                baseline_finish=_date(text(baseline, "Finish")) if baseline is not None else None,
                actual_start=_date(text(task, "ActualStart")), actual_finish=_date(text(task, "ActualFinish")),
                percent_complete=_num(text(task, "PercentComplete"))))
        for task in tasks:
            succ = by_uid.get(text(task, "UID"))
            if not succ:
                continue
            for link in task.findall(q("PredecessorLink"), ns):
                pred = by_uid.get(text(link, "PredecessorUID"))
                if not pred:
                    out.warnings.append(f"{succ}: predecessor UID {text(link, 'PredecessorUID')} not found")
                    continue
                lag_days = round(float(text(link, "LinkLag") or 0) / 600.0 / hours_per_day)
                out.dependencies.append(ImportedDependency(pred, succ, self._TYPE.get(text(link, "Type"), "FS"),
                                                           int(lag_days)))
        return out


class XerProvider(Provider):
    name = "primavera_xer"
    extensions = (".xer",)
    _TYPE = {"PR_FS": "FS", "PR_SS": "SS", "PR_FF": "FF", "PR_SF": "SF"}

    @staticmethod
    def tables(text: str) -> Dict[str, List[Dict[str, str]]]:
        tables: Dict[str, List[Dict[str, str]]] = {}
        current, fields = None, []
        for line in text.splitlines():
            parts = line.rstrip("\r\n").split("\t")
            if parts[0] == "%T":
                current, fields = parts[1].strip(), []
                tables[current] = []
            elif parts[0] == "%F":
                fields = parts[1:]
            elif parts[0] == "%R" and current:
                tables[current].append(dict(zip(fields, parts[1:])))
        return tables

    def read(self, path: Path) -> ImportedSchedule:
        raw = path.read_bytes()
        text = raw.decode("utf-8", "replace") if raw[:3] != b"\xef\xbb\xbf" else raw[3:].decode("utf-8", "replace")
        if text.count("�") > 20:
            text = raw.decode("cp1252", "replace")
        t = self.tables(text)
        out = ImportedSchedule(provider=self.name, source=str(path))
        project = (t.get("PROJECT") or [{}])[0]
        out.project_name = project.get("proj_short_name", "")
        out.data_date = _date(project.get("last_recalc_date") or project.get("next_data_date"))
        wbs = {w.get("wbs_id"): w for w in t.get("PROJWBS", [])}
        out.wbs = [{"code": w.get("wbs_short_name", ""), "name": w.get("wbs_name", "")} for w in wbs.values()]
        by_id: Dict[str, str] = {}
        for task in t.get("TASK", []):
            code = task.get("task_code") or task.get("task_id")
            by_id[task.get("task_id")] = code
            kind = _TYPES.get((task.get("task_type") or "TT_Task").lower(), "task")
            hours = _num(task.get("target_drtn_hr_cnt")) or 0
            remain = _num(task.get("remain_drtn_hr_cnt"))
            out.activities.append(ImportedActivity(
                code=code, name=task.get("task_name") or code,
                wbs_code=(wbs.get(task.get("wbs_id")) or {}).get("wbs_short_name", ""),
                wbs_name=(wbs.get(task.get("wbs_id")) or {}).get("wbs_name", ""),
                type=kind, duration=0 if kind != "task" else max(1, round(hours / 8.0)),
                planned_start=_date(task.get("target_start_date") or task.get("early_start_date")),
                planned_finish=_date(task.get("target_end_date") or task.get("early_end_date")),
                actual_start=_date(task.get("act_start_date")), actual_finish=_date(task.get("act_end_date")),
                percent_complete=_num(task.get("phys_complete_pct")),
                remaining=round(remain / 8.0) if remain is not None and task.get("act_start_date") else None))
        for link in t.get("TASKPRED", []):
            pred, succ = by_id.get(link.get("pred_task_id")), by_id.get(link.get("task_id"))
            if not pred or not succ:
                out.warnings.append(f"relationship {link.get('task_pred_id')} refers to a task not in the file")
                continue
            lag = round((_num(link.get("lag_hr_cnt")) or 0) / 8.0)
            out.dependencies.append(ImportedDependency(pred, succ, self._TYPE.get(link.get("pred_type"), "FS"),
                                                       int(lag)))
        return out


PROVIDERS = [CsvProvider, XlsxProvider, MspdiProvider, XerProvider]


def provider_for(path: Path) -> Optional[Provider]:
    for cls in PROVIDERS:
        try:
            if cls.can_read(path):
                return cls()
        except OSError:
            continue
    return None


def read_schedule(path) -> ImportedSchedule:
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"{path} does not exist")
    provider = provider_for(path)
    if provider is None:
        raise ValueError(f"no schedule provider reads '{path.suffix}' files "
                         f"(supported: CSV, XLSX, MS Project XML, Primavera XER)")
    schedule = provider.read(path)
    if not schedule.activities:
        raise ValueError(f"{path.name}: no activities could be read"
                         + (f" - {'; '.join(schedule.warnings[:3])}" if schedule.warnings else ""))
    return schedule
