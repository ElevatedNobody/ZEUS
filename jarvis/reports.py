"""
Artifacts built from the project record - never from a model's imagination.

    progress report       DOCX   status, forecast vs baseline, blockers, forecast changes with causal chains
    delay chronology      DOCX / XLSX / PDF   the formal dated chronology with evidence references
    evidence pack         ZIP    every evidence file, a manifest with SHA-256 and the event chain, the chronology
    schedule analysis     XLSX   activities, baseline vs forecast, snapshots, forecast changes, delay windows
    learning report       DOCX   the project-learning card for one activity

Every file is written through ZEUS's office helpers: reopened and parsed
before it is reported as produced (core.office.finish), and registered with
its provenance in the artifact registry. DEMO projects are labelled DEMO on
the first line of every artifact.
"""

from __future__ import annotations

import hashlib
import json
import time
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional

from jarvis import chronology, dates, events, explain, origin_label, service, store

SUBFOLDER = "ZEUS"


def _out(project: Dict[str, Any], stem: str, ext: str) -> Path:
    from core.office import documents_dir, safe_stem, unique_path
    folder = documents_dir(SUBFOLDER)
    return unique_path(folder, safe_stem(f"{project['code']}_{stem}_{time.strftime('%Y%m%d_%H%M')}"), ext)


def _label(project: Dict[str, Any]) -> str:
    label = origin_label(project["origin"])
    return "" if label == "RECORDED" else f"{label} DATA - "


def _provenance(project_id: str) -> Dict[str, Any]:
    chain = events.verify_chain()
    snap = store.one("SELECT snapshot_id, seq, captured_at, project_finish FROM jv_snapshots WHERE project_id = ? "
                     "ORDER BY seq DESC LIMIT 1", (project_id,))
    return {"generated_at": store.now_iso(), "generated_by": store.current_user() or "unknown",
            "event_chain_ok": chain["ok"], "event_chain_head": chain["head"], "events_checked": chain["events"],
            "snapshot": snap}


def _finish(path: Path, ext: str, extra: Dict[str, Any]) -> Dict[str, Any]:
    from core.office import finish
    return finish(path, ext, open_after=False, extra=extra)


# ------------------------------------------------------------ progress report --

def progress_report(project: str) -> Dict[str, Any]:
    from docx import Document
    from docx.shared import Pt
    project_id = service.resolve_project(project)
    p = service.get_project(project_id)
    ov = service.overview(project_id)
    table = service.schedule_table(project_id)
    with store.connection() as conn:
        changes = explain.list_changes(conn, project_id, only_finish=True)
        why = explain.why_moved(conn, project_id) if changes else None
        blockers = [dict(r) for r in conn.execute(
            "SELECT b.*, a.code FROM jv_blockers b LEFT JOIN jv_activities a ON a.activity_id = b.activity_id "
            "WHERE b.project_id = ? ORDER BY b.raised_on", (project_id,))]
        deadlines = chronology.deadlines(conn, project_id)
    prov = _provenance(project_id)
    doc = Document()
    doc.core_properties.author = "ZEUS"
    doc.core_properties.title = f"{p['name']} progress report"
    doc.add_heading(f"{_label(p)}{p['name']} - Progress report", level=0)
    doc.add_paragraph(f"Project {p['code']} · data date {dates.fmt(ov.get('data_date'))} · schedule version "
                      f"{ov.get('version') or '—'} · generated {prov['generated_at']} by {prov['generated_by']}")
    if _label(p):
        doc.add_paragraph(f"This report is built from {origin_label(p['origin'])} data. It is not a record of a real project.")
    doc.add_heading("Summary", level=1)
    summary = doc.add_table(rows=0, cols=2)
    summary.style = "Light Grid Accent 1"
    for label, value in (("Baseline finish", dates.fmt(ov.get("baseline_finish"))),
                         ("Forecast finish", dates.fmt(ov.get("forecast_finish"))),
                         ("Finish variance", f"{ov.get('finish_variance_days'):+d} days" if ov.get("finish_variance_days") is not None else "—"),
                         ("Actual progress", f"{ov.get('actual_progress_pct', 0):.1f}%"),
                         ("Planned progress at data date", f"{ov.get('planned_progress_pct', 0):.1f}%"),
                         ("Progress basis", ov.get("progress_basis", "")),
                         ("Critical path", " -> ".join(ov.get("critical_path") or [])),
                         ("Open blockers", str(ov["counts"]["open_blockers"])),
                         ("Updates awaiting approval", str(ov["counts"]["pending_approvals"])),
                         ("Open sync conflicts", str(ov["counts"]["open_conflicts"]))):
        cells = summary.add_row().cells
        cells[0].text, cells[1].text = label, value
    if why and why.get("found"):
        doc.add_heading("Why the finish date moved", level=1)
        doc.add_paragraph(why["answer"])
        for step in why["steps"]:
            doc.add_paragraph(f"{step['step']}: {step['title']} - {step['detail']}", style="List Bullet")
    doc.add_heading("Activities (forecast from the approved schedule and approved field events)", level=1)
    grid = doc.add_table(rows=1, cols=8)
    grid.style = "Light Grid Accent 1"
    for cell, text in zip(grid.rows[0].cells, ("Code", "Activity", "Status", "% complete", "Forecast start",
                                               "Forecast finish", "Baseline finish", "Float (d)")):
        cell.text = text
    for a in table["activities"]:
        cells = grid.add_row().cells
        values = (a["code"], a["name"], a["status"].replace("_", " "),
                  "—" if a.get("percent_complete") is None else f"{a['percent_complete']:g}", dates.fmt(a["es"]),
                  dates.fmt(a["ef"]), dates.fmt(a.get("baseline_finish")),
                  "—" if a.get("total_float") is None else str(a["total_float"]))
        for cell, text in zip(cells, values):
            cell.text = text
    doc.add_heading("Blockers", level=1)
    if not blockers:
        doc.add_paragraph("No blocker has been reported.")
    for b in blockers:
        doc.add_paragraph(f"{dates.fmt(b['raised_on'])} · {b.get('code') or '—'} · {b['status'].upper()} · "
                          f"{b.get('category') or 'other'}: {b.get('description') or ''}"
                          + (f" (expected to clear {dates.fmt(b['expected_clear'])})" if b.get("expected_clear") else "")
                          + (f"  Original: {b['original_text']}" if b.get("original_text") and
                             b["original_text"] != b.get("description") else ""), style="List Bullet")
    doc.add_heading("Contract notice deadlines", level=1)
    if not deadlines:
        doc.add_paragraph("No contract notice deadline is open.")
    for d in deadlines:
        doc.add_paragraph(f"{d['clause']} ({d['notice_type']}) for {d.get('activity_code') or 'the project'}: due "
                          f"{dates.fmt(d['due_date'])} - {d['status'].upper()}"
                          + (f", notice {d['notice_reference']} of {dates.fmt(d['notice_issued_on'])}" if d.get("notice_reference") else ""),
                          style="List Bullet")
    doc.add_heading("Provenance", level=1)
    doc.add_paragraph(f"Event chain verified: {'yes' if prov['event_chain_ok'] else 'NO - see the audit'} "
                      f"({prov['events_checked']} events, head {prov['event_chain_head'][:16]}…). "
                      f"Snapshot {((prov['snapshot'] or {}).get('snapshot_id') or '—')}. Every date above comes from "
                      f"the CPM engine; no figure was produced by a language model.")
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.size = run.font.size or Pt(10)
    path = _out(p, "progress_report", "docx")
    doc.save(str(path))
    return _finish(path, "docx", {"title": f"{p['name']} progress report", "project": p["code"],
                                  "label": origin_label(p["origin"]), "activities": len(table["activities"])})


# ------------------------------------------------------------------ chronology --

def chronology_report(project: str, fmt: str = "docx", date_from: str = "", date_to: str = "", activity: str = "",
                      kinds: Optional[List[str]] = None) -> Dict[str, Any]:
    project_id = service.resolve_project(project)
    p = service.get_project(project_id)
    with store.connection() as conn:
        activity_id = None
        if activity:
            row = conn.execute("SELECT activity_id FROM jv_activities WHERE project_id = ? AND (code = ? OR activity_id = ?)",
                               (project_id, activity, activity)).fetchone()
            activity_id = row["activity_id"] if row else None
        rows = chronology.timeline(conn, project_id, date_from=date_from or None, date_to=date_to or None,
                                   activity_id=activity_id, kinds=kinds)
        evidence = {r["evidence_id"]: dict(r) for r in conn.execute(
            "SELECT evidence_id, kind, title, sha256 FROM jv_evidence WHERE project_id = ?", (project_id,))}
        deadlines = chronology.deadlines(conn, project_id)
    header = ("No.", "Date", "Event", "Activity", "Detail", "Source", "Captured", "Status", "Evidence (SHA-256)")
    table = []
    for i, r in enumerate(rows, 1):
        refs = "; ".join(f"{e} {(evidence.get(e) or {}).get('kind', '')} {(evidence.get(e) or {}).get('sha256', '')[:12]}"
                         for e in r["evidence_ids"])
        table.append((i, dates.fmt(r["occurred_at"]), r["title"], r.get("activity_code") or "", r.get("detail") or "",
                      r.get("source_type") or "", r.get("captured_at") or "", r.get("status") or "", refs))
    prov = _provenance(project_id)
    title = f"{_label(p)}{p['name']} - Delay chronology"
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        ws = wb.active
        ws.title = "Chronology"
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append([f"Generated {prov['generated_at']} by {prov['generated_by']}; event chain "
                   f"{'verified' if prov['event_chain_ok'] else 'NOT verified'}"])
        ws.append(list(header))
        for cell in ws[3]:
            cell.font = Font(bold=True)
        for row in table:
            ws.append(list(row))
        dl = wb.create_sheet("Notice deadlines")
        dl.append(["Clause", "Notice", "Activity", "Trigger date", "Due", "Status", "Notice ref", "Notice date"])
        for d in deadlines:
            dl.append([d["clause"], d["notice_type"], d.get("activity_code"), d["trigger_date"], d["due_date"],
                       d["status"], d.get("notice_reference"), d.get("notice_issued_on")])
        wb.properties.creator = "ZEUS"
        path = _out(p, "delay_chronology", "xlsx")
        wb.save(str(path))
    elif fmt == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        path = _out(p, "delay_chronology", "pdf")
        styles = getSampleStyleSheet()
        small = styles["BodyText"].clone("small", fontSize=7, leading=9)
        story = [Paragraph(title, styles["Title"]),
                 Paragraph(f"Generated {prov['generated_at']} by {prov['generated_by']}. Event chain "
                           f"{'verified' if prov['event_chain_ok'] else 'NOT verified'}.", styles["BodyText"]),
                 Spacer(1, 8)]
        data = [[Paragraph(h, small) for h in header]] + [[Paragraph(str(c), small) for c in row] for row in table]
        grid = Table(data, repeatRows=1, colWidths=[22, 52, 110, 60, 200, 50, 70, 50, 150])
        grid.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                                  ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story.append(grid)
        SimpleDocTemplate(str(path), pagesize=landscape(A4), title=title, author="ZEUS").build(story)
    else:
        from docx import Document
        doc = Document()
        doc.core_properties.author = "ZEUS"
        doc.add_heading(title, level=0)
        doc.add_paragraph(f"Generated {prov['generated_at']} by {prov['generated_by']}. Event chain "
                          f"{'verified' if prov['event_chain_ok'] else 'NOT verified'} ({prov['events_checked']} events). "
                          f"Every entry below is a recorded event with its evidence reference.")
        grid = doc.add_table(rows=1, cols=len(header))
        grid.style = "Light Grid Accent 1"
        for cell, text in zip(grid.rows[0].cells, header):
            cell.text = text
        for row in table:
            for cell, text in zip(grid.add_row().cells, row):
                cell.text = str(text)
        doc.add_heading("Contract notice deadlines", level=1)
        for d in deadlines:
            doc.add_paragraph(f"{d['clause']} {d['notice_type']} - {d.get('activity_code') or 'project'}: triggered "
                              f"{dates.fmt(d['trigger_date'])}, due {dates.fmt(d['due_date'])}, {d['status'].upper()}",
                              style="List Bullet")
        path = _out(p, "delay_chronology", "docx")
        doc.save(str(path))
    return _finish(path, fmt, {"title": title, "entries": len(table), "project": p["code"],
                               "label": origin_label(p["origin"])})


# --------------------------------------------------------------- evidence pack --

def evidence_pack(project: str, activity: str = "") -> Dict[str, Any]:
    project_id = service.resolve_project(project)
    p = service.get_project(project_id)
    items = service.evidence_list(project_id, activity=activity, limit=5000)
    chron = chronology_report(project_id, "xlsx", activity=activity)
    path = _out(p, "evidence_pack" + (f"_{activity}" if activity else ""), "zip")
    manifest = {"project": {k: p[k] for k in ("project_id", "code", "name", "origin")}, "label": origin_label(p["origin"]),
                "activity": activity or None, "provenance": _provenance(project_id), "evidence": [], "events": []}
    with store.connection() as conn:
        for item in items:
            row = dict(conn.execute("SELECT * FROM jv_evidence WHERE evidence_id = ?", (item["evidence_id"],)).fetchone())
            entry = {k: row[k] for k in ("evidence_id", "kind", "title", "filename", "sha256", "bytes", "occurred_at",
                                         "captured_at", "captured_by", "device_id", "source_type", "language", "origin")}
            entry["original_text"] = row["original_text"]
            entry["activities"] = item.get("activities")
            manifest["evidence"].append(entry)
            for ev in conn.execute("SELECT e.event_id, e.event_type, e.activity_id, e.occurred_at, e.captured_at, "
                                   "e.device_id, e.actor, e.content_hash, e.chain_hash, s.status, s.decided_by, "
                                   "s.decided_at FROM jv_events e LEFT JOIN jv_event_status s ON s.event_id = e.event_id "
                                   "WHERE e.evidence_ids LIKE ?", (f'%"{item["evidence_id"]}"%',)):
                if ev["event_id"] not in {e["event_id"] for e in manifest["events"]}:
                    manifest["events"].append(dict(ev))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for entry in manifest["evidence"]:
            row = store.one("SELECT stored_path FROM jv_evidence WHERE evidence_id = ?", (entry["evidence_id"],))
            if row and row["stored_path"] and Path(row["stored_path"]).is_file():
                arc = f"evidence/{entry['evidence_id']}{Path(row['stored_path']).suffix}"
                zf.write(row["stored_path"], arc)
                entry["file_in_pack"] = arc
            elif entry.get("original_text"):
                arc = f"evidence/{entry['evidence_id']}.txt"
                zf.writestr(arc, entry["original_text"])
                entry["file_in_pack"] = arc
        if chron.get("saved"):
            zf.write(chron["saved"], "chronology.xlsx")
        zf.writestr("manifest.json", json.dumps(manifest, indent=2, ensure_ascii=False, default=str))
        zf.writestr("README.txt", f"{_label(p)}ZEUS evidence pack for {p['name']}.\nEvery file's SHA-256 is in "
                                  f"manifest.json; text evidence is stored as captured (original language).\n")
    verified = _verify_pack(path, manifest)
    try:
        from core.artifacts import register_artifact
        register_artifact(str(path), kind="zip", tool="export_evidence_pack", request=f"evidence pack {p['code']}",
                          sources=[e["evidence_id"] for e in manifest["evidence"]][:200], verified=verified["ok"])
    except Exception:
        pass
    if not verified["ok"]:
        return {"error": f"evidence pack failed verification: {verified['problems'][:3]}", "path": str(path)}
    return {"saved": str(path), "format": "zip", "verified": True, "evidence": len(manifest["evidence"]),
            "events": len(manifest["events"]), "bytes": path.stat().st_size, "label": origin_label(p["origin"]),
            "chronology": chron.get("saved"), "manifest_hashes_checked": len(manifest["evidence"])}


def _verify_pack(path: Path, manifest: Dict[str, Any]) -> Dict[str, Any]:
    problems = []
    with zipfile.ZipFile(path) as zf:
        names = set(zf.namelist())
        for entry in manifest["evidence"]:
            arc = entry.get("file_in_pack")
            if not arc or arc not in names:
                problems.append(f"{entry['evidence_id']} missing from the pack")
                continue
            if hashlib.sha256(zf.read(arc)).hexdigest() != entry["sha256"]:
                problems.append(f"{entry['evidence_id']} hash does not match")
    return {"ok": not problems, "problems": problems}


# ------------------------------------------------------------ schedule analysis --

def schedule_analysis(project: str) -> Dict[str, Any]:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    project_id = service.resolve_project(project)
    p = service.get_project(project_id)
    table = service.schedule_table(project_id)
    ov = service.overview(project_id)
    with store.connection() as conn:
        snaps = [dict(r) for r in conn.execute("SELECT seq, captured_at, data_date, project_finish, reason FROM jv_snapshots "
                                               "WHERE project_id = ? ORDER BY seq", (project_id,))]
        changes = [dict(r) for r in conn.execute("SELECT * FROM jv_forecast_changes WHERE project_id = ? ORDER BY created_at",
                                                 (project_id,))]
        windows = chronology.delay_windows(conn, project_id)
        blockers = [dict(r) for r in conn.execute("SELECT b.*, a.code FROM jv_blockers b LEFT JOIN jv_activities a ON "
                                                  "a.activity_id = b.activity_id WHERE b.project_id = ?", (project_id,))]
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    bold = Font(bold=True)
    ws.append([f"{_label(p)}{p['name']} - schedule analysis"])
    ws["A1"].font = Font(bold=True, size=13)
    for label, value in (("Data date", ov.get("data_date")), ("Schedule version", ov.get("version")),
                         ("Baseline finish", ov.get("baseline_finish")), ("Forecast finish", ov.get("forecast_finish")),
                         ("Finish variance (days)", ov.get("finish_variance_days")),
                         ("Actual progress %", ov.get("actual_progress_pct")),
                         ("Planned progress %", ov.get("planned_progress_pct")),
                         ("Critical path", " -> ".join(ov.get("critical_path") or []))):
        ws.append([label, value])
    sheet = wb.create_sheet("Activities")
    cols = ["code", "name", "type", "status", "percent_complete", "duration", "remaining", "es", "ef", "ls", "lf",
            "total_float", "free_float", "critical", "near_critical", "baseline_start", "baseline_finish",
            "finish_variance_days", "evidence_count"]
    sheet.append(cols + ["driver"])
    for c in sheet[1]:
        c.font = bold
    red = PatternFill("solid", fgColor="FFE4E6")
    for a in table["activities"]:
        drv = a.get("driver") or {}
        sheet.append([a.get(c) for c in cols] + [f"{drv.get('kind')}{(' ' + str(drv.get('pred_code'))) if drv.get('pred_code') else ''}"])
        if a.get("critical"):
            for cell in sheet[sheet.max_row]:
                cell.fill = red
    s2 = wb.create_sheet("Snapshots")
    s2.append(["seq", "captured_at", "data_date", "project_finish", "reason"])
    for s in snaps:
        s2.append([s["seq"], s["captured_at"], s["data_date"], s["project_finish"], s["reason"]])
    s3 = wb.create_sheet("Forecast changes")
    s3.append(["change_id", "captured_at", "previous_finish", "new_finish", "delta_days", "trigger", "causal chain"])
    for c in changes:
        chain = store.loads(c["chain"], [])
        s3.append([c["change_id"], c["captured_at"], c["previous_finish"], c["new_finish"], c["delta_days"],
                   c["trigger_kind"], " -> ".join(f"{x['code']}({x['delta_days']:+d})" for x in chain)])
    s4 = wb.create_sheet("Delay windows")
    s4.append(["window", "finish_before", "finish_after", "movement_days", "driving_activity", "driving_path",
               "field_events_in_window", "schedule_version_changed"])
    for w in windows:
        s4.append([w["window"], w["finish_before"], w["finish_after"], w["movement_days"], w["driving_activity"],
                   " -> ".join(w["driving_path"]), w["field_events_in_window"], w["schedule_version_changed"]])
    s5 = wb.create_sheet("Blockers")
    s5.append(["activity", "raised_on", "category", "status", "expected_clear", "cleared_on", "description", "original"])
    for b in blockers:
        s5.append([b.get("code"), b["raised_on"], b.get("category"), b["status"], b.get("expected_clear"),
                   b.get("cleared_on"), b.get("description"), b.get("original_text")])
    wb.properties.creator = "ZEUS"
    path = _out(p, "schedule_analysis", "xlsx")
    wb.save(str(path))
    return _finish(path, "xlsx", {"title": f"{p['name']} schedule analysis", "project": p["code"],
                                  "label": origin_label(p["origin"]), "windows": len(windows)})


# --------------------------------------------------------------- learning report --

def learning_report(activity: str, project: Optional[str] = None) -> Dict[str, Any]:
    from docx import Document
    from jarvis import learning
    card = learning.cards(activity, project)
    row = store.one("SELECT project_id FROM jv_activities WHERE activity_id = ? OR code = ?", (activity, activity))
    p = service.get_project(row["project_id"])
    doc = Document()
    doc.core_properties.author = "ZEUS"
    t = card["target"]
    doc.add_heading(f"{_label(p)}Project-learning card - {t['code']} {t['name']}", level=0)
    if not card["comparables"]:
        doc.add_paragraph(card["message"])
    else:
        for section in ("FACT", "INFERENCE", "PREDICTION", "RECOMMENDATION"):
            doc.add_heading(section, level=1)
            for line in card.get(section) or ["—"]:
                doc.add_paragraph(line, style="List Bullet")
        doc.add_paragraph(f"Sources: {', '.join(card['stats']['origins'])}. Historical values are reported as recorded; "
                          f"none is estimated.")
    path = _out(p, f"learning_{t['code']}", "docx")
    doc.save(str(path))
    return _finish(path, "docx", {"title": f"learning card {t['code']}", "comparables": len(card["comparables"]),
                                  "label": origin_label(p["origin"])})
