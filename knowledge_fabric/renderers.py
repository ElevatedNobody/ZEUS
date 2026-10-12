"""
ONE VERIFIED ANALYSIS -> MULTIPLE OUTPUT FORMATS.

Every renderer reads the same saved AnalysisObject - never the question again -
so the Excel workbook, the PowerPoint deck, the Word report and the PDF carry
the same numbers, the same evidence ids and the same recommendations.

After writing, each file is:
  1. reopened and checked by core.orchestration.validators (opens, has content,
     nothing runs off a slide, formulas resolve);
  2. checked for consistency with the analysis: its id is inside, and every
     number in it comes from the analysis (facts, calculations, evidence text);
  3. registered in ZEUS's artifact list with its sources and verification.

A file that fails a check is still kept (so it can be inspected) but is marked
NOT verified - the result says which check failed.
"""

from __future__ import annotations

import re
import time
from contextlib import contextmanager
from contextvars import ContextVar
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from core.logger import get_logger
from knowledge_fabric import analysis as analysis_mod, contradictions
from knowledge_fabric.analysis import AnalysisObject

logger = get_logger("knowledge_fabric.renderers")

FORMATS = ("xlsx", "docx", "pptx", "pdf", "brief", "action_plan", "md")
AUTHOR = "ZEUS"

# Petroleum palette (docs/UI_DESIGN_SYSTEM.md): deep crude, steel, flare amber.
INK = "0B1A22"
STEEL = "3A4E5A"
AMBER = "E39B2B"
PALE = "EEF2F4"
LINE = "C9D3D9"


# --------------------------------------------------------------- helpers ----

_output_override: ContextVar[Optional[Path]] = ContextVar("knowledge_fabric_artifacts", default=None)


@contextmanager
def writing_to(folder: Path):
    """Evaluation runs write their files to a temporary folder and do not
    register them as user artifacts (context-local, like store.using_database)."""
    token = _output_override.set(Path(folder))
    try:
        yield Path(folder)
    finally:
        _output_override.reset(token)


def output_dir() -> Path:
    override = _output_override.get()
    if override is not None:
        override.mkdir(parents=True, exist_ok=True)
        return override
    from core.paths import documents_root
    folder = documents_root() / "Company Intelligence"
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def _slug(text: str, limit: int = 48) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "_", text or "analysis").strip("_")
    return (slug[:limit].rstrip("_") or "analysis")


def target_path(obj: AnalysisObject, fmt: str) -> Path:
    ext = {"brief": "docx", "action_plan": "xlsx"}.get(fmt, fmt)
    suffix = {"brief": "_brief", "action_plan": "_action_plan"}.get(fmt, "")
    return output_dir() / f"{_slug(obj.question)}{suffix}_{obj.analysis_id}.{ext}"


def _fmt(value: Any) -> str:
    if value is None:
        return "n/a"
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        text = f"{value:,.2f}"
        return text.rstrip("0").rstrip(".") if "." in text else text
    return str(value)


_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _clean(text: Any, limit: Optional[int] = None) -> str:
    text = _ILLEGAL.sub(" ", str(text if text is not None else ""))
    text = re.sub(r"[ \t]+", " ", text).strip()
    if limit and len(text) > limit:
        text = text[:limit].rsplit(" ", 1)[0] + " ..."
    return text


def _cell_text(text: Any, limit: int = 32000) -> str:
    """Document text written into a spreadsheet is DATA: a leading '=' would make
    Excel evaluate it as a formula (formula injection), so it is defused."""
    text = _clean(text, limit)
    return (" " + text) if text[:1] in ("=", "+", "-", "@") else text


def _calc_value_text(calc: Dict[str, Any]) -> str:
    value = calc.get("value")
    if value is None:
        return "not calculated - " + "; ".join(calc.get("notes", [])[:1])
    if isinstance(value, dict):
        return ", ".join(f"{k}: {_fmt(v)}" for k, v in value.items() if not isinstance(v, (list, dict)))
    if isinstance(value, list):
        rows = []
        for row in value[:12]:
            if isinstance(row, dict):
                key = row.get("year") or row.get("scenario") or row.get("entity")
                val = row.get("value", row.get("total", row.get("percent")))
                rows.append(f"{key}: {_fmt(val)}")
        return "; ".join(rows) + (" ..." if len(value) > 12 else "")
    return _fmt(value)


def _status_line(obj: AnalysisObject) -> str:
    verified = {"VERIFIED": "answer VERIFIED against company records",
                "VALIDATED": "answer VALIDATED against the evidence (not company records)",
                "INSUFFICIENT_DATA": "INSUFFICIENT DATA"}.get(obj.verification_state, "answer NOT verified - see validation")
    origin = obj.gaps.get("origin_warning")
    return verified + (f"; {origin}" if origin else "")


def _series_table(obj: AnalysisObject) -> Tuple[List[str], List[List[Any]]]:
    """Years down, one column per series (label + metric) - the data behind the charts."""
    series = obj.series[:8]
    years = sorted({p["year"] for s in series for p in s["points"]})
    headers = ["Year"] + [f"{s['label']} - {s['metric']}" + (f" ({s['unit']})" if s.get("unit") else "") for s in series]
    rows = []
    for year in years:
        row: List[Any] = [year]
        for s in series:
            point = next((p for p in s["points"] if p["year"] == year), None)
            row.append(point["value"] if point else None)
        rows.append(row)
    return headers, rows


def _created(obj: AnalysisObject) -> str:
    return time.strftime("%Y-%m-%d %H:%M", time.localtime(obj.created_at))


# ------------------------------------------------------------------ XLSX ----

def _xlsx(obj: AnalysisObject, path: Path, action_plan_only: bool = False) -> None:
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    book = Workbook()
    header_fill = PatternFill("solid", fgColor=INK)
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color=LINE)
    wrap = Alignment(wrap_text=True, vertical="top")

    def sheet(title: str, headers: List[str], rows: List[List[Any]], widths: Optional[List[int]] = None):
        ws = book.create_sheet(title[:31])
        ws.append(headers)
        for cell in ws[1]:
            cell.fill, cell.font = header_fill, header_font
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for row in rows:
            ws.append([_cell_text(v) if isinstance(v, str) else v for v in row])
        for index, width in enumerate(widths or [], 1):
            ws.column_dimensions[get_column_letter(index)].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = wrap
                cell.border = Border(bottom=thin)
        ws.freeze_panes = "A2"
        return ws

    summary = book.active
    summary.title = "Summary"
    summary_rows = [("ZEUS company intelligence", ""), ("Analysis id", obj.analysis_id), ("Question", obj.question),
                    ("Created", _created(obj)), ("Status", _status_line(obj)),
                    ("Answer written by", obj.narrative_source), ("Summary", obj.summary)]
    if obj.gaps.get("missing_years"):
        from knowledge_fabric.coverage import _ranges
        summary_rows.append(("Missing years", ", ".join(_ranges(obj.gaps["missing_years"]))))
    for key, value in summary_rows:
        summary.append([key, _cell_text(value)])
    summary["A1"].font = Font(bold=True, size=14, color=INK)
    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 110
    for row in summary.iter_rows(min_row=2):
        row[0].font = Font(bold=True, color=STEEL)
        row[1].alignment = wrap

    if obj.records.get("records") is not None and not action_plan_only:
        # Every matching register row (the chat shows the first 25), and live
        # formulas beside the engine's count and total - the two must agree.
        from knowledge_fabric.tables import parse_number
        rec = obj.records
        cols, measure = rec.get("shown_columns", []), rec.get("measure")
        out_rows = []
        for r in rec["records"]:
            row = [(parse_number(r["data"].get(c)) if c == measure and parse_number(r["data"].get(c)) is not None
                    else r["data"].get(c, "")) for c in cols]
            out_rows.append(row + [r.get("evidence_id") or ""])
        ws = sheet("Records", cols + ["Evidence"], out_rows, [22] * len(cols) + [10])
        last = len(out_rows) + 1
        engine = {c["name"].split(": ", 1)[-1]: c for c in obj.calculations}
        ws.append([])
        ws.append(["Records matching", f"=COUNTA(A2:A{last})" if out_rows else 0, "engine",
                   (engine.get("records matching") or {}).get("value"), (engine.get("records matching") or {}).get("id")])
        if measure in cols and out_rows:
            letter = get_column_letter(cols.index(measure) + 1)
            total = engine.get(f"total {measure}") or {}
            ws.append([f"Total {measure}", f"=SUM({letter}2:{letter}{last})", "engine", total.get("value"), total.get("id")])
            for cell in ws[letter][1:last]:
                cell.number_format = "#,##0.00"
        ws.append(["Conditions", _cell_text("; ".join(rec.get("conditions") or ["all records"]))])
        ws.append(["Dates compared with", rec.get("as_of")])
        ws.append(["Register", _cell_text(f"{rec['register'].get('name')} - {rec['register'].get('title')}")])

    if not action_plan_only:
        headers, rows = _series_table(obj)
        if rows:
            ws = sheet("Series", headers, rows, [8] + [24] * (len(headers) - 1))
            for row in ws.iter_rows(min_row=2, min_col=2):
                for cell in row:
                    cell.number_format = "#,##0.00"
            if len(rows) >= 2:
                chart = LineChart()
                chart.title = "Figures by year (from the records)"
                chart.y_axis.title = "value"
                chart.x_axis.title = "year"
                data = Reference(ws, min_col=2, min_row=1, max_col=len(headers), max_row=len(rows) + 1)
                chart.add_data(data, titles_from_data=True)
                chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(rows) + 1))
                chart.height, chart.width = 9, 18
                ws.add_chart(chart, f"{get_column_letter(len(headers) + 2)}2")
            # Live formulas beside the engine's numbers: the workbook recomputes,
            # and the two columns must agree.
            live = book.create_sheet("Checks")
            live.append(["Series", "Measure", "Excel formula", "Analysis engine", "Engine calc id"])
            for cell in live[1]:
                cell.fill, cell.font = header_fill, header_font
            last = len(rows) + 1
            for col_index, s in enumerate(obj.series[:8], start=2):
                col = get_column_letter(col_index)
                rng = f"Series!{col}2:{col}{last}"
                label = s["label"] if s["label"] != "Company" else ""
                prefix = f"{label}: " if label else ""
                engine = {c["name"]: c for c in obj.calculations}
                for measure, formula, calc_name in (
                        ("total", f"=SUM({rng})", f"{prefix}total {s['metric']}"),
                        ("average", f"=AVERAGE({rng})", f"{prefix}average {s['metric']}"),
                        ("median", f"=MEDIAN({rng})", f"{prefix}median {s['metric']}"),
                        ("minimum", f"=MIN({rng})", f"{prefix}minimum {s['metric']}"),
                        ("maximum", f"=MAX({rng})", f"{prefix}maximum {s['metric']}")):
                    calc = engine.get(calc_name)
                    live.append([f"{s['label']} - {s['metric']}", measure, formula,
                                 calc["value"] if calc else None, calc["id"] if calc else ""])
            for column, width in zip("ABCDE", (34, 12, 30, 18, 14)):
                live.column_dimensions[column].width = width

        fact_rows = [[f.get("id", ""), f.get("year"), f.get("period"), f.get("entity_label") or "company",
                      f.get("metric"), f.get("value"), f.get("unit"), f.get("title"), f.get("page"), f.get("cell_ref"),
                      f.get("origin"), f.get("document_id")] for f in obj.facts]
        if fact_rows:
            sheet("Facts", ["Id", "Year", "Period", "Entity", "Metric (as written)", "Value", "Unit", "Document",
                            "Page", "Cell", "Origin", "Document id"], fact_rows,
                  [6, 7, 10, 18, 26, 14, 10, 36, 6, 8, 10, 30])
        sheet("Calculations", ["Id", "Calculation", "Result", "Unit", "Method", "Inputs (fact ids / years)", "Notes"],
              [[c["id"], c["name"], _calc_value_text(c), c.get("unit"), c.get("method"),
                ", ".join(f"{i.get('title', '')[:24]} {i.get('year')}" for i in c.get("inputs", [])[:12]),
                "; ".join(c.get("notes", []))] for c in obj.calculations],
              [6, 36, 40, 10, 30, 50, 50])
        sheet("Evidence", ["Id", "Citation", "Authority", "Origin", "Classification", "Year", "Page", "Section",
                           "Excerpt", "Flags"],
              [[e["evidence_id"], e["citation"], e["authority_label"], e["origin"], e["classification"], e.get("year"),
                e.get("page"), e.get("section_path"), _clean(e["text"], 1500), "; ".join(e.get("flags", []))]
               for e in obj.evidence], [6, 40, 22, 10, 12, 7, 6, 30, 90, 30])
        if obj.drivers:
            sheet("Drivers", ["Driver", "How it is known", "Statement", "Evidence", "Calculation"],
                  [[d["driver"], d["kind"], d["statement"], ", ".join(d.get("evidence_ids", [])),
                    d.get("calculation_id", "")] for d in obj.drivers], [40, 14, 80, 12, 12])
        if obj.contradictions:
            rows = []
            for c in obj.contradictions:
                if c["type"] in ("numeric", "reconciliation"):
                    rows.append([c["metric"], c["period"], f"{c['source_a']['title']}: {_fmt(c['source_a']['value'])}",
                                 f"{c['source_b']['title']}{' (summed rows)' if c['source_b'].get('derived') else ''}: "
                                 f"{_fmt(c['source_b']['value'])}",
                                 "; ".join(c["explanations"]), c["preference_reason"]])
                elif c["type"] == "reconciliation_summary":
                    rows.append([c["metric"], f"{c['first_period']} to {c['last_period']}", c["stated_title"],
                                 f"{c['summed_title']} (summed rows)",
                                 f"{c['count']} periods differ; largest gap {_fmt(c['largest_gap_percent'])}% "
                                 f"({c['largest_gap_period']})", c["preference_reason"]])
                elif c["type"] == "record":
                    rows.append([c["register"], c["record_key"], c.get("title") or "", c.get("evidence_id") or "",
                                 c["detail"], "reported for a person to check - not corrected"])
                else:
                    rows.append([c["relation"], c["subject"], "", "",
                                 ", ".join(s["cause"] or "?" for s in c["stated"]), c["note"]])
            sheet("Conflicts", ["Metric", "Period", "Source A", "Source B", "Conflict", "Preference"], rows,
                  [20, 10, 40, 40, 60, 30])

    rec_fields = ["id", "action", "owner", "timeline", "dependencies", "cost_category", "expected_impact", "risk",
                  "kpi", "data_required", "success_condition", "monitoring_frequency", "confidence", "basis",
                  "evidence_ids"]
    if obj.recommendations:
        sheet("Recommendations", [f.replace("_", " ").upper() for f in rec_fields],
              [[", ".join(r.get(f, [])) if isinstance(r.get(f), list) else r.get(f) for f in rec_fields]
               for r in obj.recommendations], [5, 60, 22, 12, 30, 18, 36, 30, 30, 36, 40, 12, 10, 30, 12])
    if obj.implementation_plan:
        rows = [[phase["phase"], a["id"], a["action"], a["owner"], ", ".join(a.get("dependencies", [])), a.get("kpi")]
                for phase in obj.implementation_plan for a in phase["actions"]]
        sheet("Implementation Plan", ["Phase", "Action id", "Action", "Owner", "Dependencies", "KPI"], rows,
              [26, 9, 60, 22, 36, 30])
    if obj.kpis:
        sheet("KPIs", ["KPI", "Owner", "Linked to", "Baseline", "Target", "Monitoring"],
              [[k["kpi"], k["owner"], k["linked_to"],
                (f"{_fmt(k['baseline']['value'])} {k['baseline'].get('unit') or ''} ({k['baseline']['year']})"
                 if isinstance(k["baseline"], dict) else k["baseline"]), k["target"], k["monitoring_frequency"]]
               for k in obj.kpis], [36, 22, 10, 40, 30, 14])
    if obj.risks:
        sheet("Risks", ["Risk", "Linked to", "Mitigation"],
              [[r["risk"], r["linked_to"], r["mitigation"]] for r in obj.risks], [70, 14, 60])
    docs: Dict[str, Dict[str, Any]] = {}
    for e in obj.evidence:
        docs.setdefault(e["document_id"], e)
    sheet("Provenance", ["Document id", "Title", "File", "Type", "Authority", "Origin", "Classification", "Source hash"],
          [[d, e["title"], e.get("filename"), e.get("document_type"), e["authority_label"], e["origin"],
            e["classification"], e.get("source_hash")] for d, e in docs.items()],
          [30, 40, 30, 18, 22, 10, 12, 20])
    props = book.properties
    props.creator, props.lastModifiedBy = AUTHOR, AUTHOR
    props.title = _clean(obj.question, 200)
    props.subject = f"analysis {obj.analysis_id}"
    props.keywords = obj.analysis_id
    book.save(path)


# ------------------------------------------------------------------ DOCX ----

def _docx(obj: AnalysisObject, path: Path, brief: bool = False) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)
    for name, size in (("Heading 1", 16), ("Heading 2", 13)):
        styles[name].font.color.rgb = RGBColor.from_string(INK)
        styles[name].font.size = Pt(size)

    def table(headers: List[str], rows: List[List[Any]]):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Light Grid Accent 1"
        for cell, text in zip(t.rows[0].cells, headers):
            cell.text = text
            for run in cell.paragraphs[0].runs:
                run.font.bold = True
        for row in rows:
            cells = t.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = _clean(_fmt(value) if isinstance(value, (int, float)) else value, 700)
        return t

    title = doc.add_heading("Executive brief" if brief else "Company intelligence report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    meta = doc.add_paragraph()
    meta.add_run("Question: ").bold = True
    meta.add_run(_clean(obj.question))
    meta = doc.add_paragraph()
    meta.add_run("Analysis: ").bold = True
    meta.add_run(f"{obj.analysis_id} - created {_created(obj)} - {_status_line(obj)}")

    doc.add_heading("Summary", level=1)
    doc.add_paragraph(_clean(obj.summary))
    if obj.gaps.get("origin_warning"):
        p = doc.add_paragraph()
        run = p.add_run("Data origin: " + obj.gaps["origin_warning"])
        run.bold = True
        run.font.color.rgb = RGBColor.from_string("9A5B00")

    if obj.records.get("records") is not None:
        rec = obj.records
        cols = rec.get("shown_columns", [])
        limit = 15 if brief else 60
        doc.add_heading(f"Records - {rec['register'].get('name')}", level=1)
        doc.add_paragraph(_clean(f"{len(rec['records'])} record(s) in {rec['register'].get('title')} match: "
                                 + ("; ".join(rec.get("conditions") or []) or "all records")
                                 + f". Dates compared with {rec.get('as_of')}."))
        if rec["records"] and cols:
            table(cols + ["Evidence"], [[r["data"].get(c, "") for c in cols] + [r.get("evidence_id") or ""]
                                        for r in rec["records"][:limit]])
            if len(rec["records"]) > limit:
                doc.add_paragraph(f"... and {len(rec['records']) - limit} more - every record is in the Excel export.")
    headers, rows = _series_table(obj)
    if rows:
        doc.add_heading("Figures from the records", level=1)
        table(headers, rows[:40])
    key_calcs = [c for c in obj.calculations if c.get("ok")][: (8 if brief else 40)]
    if key_calcs:
        doc.add_heading("Key calculations" if brief else "Calculations", level=1)
        table(["Id", "Calculation", "Result", "Unit"] + ([] if brief else ["Method"]),
              [[c["id"], c["name"], _calc_value_text(c), c.get("unit") or ""] + ([] if brief else [c.get("method")])
               for c in key_calcs])
    if obj.drivers:
        doc.add_heading("Possible drivers", level=1)
        for d in obj.drivers[: (4 if brief else 12)]:
            kind = {"evidenced": "stated in the records", "correlated": "correlation, not proof of cause",
                    "hypothesis": "hypothesis"}.get(d["kind"], d["kind"])
            ids = ", ".join(d.get("evidence_ids", []) + ([d["calculation_id"]] if d.get("calculation_id") else []))
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(_clean(d["driver"])).bold = True
            p.add_run(f" ({kind}) - {_clean(d['statement'], 400)}" + (f" [{ids}]" if ids else ""))
    if obj.contradictions and not brief:
        doc.add_heading("Conflicts between sources", level=1)
        for c in obj.contradictions[:12]:
            doc.add_paragraph(_clean(contradictions.line(c), 900), style="List Bullet")
    if obj.recommendations:
        doc.add_heading("Recommended actions", level=1)
        doc.add_paragraph("From the practice library, matched to the drivers the records support. "
                          "Expected impact is described, not invented: quantify it from the data listed as required.")
        for r in obj.recommendations[: (3 if brief else 20)]:
            doc.add_heading(f"{r['id']}. {_clean(r['action'], 160)}", level=2)
            fields = [("Owner", r.get("owner")), ("Timeline", r.get("timeline")),
                      ("Dependencies", ", ".join(r.get("dependencies", []))), ("Cost category", r.get("cost_category")),
                      ("Expected impact", r.get("expected_impact")), ("Risk", r.get("risk")), ("KPI", r.get("kpi")),
                      ("Data required", r.get("data_required")), ("Success condition", r.get("success_condition")),
                      ("Monitoring frequency", r.get("monitoring_frequency")), ("Confidence", r.get("confidence")),
                      ("Basis", "; ".join(r.get("basis", [])))]
            if brief:
                fields = fields[:3] + [fields[6], fields[10]]
            table(["Field", "Detail"], [[k, v or ""] for k, v in fields])
    if obj.implementation_plan and not brief:
        doc.add_heading("Implementation plan", level=1)
        for phase in obj.implementation_plan:
            doc.add_heading(phase["phase"], level=2)
            for a in phase["actions"]:
                doc.add_paragraph(f"{a['id']} - {a['action']} (owner: {a['owner']}; KPI: {a.get('kpi')})",
                                  style="List Bullet")
    if obj.kpis and not brief:
        doc.add_heading("KPIs to monitor", level=1)
        table(["KPI", "Owner", "Linked to", "Baseline", "Monitoring"],
              [[k["kpi"], k["owner"], k["linked_to"],
                (f"{_fmt(k['baseline']['value'])} {k['baseline'].get('unit') or ''} ({k['baseline']['year']})"
                 if isinstance(k["baseline"], dict) else k["baseline"]), k["monitoring_frequency"]] for k in obj.kpis])
    if obj.risks and not brief:
        doc.add_heading("Risks", level=1)
        for r in obj.risks:
            doc.add_paragraph(f"{r['risk']} - mitigation: {r['mitigation']}", style="List Bullet")
    gaps = _gap_lines(obj)
    if gaps:
        doc.add_heading("Data gaps", level=1)
        for line in gaps:
            doc.add_paragraph(line, style="List Bullet")
    if obj.assumptions and not brief:
        doc.add_heading("Assumptions and limits", level=1)
        for a in obj.assumptions:
            doc.add_paragraph(_clean(a), style="List Bullet")
    if obj.evidence:
        doc.add_heading("Sources", level=1)
        for e in obj.evidence[: (5 if brief else 20)]:
            p = doc.add_paragraph(style="List Number")
            p.add_run(f"[{e['evidence_id']}] {e['citation']}").bold = True
            p.add_run(f" - {e['authority_label']}, {e['origin']}"
                      + ("" if brief else f"\n\"{_clean(e['text'], 400)}\""))
    core = doc.core_properties
    core.author, core.last_modified_by = AUTHOR, AUTHOR
    core.title = _clean(obj.question, 200)
    core.subject = f"analysis {obj.analysis_id}"
    core.keywords = obj.analysis_id
    doc.save(path)


def _gap_lines(obj: AnalysisObject) -> List[str]:
    from knowledge_fabric.coverage import _ranges
    lines = []
    if obj.gaps.get("missing_years"):
        lines.append("No records for " + ", ".join(_ranges(obj.gaps["missing_years"])))
    if obj.gaps.get("sample_only_years"):
        lines.append("Only sample data for " + ", ".join(_ranges(obj.gaps["sample_only_years"])))
    return lines


# ------------------------------------------------------------------ PPTX ----

def _pptx(obj: AnalysisObject, path: Path) -> None:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    blank = deck.slide_layouts[6]
    ink, amber, steel = RGBColor.from_string(INK), RGBColor.from_string(AMBER), RGBColor.from_string(STEEL)

    def text_box(slide, left, top, width, height, lines, size=14, bold_first=False, color=None):
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        frame = box.text_frame
        frame.word_wrap = True
        for index, line in enumerate(lines):
            para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            run = para.add_run()
            run.text = line
            run.font.size = Pt(size)
            run.font.color.rgb = color or ink
            run.font.bold = bold_first and index == 0
            para.space_after = Pt(6)
        return box

    def new_slide(title: str, kicker: str = ""):
        slide = deck.slides.add_slide(blank)
        bar = slide.shapes.add_shape(1, 0, 0, deck.slide_width, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = amber
        bar.line.fill.background()
        text_box(slide, 0.6, 0.35, 12.1, 0.8, [title], size=26, bold_first=True)
        if kicker:
            text_box(slide, 0.6, 1.05, 12.1, 0.45, [kicker], size=12, color=steel)
        footer = text_box(slide, 0.6, 7.0, 12.1, 0.35, [f"ZEUS - analysis {obj.analysis_id}"], size=9, color=steel)
        return slide

    # 1 title
    slide = deck.slides.add_slide(blank)
    bg = slide.shapes.add_shape(1, 0, 0, deck.slide_width, deck.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = ink
    bg.line.fill.background()
    text_box(slide, 0.8, 2.2, 11.7, 1.6, [_clean(obj.question, 150)], size=30, bold_first=True,
             color=RGBColor(0xFF, 0xFF, 0xFF))
    text_box(slide, 0.8, 4.2, 11.7, 1.2, [f"Company intelligence - analysis {obj.analysis_id}", _created(obj),
                                          _clean(_status_line(obj), 160)], size=13, color=amber)

    # 2 summary
    slide = new_slide("Summary", _clean(obj.gaps.get("origin_warning", ""), 150))
    findings = [_clean(obj.summary, 260)]
    for c in [c for c in obj.calculations if c.get("ok") and not isinstance(c.get("value"), (list, dict))][:4]:
        findings.append(f"{_clean(c['name'], 70)}: {_fmt(c['value'])}{(' ' + c['unit']) if c.get('unit') else ''} [{c['id']}]")
    text_box(slide, 0.6, 1.6, 12.1, 5.2, findings, size=16)

    # 3 chart
    headers, rows = _series_table(obj)
    if len(rows) >= 2:
        slide = new_slide("Figures by year", "values as reported in the records - see the Sources slide")
        data = CategoryChartData()
        data.categories = [str(r[0]) for r in rows]
        for col, header in enumerate(headers[1:5], start=1):
            data.add_series(header[:60], [r[col] for r in rows])
        frame = slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.6), Inches(1.6), Inches(12.1), Inches(5.2), data)
        chart = frame.chart
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(11)
        chart.category_axis.tick_labels.font.size = Pt(11)
        chart.value_axis.tick_labels.font.size = Pt(11)

    # 4 calculations table
    calcs = [c for c in obj.calculations if c.get("ok")][:8]
    if calcs:
        slide = new_slide("Key calculations", "computed by the analysis engine from the table facts")
        shape = slide.shapes.add_table(len(calcs) + 1, 3, Inches(0.6), Inches(1.6), Inches(12.1), Inches(0.4 * (len(calcs) + 1)))
        t = shape.table
        t.columns[0].width, t.columns[1].width, t.columns[2].width = Inches(0.9), Inches(5.2), Inches(6.0)
        for col, text in enumerate(("Id", "Calculation", "Result")):
            t.cell(0, col).text = text
        for row, c in enumerate(calcs, start=1):
            for col, text in enumerate((c["id"], _clean(c["name"], 60),
                                        _clean(_calc_value_text(c) + (f" {c['unit']}" if c.get("unit") else ""), 80))):
                t.cell(row, col).text = text
        for r in range(len(calcs) + 1):
            for col in range(3):
                for para in t.cell(r, col).text_frame.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(12)

    # 5 drivers
    if obj.drivers:
        slide = new_slide("Possible drivers", "evidenced = stated in the records; correlated = moved together, not proof")
        lines = []
        for d in obj.drivers[:5]:
            ids = ", ".join(d.get("evidence_ids", []) + ([d["calculation_id"]] if d.get("calculation_id") else []))
            lines.append(f"{_clean(d['driver'], 70)} ({d['kind']})" + (f" [{ids}]" if ids else ""))
        text_box(slide, 0.6, 1.6, 12.1, 5.2, lines, size=16)

    # 6 recommendations (3 per slide)
    recs = obj.recommendations[:9]
    for start in range(0, len(recs), 3):
        slide = new_slide("Recommended actions" + (" (continued)" if start else ""),
                          "practice library matched to the evidence - impact to be quantified from company data")
        for index, r in enumerate(recs[start:start + 3]):
            top = 1.6 + index * 1.75
            text_box(slide, 0.6, top, 12.1, 1.7,
                     [f"{r['id']}. {_clean(r['action'], 130)}",
                      f"Owner: {r['owner']} | Timeline: {r['timeline']} | Cost: {r.get('cost_category')} | Confidence: {r['confidence']}",
                      f"KPI: {_clean(r.get('kpi'), 80)} | Success: {_clean(r.get('success_condition'), 90)}"],
                     size=12, bold_first=True)

    # 7 roadmap
    if obj.implementation_plan:
        slide = new_slide("Implementation roadmap")
        count = len(obj.implementation_plan[:3])
        width = 12.1 / max(count, 1)
        for index, phase in enumerate(obj.implementation_plan[:3]):
            lines = [phase["phase"]] + [f"{a['id']} {_clean(a['action'], 70)}" for a in phase["actions"][:4]]
            text_box(slide, 0.6 + index * width, 1.6, width - 0.2, 5.2, lines, size=12, bold_first=True)

    # 8 gaps, conflicts, risks
    lines = _gap_lines(obj)
    lines += [f"Conflict: {c['metric']} {c['period']} - {_fmt(c['source_a']['value'])} vs {_fmt(c['source_b']['value'])}"
              for c in obj.contradictions[:3] if c["type"] == "numeric"]
    lines += [f"Risk: {_clean(r['risk'], 110)}" for r in obj.risks[:4]]
    if lines:
        slide = new_slide("Gaps, conflicts and risks")
        text_box(slide, 0.6, 1.6, 12.1, 5.2, lines[:8], size=14)

    # 9 sources
    if obj.evidence:
        slide = new_slide("Sources")
        text_box(slide, 0.6, 1.4, 12.1, 5.5,
                 [f"[{e['evidence_id']}] {_clean(e['citation'], 90)} - {e['authority_label']}, {e['origin']}"
                  for e in obj.evidence[:10]], size=12)
    props = deck.core_properties
    props.author, props.last_modified_by = AUTHOR, AUTHOR
    props.title = _clean(obj.question, 200)
    props.subject = f"analysis {obj.analysis_id}"
    props.keywords = obj.analysis_id
    deck.save(path)


# ------------------------------------------------------------------- PDF ----

def _pdf(obj: AnalysisObject, path: Path) -> None:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=9.5, leading=12.5)
    small = ParagraphStyle("small", parent=body, fontSize=8.5, leading=11)
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], textColor=colors.HexColor("#" + INK), fontSize=15)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], textColor=colors.HexColor("#" + STEEL), fontSize=12)

    def para(text: Any, style=body) -> Paragraph:
        return Paragraph(escape(_clean(text)), style)

    def table(headers: List[str], rows: List[List[Any]], widths: List[float]):
        data = [[para(h, small) for h in headers]] + \
               [[para(_fmt(v) if isinstance(v, (int, float)) else (v or ""), small) for v in row] for row in rows]
        t = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + PALE)),
                               ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#" + LINE)),
                               ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        return t

    story: List[Any] = [para("Company intelligence report", h1), para(f"Question: {obj.question}"),
                        para(f"Analysis {obj.analysis_id} - created {_created(obj)} - {_status_line(obj)}", small),
                        Spacer(1, 6), para("Summary", h2), para(obj.summary)]
    headers, rows = _series_table(obj)
    if rows:
        story += [para("Figures from the records", h2),
                  table(headers[:6], [r[:6] for r in rows[:40]], [18] + [34] * (min(len(headers), 6) - 1))]
    calcs = [c for c in obj.calculations if c.get("ok")][:30]
    if calcs:
        story += [para("Calculations", h2),
                  table(["Id", "Calculation", "Result", "Method"],
                        [[c["id"], c["name"], _calc_value_text(c) + (f" {c['unit']}" if c.get("unit") else ""),
                          c.get("method")] for c in calcs], [12, 60, 60, 48])]
    if obj.drivers:
        story.append(para("Possible drivers", h2))
        for d in obj.drivers[:10]:
            ids = ", ".join(d.get("evidence_ids", []) + ([d["calculation_id"]] if d.get("calculation_id") else []))
            story.append(para(f"- {d['driver']} ({d['kind']}): {_clean(d['statement'], 300)}" + (f" [{ids}]" if ids else "")))
    if obj.recommendations:
        story.append(para("Recommended actions", h2))
        story.append(table(["Id", "Action", "Owner", "Timeline", "KPI", "Confidence"],
                           [[r["id"], r["action"], r["owner"], r["timeline"], r.get("kpi"), r["confidence"]]
                            for r in obj.recommendations[:15]], [10, 70, 30, 22, 36, 18]))
    gap_lines = _gap_lines(obj)
    if gap_lines or obj.gaps.get("origin_warning"):
        story.append(para("Data gaps and origin", h2))
        for line in gap_lines + ([obj.gaps["origin_warning"]] if obj.gaps.get("origin_warning") else []):
            story.append(para("- " + line))
    if obj.evidence:
        story.append(para("Sources", h2))
        for e in obj.evidence[:20]:
            story.append(para(f"[{e['evidence_id']}] {e['citation']} - {e['authority_label']}, {e['origin']}", small))

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.drawString(15 * mm, 10 * mm, f"ZEUS - analysis {obj.analysis_id} - page {document.page}")
        canvas.restoreState()

    document = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
                                 topMargin=15 * mm, bottomMargin=18 * mm, title=_clean(obj.question, 200),
                                 author=AUTHOR, subject=f"analysis {obj.analysis_id}", keywords=obj.analysis_id)
    document.build(story, onFirstPage=footer, onLaterPages=footer)


def _markdown(obj: AnalysisObject, path: Path) -> None:
    header = (f"# {obj.question}\n\nAnalysis {obj.analysis_id} - created {_created(obj)} - {_status_line(obj)}\n\n")
    path.write_text(header + obj.answer + "\n", encoding="utf-8")


# ----------------------------------------------------------- verification ---

def extract_text(path: Path) -> Optional[str]:
    ext = path.suffix.lower()
    try:
        if ext == ".docx":
            from docx import Document
            doc = Document(str(path))
            parts = [p.text for p in doc.paragraphs]
            for t in doc.tables:
                for row in t.rows:
                    parts += [c.text for c in row.cells]
            return "\n".join(parts + [doc.core_properties.keywords or ""])
        if ext == ".pptx":
            from pptx import Presentation
            deck = Presentation(str(path))
            parts = []
            for slide in deck.slides:
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False):
                        parts.append(shape.text_frame.text)
                    if getattr(shape, "has_table", False):
                        parts += [c.text for row in shape.table.rows for c in row.cells]
            return "\n".join(parts)
        if ext == ".xlsx":
            from openpyxl import load_workbook
            book = load_workbook(str(path))
            parts = []
            for ws in book.worksheets:
                for row in ws.iter_rows(values_only=True):
                    for value in row:
                        if value is None or (isinstance(value, str) and value.startswith("=")):
                            continue
                        parts.append(str(value))
            return "\n".join(parts)
        if ext == ".pdf":
            import pymupdf
            with pymupdf.open(str(path)) as document:
                return "\n".join(page.get_text() for page in document)
        if ext == ".md":
            return path.read_text(encoding="utf-8")
    except Exception as error:
        logger.warning(f"could not read back {path.name}: {error}")
    return None


_HEX_ID = re.compile(r"\b[0-9a-f]{12,}(?::\d+)*\b")


def consistency(obj: AnalysisObject, path: Path) -> Dict[str, Any]:
    """The file belongs to this analysis and states no number the analysis does not hold."""
    text = extract_text(path)
    if text is None:
        return {"checked": False, "ok": False, "reason": "the file could not be read back"}
    text = _HEX_ID.sub(" ", text)           # document ids and content hashes are not figures
    allowed = analysis_mod.allowed_numbers(obj)
    created = time.localtime(obj.created_at)
    for value in (created.tm_year, created.tm_mon, created.tm_mday, created.tm_hour, created.tm_min):
        allowed.add(str(value))
    unknown = []
    for m in analysis_mod._NUMBER.finditer(text):
        raw = m.group(0)
        try:
            value = float(raw.replace(",", ""))
        except ValueError:
            continue
        if abs(value) <= 12 and value == int(value):
            continue
        if analysis_mod._canon(value) not in allowed and analysis_mod._canon(abs(value)) not in allowed:
            unknown.append(raw)
    id_present = obj.analysis_id in text
    rec_present = all(r["id"] in text for r in obj.recommendations[:3]) if obj.recommendations and \
        path.suffix.lower() != ".pdf" else True
    ok = id_present and not unknown and rec_present
    return {"checked": True, "ok": ok, "analysis_id_present": id_present,
            "unsupported_numbers": sorted(set(unknown))[:20], "recommendations_present": rec_present}


# ----------------------------------------------------------------- create ---

_WRITERS = {"xlsx": lambda o, p: _xlsx(o, p), "action_plan": lambda o, p: _xlsx(o, p, action_plan_only=True),
            "docx": lambda o, p: _docx(o, p), "brief": lambda o, p: _docx(o, p, brief=True),
            "pptx": _pptx, "pdf": _pdf, "md": _markdown}


def create(obj: AnalysisObject, fmt: str) -> Dict[str, Any]:
    fmt = (fmt or "").lower().strip()
    if fmt not in _WRITERS:
        return {"ok": False, "format": fmt, "error": f"unknown format - one of {', '.join(FORMATS)}"}
    if not obj.evidence and not obj.facts:
        return {"ok": False, "format": fmt, "error": "the analysis has no evidence - nothing to put in a file"}
    path = target_path(obj, fmt)
    try:
        _WRITERS[fmt](obj, path)
    except Exception as error:
        logger.exception(f"{fmt} renderer failed")
        return {"ok": False, "format": fmt, "error": f"could not write the {fmt} file: {error}"}
    kind = path.suffix.lstrip(".")
    if kind == "md":
        try:
            body = path.read_text(encoding="utf-8")
            checks = {"exists": True, "utf8": True, "has_text": len(body.strip()) > 40}
        except Exception as error:
            body, checks = "", {"exists": path.exists(), "utf8": False, "has_text": False}
        file_check = {"ok": all(checks.values()), "checks": checks, "detail": f"{path.name}: {len(body)} characters"}
    else:
        try:
            from core.orchestration.validators import validate
            report = validate(kind, str(path))
            file_check = {"ok": bool(report.ok), "checks": dict(report.checks), "detail": report.detail}
        except Exception as error:
            file_check = {"ok": False, "checks": {}, "detail": f"validator unavailable: {error}"}
    match = consistency(obj, path)
    verified = bool(file_check["ok"] and match["ok"] and obj.verified)
    artifact_id = None
    if _output_override.get() is None:          # evaluation output is not a user artifact
        try:
            from core.artifacts import register_artifact
            artifact_id = register_artifact(str(path), kind=kind, creator_agent="ArtifactAgent",
                                            request=obj.question[:500], model=obj.narrative_source,
                                            tool="knowledge_fabric.renderers",
                                            sources=[f"analysis:{obj.analysis_id}"]
                                            + sorted({e["citation"] for e in obj.evidence})[:40],
                                            verified=verified)
        except Exception as error:
            logger.warning(f"artifact not registered: {error}")
    return {"ok": True, "format": fmt, "path": str(path), "bytes": path.stat().st_size, "verified": verified,
            "file_check": file_check, "consistency": match, "analysis_id": obj.analysis_id,
            "artifact_id": artifact_id}
