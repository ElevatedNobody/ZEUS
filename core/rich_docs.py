"""
Rich visual document builders for ZEUS.

Makes Word / Excel look executive-grade: colored bands, KPI-style callouts,
charts, multi-sheet workbooks, cover pages. Pure local, no cloud.
"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

from core.logger import get_logger

logger = get_logger("rich_docs")

NAVY = "10263F"
TEAL = "0E7490"
TEAL_LIGHT = "E0F7FA"
INK = "16222E"
MUTED = "62717F"
BAND = "EEF3F8"
WHITE = "FFFFFF"
AMBER = "B26A00"
GREEN = "1B7A46"
RED = "B3261E"


def _hex_rgb(h: str):
    h = h.lstrip("#")
    return tuple(int(h[i:i+2], 16) for i in (0, 2, 4))


# ------------------------------------------------------------------ WORD --

def polish_word_document(
    topic: str,
    sections: List[tuple],
    out_path: Path,
    *,
    subtitle: str = "",
    callouts: Optional[List[str]] = None,
) -> Path:
    """Build a visually strong .docx with cover band, styled headings, callouts."""
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor, Cm, Twips
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn, nsmap
    from docx.oxml import OxmlElement

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)

    def _set_run_font(run, name="Calibri", size=11, bold=False, color=INK, italic=False):
        run.font.name = name
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        r, g, b = _hex_rgb(color)
        run.font.color.rgb = RGBColor(r, g, b)
        try:
            run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
        except Exception:
            pass

    def _shade_cell(cell, hex_color: str):
        tc = cell._tePr if hasattr(cell, "_tePr") else cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), hex_color.lstrip("#"))
        shd.set(qn("w:val"), "clear")
        tcPr.append(shd)

    # --- Cover table band ---
    cover = doc.add_table(rows=3, cols=1)
    cover.autofit = True
    c0 = cover.cell(0, 0)
    _shade_cell(c0, NAVY)
    p = c0.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("  ZEUS")
    _set_run_font(run, size=12, bold=True, color=TEAL)
    run2 = p.add_run("  ·  DOCUMENT")
    _set_run_font(run2, size=10, color="9FB3C6")

    c1 = cover.cell(1, 0)
    _shade_cell(c1, NAVY)
    p = c1.paragraphs[0]
    run = p.add_run(f"  {(topic or 'Document').strip().title()}")
    _set_run_font(run, size=26, bold=True, color=WHITE)

    c2 = cover.cell(2, 0)
    _shade_cell(c2, TEAL)
    p = c2.paragraphs[0]
    meta = subtitle or f"Prepared {datetime.now().strftime('%B %d, %Y')}"
    run = p.add_run(f"  {meta}")
    _set_run_font(run, size=10, color=WHITE)

    doc.add_paragraph()

    # Section roadmap under cover
    if sections:
        road = doc.add_paragraph()
        run = road.add_run("In this document")
        _set_run_font(run, size=11, bold=True, color=TEAL)
        for i, (heading, _) in enumerate(sections, 1):
            if not heading:
                continue
            p = doc.add_paragraph(style="List Number")
            run = p.add_run(str(heading))
            _set_run_font(run, size=11, color=INK)
        doc.add_paragraph()

    # Callout strip
    if callouts:
        t = doc.add_table(rows=1, cols=min(3, len(callouts)))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, text in enumerate(callouts[:3]):
            cell = t.cell(0, i)
            _shade_cell(cell, TEAL_LIGHT)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(text)[:120])
            _set_run_font(run, size=10, bold=True, color=TEAL)
        doc.add_paragraph()

    # Body sections
    for heading, body in sections:
        if heading:
            h = doc.add_heading(str(heading), level=2)
            for run in h.runs:
                _set_run_font(run, size=14, bold=True, color=TEAL)
        for block in re.split(r"\n\s*\n", (body or "").strip()):
            block = block.strip()
            if not block:
                continue
            if re.match(r"^(\d+[\.\)]\s+|[-•]\s+)", block):
                for line in block.split("\n"):
                    line = line.strip()
                    if not line:
                        continue
                    p = doc.add_paragraph(style="List Bullet")
                    clean = re.sub(r"^(\d+[\.\)]\s+|[-•]\s+)", "", line)
                    run = p.add_run(clean)
                    _set_run_font(run, size=11, color=INK)
            else:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.space_after = Pt(10)
                p.paragraph_format.line_spacing = 1.35
                run = p.add_run(block)
                _set_run_font(run, size=11, color=INK)

    # Footer line
    doc.add_paragraph()
    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = foot.add_run("— Generated by ZEUS  ·  Offline · Local AI —")
    _set_run_font(run, size=9, italic=True, color=MUTED)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


# ----------------------------------------------------------------- EXCEL --

def polish_spreadsheet(
    topic: str,
    headers: List[str],
    rows: List[Dict[str, Any]],
    out_path: Path,
) -> Path:
    """Multi-sheet workbook: Cover + Data (chart) + Summary."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import BarChart, PieChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.formatting.rule import ColorScaleRule

    wb = Workbook()

    thin = Border(
        left=Side(style="thin", color="D7E0E8"),
        right=Side(style="thin", color="D7E0E8"),
        top=Side(style="thin", color="D7E0E8"),
        bottom=Side(style="thin", color="D7E0E8"),
    )
    navy_fill = PatternFill("solid", fgColor=NAVY)
    teal_fill = PatternFill("solid", fgColor=TEAL)
    band_fill = PatternFill("solid", fgColor=BAND)
    white_font = Font(name="Calibri", color=WHITE, bold=True, size=12)
    title_font = Font(name="Calibri", color=WHITE, bold=True, size=20)
    header_font = Font(name="Calibri", color=WHITE, bold=True, size=11)
    body_font = Font(name="Calibri", color=INK, size=11)
    muted_font = Font(name="Calibri", color=MUTED, size=10, italic=True)

    # ---- Cover sheet ----
    cover = wb.active
    cover.title = "Cover"
    cover.sheet_view.showGridLines = False
    for col in range(1, 10):
        cover.column_dimensions[get_column_letter(col)].width = 14
    for r in range(1, 5):
        for c in range(1, 9):
            cell = cover.cell(r, c, value="")
            cell.fill = navy_fill
    cover.merge_cells("B2:G2")
    cover["B2"] = "ZEUS"
    cover["B2"].font = Font(name="Calibri", color=TEAL, bold=True, size=14)
    cover.merge_cells("B3:G3")
    cover["B3"] = (topic or "Workbook").strip().title()
    cover["B3"].font = title_font
    cover.merge_cells("B4:G4")
    cover["B4"] = f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')}  ·  {len(rows)} rows  ·  {len(headers)} columns"
    cover["B4"].font = Font(name="Calibri", color="9FB3C6", size=11)

    cover.merge_cells("B6:G6")
    cover["B6"] = "CONTENTS"
    cover["B6"].font = Font(name="Calibri", color=TEAL, bold=True, size=12)
    cover["B7"] = "1. Data — full table with chart"
    cover["B8"] = "2. Cover — this page"
    cover["B9"] = "3. Summary — quick stats"
    for r in range(7, 10):
        cover.cell(r, 2).font = body_font

    # KPI boxes on cover
    numeric_cols = []
    for h in headers:
        vals = []
        for row in rows:
            v = row.get(h)
            if isinstance(v, (int, float)):
                vals.append(float(v))
            else:
                try:
                    vals.append(float(str(v).replace(",", "")))
                except Exception:
                    pass
        if vals:
            numeric_cols.append((h, vals))

    kpi_start = 11
    for i, (h, vals) in enumerate(numeric_cols[:3]):
        c = 2 + i * 2
        cover.cell(kpi_start, c).fill = teal_fill
        cover.cell(kpi_start, c).value = h[:18]
        cover.cell(kpi_start, c).font = white_font
        cover.cell(kpi_start + 1, c).value = round(sum(vals) / len(vals), 2) if vals else 0
        cover.cell(kpi_start + 1, c).font = Font(name="Calibri", color=NAVY, bold=True, size=16)
        cover.cell(kpi_start + 2, c).value = f"avg of {len(vals)} values"
        cover.cell(kpi_start + 2, c).font = muted_font

    # ---- Data sheet ----
    # The DATA is the first and active sheet with its headers in row 1, so the
    # workbook opens on the table and behaves like any spreadsheet (filters,
    # frozen header row). The cover and notes follow it.
    data = wb.create_sheet("Data", 0)
    wb.active = 0
    data.sheet_view.showGridLines = False

    money_words = ("cost", "price", "amount", "revenue", "budget", "salary", "spend",
                   "expense", "income", "profit", "value (", "usd", "inr", "eur", "$", "₹")
    money_cols = {h for h in headers if any(w in str(h).lower() for w in money_words)}

    for col, h in enumerate(headers, 1):
        cell = data.cell(1, col, value=h)
        cell.font = header_font
        cell.fill = navy_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin
    data.row_dimensions[1].height = 22
    data.freeze_panes = "A2"
    data.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{1 + len(rows)}"

    for r_idx, row in enumerate(rows, 2):
        for c_idx, h in enumerate(headers, 1):
            val = row.get(h, "")
            cell = data.cell(r_idx, c_idx, value=val)
            cell.font = body_font
            cell.border = thin
            cell.alignment = Alignment(vertical="center")
            if r_idx % 2 == 1:
                cell.fill = band_fill
            if isinstance(val, (int, float)):
                if h in money_cols or isinstance(val, float):
                    cell.number_format = "#,##0.00"
                else:
                    cell.number_format = "#,##0"

    for col, h in enumerate(headers, 1):
        width = max(12, min(28, len(str(h)) + 4))
        for row in rows:
            width = max(width, min(28, len(str(row.get(h, ""))) + 2))
        data.column_dimensions[get_column_letter(col)].width = width

    # Color scale on first numeric column
    if numeric_cols:
        first_h = numeric_cols[0][0]
        col_idx = headers.index(first_h) + 1
        col_letter = get_column_letter(col_idx)
        data.conditional_formatting.add(
            f"{col_letter}2:{col_letter}{1 + len(rows)}",
            ColorScaleRule(
                start_type="min", start_color="F8696B",
                mid_type="percentile", mid_value=50, mid_color="FFEB84",
                end_type="max", end_color="63BE7B",
            ),
        )

    # Chart
    if numeric_cols and rows:
        label_col = 1
        for i, h in enumerate(headers):
            if h != numeric_cols[0][0]:
                label_col = i + 1
                break
        value_col = headers.index(numeric_cols[0][0]) + 1
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = str(numeric_cols[0][0])[:40]
        chart.y_axis.title = None
        chart.x_axis.title = None
        data_ref = Reference(data, min_col=value_col, min_row=1, max_row=1 + len(rows))
        cats = Reference(data, min_col=label_col, min_row=2, max_row=1 + len(rows))
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats)
        chart.shape = 4
        chart.width = 15
        chart.height = 8
        data.add_chart(chart, f"A{3 + len(rows)}")

        # Pie if few categories
        if len(rows) <= 8:
            pie = PieChart()
            pie.title = f"Share — {numeric_cols[0][0][:30]}"
            labels = Reference(data, min_col=label_col, min_row=2, max_row=1 + len(rows))
            pie_data = Reference(data, min_col=value_col, min_row=1, max_row=1 + len(rows))
            pie.add_data(pie_data, titles_from_data=True)
            pie.set_categories(labels)
            pie.dataLabels = DataLabelList()
            pie.dataLabels.showPercent = True
            pie.width = 12
            pie.height = 8
            data.add_chart(pie, f"J{3 + len(rows)}")

    # ---- Summary sheet ----
    summary = wb.create_sheet("Summary")
    summary.sheet_view.showGridLines = False
    summary["A1"] = "SUMMARY"
    summary["A1"].font = Font(name="Calibri", color=WHITE, bold=True, size=16)
    summary["A1"].fill = navy_fill
    summary.merge_cells("A1:D1")
    summary.row_dimensions[1].height = 28
    summary["A3"] = "Column"
    summary["B3"] = "Count"
    summary["C3"] = "Sum / Total"
    summary["D3"] = "Average"
    for col in range(1, 5):
        summary.cell(3, col).font = header_font
        summary.cell(3, col).fill = teal_fill
        summary.cell(3, col).border = thin

    r = 4
    for h, vals in numeric_cols:
        summary.cell(r, 1, h).font = body_font
        summary.cell(r, 2, len(vals)).font = body_font
        summary.cell(r, 3, round(sum(vals), 2)).font = body_font
        summary.cell(r, 4, round(sum(vals) / len(vals), 2)).font = body_font
        for c in range(1, 5):
            summary.cell(r, c).border = thin
            if r % 2 == 0:
                summary.cell(r, c).fill = band_fill
        r += 1

    if not numeric_cols:
        summary["A4"] = "No numeric columns detected — data sheet has the full table."
        summary["A4"].font = muted_font

    for col in range(1, 5):
        summary.column_dimensions[get_column_letter(col)].width = 18

    # Notes sheet - professional touch
    notes = wb.create_sheet("Notes")
    notes.sheet_view.showGridLines = False
    notes["A1"] = "NOTES & LEGEND"
    notes["A1"].font = Font(name="Calibri", color=WHITE, bold=True, size=16)
    notes["A1"].fill = navy_fill
    notes.merge_cells("A1:C1")
    notes.row_dimensions[1].height = 28
    notes["A3"] = "How to use this workbook"
    notes["A3"].font = Font(name="Calibri", color=TEAL, bold=True, size=12)
    notes["A4"] = "1. Data — full table, filters, conditional colour, charts"
    notes["A5"] = "2. Cover — title and quick KPI averages"
    notes["A6"] = "3. Summary — count, sum and average for numeric columns"
    notes["A7"] = "4. Notes — this page"
    for r in range(4, 8):
        notes.cell(r, 1).font = body_font
    notes["A9"] = "Generated offline by ZEUS. Charts update if you change values on the Data sheet."
    notes["A9"].font = muted_font
    notes.column_dimensions["A"].width = 70

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))
    return out_path
