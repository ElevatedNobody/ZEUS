"""
Phase C - local industrial document extraction pipeline.

    INPUT DOCUMENT
        v
    FILE TYPE DETECTION
        v
    DIGITAL TEXT EXTRACTION  ---or---  SCANNED DOCUMENT DETECTION
        v                                   v
        |                                  OCR  (PaddleOCR -> Tesseract -> EasyOCR)
        +-----------------+-----------------+
                          v
                   LAYOUT ANALYSIS
                          v
                   TABLE EXTRACTION
                          v
                   IMAGE EXTRACTION
                          v
                   METADATA EXTRACTION
                          v
             (handed to extraction.py for normalisation,
              chunking, embedding and the knowledge base)

Layering rules (never collapsed into "just OCR everything"):

  1. Native/digital text is always tried first - it is exact, fast and carries
     no OCR confidence penalty.
  2. PyMuPDF is the primary PDF engine; pypdf is the fallback if PyMuPDF is
     not installed.
  3. OCR runs only for pages that have no usable digital text layer.
  4. PaddleOCR provides document structure/layout when available.
  5. Tesseract remains the reliable OCR fallback - never removed.
  6. The vision model is used only for pages/images that still need
     multimodal understanding, and anything it produces is marked as an
     AI inference, never as document text.

Everything is local. No network calls except loopback to Ollama, and only
when vision analysis is explicitly requested.
"""

from __future__ import annotations

import hashlib
import io
import os
import re
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from core.logger import get_logger
from industrial import ASSETS_DIR
from industrial import ocr_engines

logger = get_logger("industrial.pipeline")

# --------------------------------------------------------------------------- types

TEXT_EXTENSIONS = {".txt", ".md", ".log", ".rst", ".ini", ".cfg", ".json", ".yaml", ".yml"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}
SUPPORTED_EXTENSIONS = ({".pdf", ".docx", ".xlsx", ".xlsm", ".pptx", ".csv"}
                        | TEXT_EXTENSIONS | IMAGE_EXTENSIONS)
UNSTRUCTURED_EXTENSIONS = {".html", ".htm", ".eml", ".msg", ".rtf", ".odt", ".epub", ".xml"}

# Rows read from one CSV file or one sheet. A daily ledger for 37 years is ~13,500
# rows; above this bound the file is cut and the document carries a warning saying so.
MAX_TABLE_ROWS = 100_000

_MAGIC = [
    (b"%PDF-", "pdf"),
    (b"PK\x03\x04", "zip_container"),   # docx/xlsx/pptx
    (b"\x89PNG\r\n\x1a\n", "png"),
    (b"\xff\xd8\xff", "jpeg"),
    (b"II*\x00", "tiff"),
    (b"MM\x00*", "tiff"),
    (b"\xd0\xcf\x11\xe0", "ole"),       # legacy doc/xls/ppt
]


@dataclass
class DocElement:
    """One extracted piece of a document, with its provenance intact."""
    element_type: str = "text"            # text|heading|table|figure|list|header|footer|sheet|slide
    text: str = ""
    page_number: Optional[int] = None
    section: Optional[str] = None
    heading: Optional[str] = None
    order: int = 0
    bbox: Optional[List[float]] = None
    extraction_method: str = "unknown"
    ocr_confidence: Optional[float] = None
    ocr_engine: Optional[str] = None
    is_inference: bool = False            # True only for vision-model output
    source_location: Optional[str] = None
    extra: Dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> dict:
        d = dict(self.__dict__)
        d["extra"] = dict(self.extra)
        return d


@dataclass
class DocAsset:
    asset_type: str                        # image|figure|page_render
    path: str
    page_number: Optional[int] = None
    caption: str = ""
    extraction_method: str = ""


@dataclass
class ExtractionOptions:
    ocr: bool = True
    layout: bool = True
    tables: bool = True
    images: bool = True
    vision: bool = False                  # requires Ollama + a vision model
    max_pages: int = 400
    ocr_dpi: int = 200
    max_images_per_doc: int = 40
    min_image_bytes: int = 6_000          # skip logos/rules
    ocr_language: str = "en"
    native_text_min_chars: int = 40       # below this a PDF page counts as scanned
    save_assets: bool = True


@dataclass
class ExtractedDocument:
    source_path: str
    filename: str
    file_type: str
    page_count: int = 0
    elements: List[DocElement] = field(default_factory=list)
    assets: List[DocAsset] = field(default_factory=list)
    raw_metadata: Dict[str, Any] = field(default_factory=dict)
    extraction_methods: List[str] = field(default_factory=list)
    scanned_pages: List[int] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)
    error: Optional[str] = None
    file_hash: str = ""
    duration_s: float = 0.0

    @property
    def text(self) -> str:
        return "\n\n".join(e.text for e in self.elements if e.text and not e.is_inference)

    @property
    def is_scanned(self) -> bool:
        return bool(self.scanned_pages)

    def mean_ocr_confidence(self) -> Optional[float]:
        vals = [e.ocr_confidence for e in self.elements if e.ocr_confidence is not None]
        return sum(vals) / len(vals) if vals else None

    def summary(self) -> dict:
        return {
            "filename": self.filename, "file_type": self.file_type, "pages": self.page_count,
            "elements": len(self.elements),
            "tables": sum(1 for e in self.elements if e.element_type == "table"),
            "figures": sum(1 for e in self.elements if e.element_type == "figure"),
            "assets": len(self.assets), "scanned_pages": self.scanned_pages,
            "extraction_methods": sorted(set(self.extraction_methods)),
            "mean_ocr_confidence": self.mean_ocr_confidence(),
            "warnings": self.warnings, "error": self.error,
        }


# ------------------------------------------------------------- file type detection

def detect_file_type(path: Path) -> str:
    """Extension first, magic bytes to confirm/correct. Returns a normalised
    type name: pdf, docx, xlsx, pptx, csv, text, image, or unknown."""
    ext = path.suffix.lower()
    head = b""
    try:
        with open(path, "rb") as fh:
            head = fh.read(16)
    except OSError:
        pass

    magic = None
    for sig, name in _MAGIC:
        if head.startswith(sig):
            magic = name
            break

    if ext == ".pdf" or magic == "pdf":
        return "pdf"
    if ext == ".docx":
        return "docx"
    if ext in (".xlsx", ".xlsm"):
        return "xlsx"
    if ext == ".pptx":
        return "pptx"
    if ext == ".csv":
        return "csv"
    if ext in IMAGE_EXTENSIONS or magic in ("png", "jpeg", "tiff"):
        return "image"
    if ext in TEXT_EXTENSIONS:
        return "text"
    if ext in UNSTRUCTURED_EXTENSIONS:
        return "unstructured"
    if magic == "ole":
        return "legacy_office"
    return "unknown"


def file_hash(path: Path, blocks: int = 2048) -> str:
    h = hashlib.sha256()
    try:
        with open(path, "rb") as fh:
            for chunk in iter(lambda: fh.read(65536), b""):
                h.update(chunk)
                blocks -= 1
                if blocks <= 0:
                    break
    except OSError:
        return ""
    return h.hexdigest()[:32]


def make_document_id(path: Path) -> str:
    """Stable id from path + content hash, so re-ingesting the same file
    replaces its chunks instead of duplicating them, while an edited file
    becomes a new revision of the record."""
    return hashlib.sha256(f"{path.resolve()}::{file_hash(path)}".encode("utf-8")).hexdigest()[:24]


# ------------------------------------------------------------------------ PDF

def _load_pymupdf():
    try:
        import pymupdf  # PyMuPDF >= 1.24 preferred import name
        return pymupdf
    except ImportError:
        try:
            import fitz
            return fitz
        except ImportError:
            return None


def _page_image(page, dpi: int):
    from PIL import Image
    pix = page.get_pixmap(dpi=dpi)
    return Image.open(io.BytesIO(pix.tobytes("png")))


def _text_table_is_plausible(rows: List[List[str]], min_fill: float = 0.65) -> bool:
    """Quality gate for whitespace-detected tables: a real table is densely
    filled and has a consistent number of populated columns per row."""
    if len(rows) < 2:
        return False
    cells = sum(len(r) for r in rows)
    filled = sum(1 for r in rows for c in r if c)
    if not cells or filled / cells < min_fill:
        return False
    widths = [sum(1 for c in r if c) for r in rows]
    widths = [w for w in widths if w]
    if len(widths) < 2 or max(widths) < 2:
        return False
    modal = max(set(widths), key=widths.count)
    return widths.count(modal) / len(widths) >= 0.7


def _extract_pdf_tables(page, page_no: int, order_start: int) -> List[DocElement]:
    """PyMuPDF's table finder. Returns table elements rendered as pipe-separated
    text so the content stays searchable while the grid stays readable."""
    elements: List[DocElement] = []
    method = "pymupdf_table"
    try:
        finder = page.find_tables()
        tables = list(getattr(finder, "tables", finder) or [])
    except Exception as e:
        logger.debug(f"table detection unavailable on page {page_no}: {e}")
        return elements

    if not tables:
        # Many industrial forms and reports use borderless tables, which the
        # default line-based finder cannot see. Retry with the whitespace/text
        # strategy before giving up.
        try:
            finder = page.find_tables(strategy="text")
            tables = list(getattr(finder, "tables", finder) or [])
            if tables:
                method = "pymupdf_table_text"
        except Exception:
            tables = []

    for i, table in enumerate(tables):
        try:
            rows = table.extract()
        except Exception:
            continue
        if not rows:
            continue
        cleaned = [[("" if c is None else str(c).replace("\n", " ").strip()) for c in row]
                   for row in rows]
        if not any(any(c for c in row) for row in cleaned):
            continue
        if method == "pymupdf_table_text" and not _text_table_is_plausible(cleaned):
            # The whitespace strategy will happily call an entire page of prose
            # a table. A bogus table is worse than a missed one - it would be
            # cited as a measurement grid - so reject low-density results.
            continue
        text = "\n".join(" | ".join(row) for row in cleaned)
        bbox = list(getattr(table, "bbox", []) or []) or None
        elements.append(DocElement(
            element_type="table", text=text, page_number=page_no, order=order_start + i,
            bbox=[float(v) for v in bbox] if bbox else None,
            extraction_method=method,
            source_location=f"page={page_no}&table={i + 1}",
            extra={"rows": len(cleaned), "columns": max(len(r) for r in cleaned),
                   "grid": cleaned[:60]},
        ))
    return elements


def _extract_pdf_images(doc, page, page_no: int, out_dir: Path, opts: ExtractionOptions,
                        budget: List[int]) -> Tuple[List[DocAsset], List[DocElement]]:
    assets: List[DocAsset] = []
    elements: List[DocElement] = []
    try:
        images = page.get_images(full=True)
    except Exception:
        return assets, elements

    for idx, info in enumerate(images):
        if budget[0] <= 0:
            break
        xref = info[0]
        try:
            raw = doc.extract_image(xref)
        except Exception:
            continue
        data = raw.get("image")
        if not data or len(data) < opts.min_image_bytes:
            continue
        ext = raw.get("ext", "png")
        name = f"p{page_no:04d}_img{idx + 1}.{ext}"
        target = out_dir / name
        if opts.save_assets:
            try:
                out_dir.mkdir(parents=True, exist_ok=True)
                target.write_bytes(data)
            except OSError as e:
                logger.warning(f"could not save extracted image: {e}")
                continue
        budget[0] -= 1
        assets.append(DocAsset(asset_type="image", path=str(target), page_number=page_no,
                               extraction_method="pymupdf_image"))
        elements.append(DocElement(
            element_type="figure", text=f"[Figure on page {page_no}: {name}]",
            page_number=page_no, order=10_000 + idx, extraction_method="pymupdf_image",
            source_location=f"page={page_no}&image={idx + 1}",
            extra={"asset_path": str(target), "bytes": len(data)},
        ))
    return assets, elements


def extract_pdf(path: Path, opts: ExtractionOptions, doc_out: ExtractedDocument) -> None:
    mupdf = _load_pymupdf()
    if mupdf is None:
        _extract_pdf_pypdf_fallback(path, doc_out)
        return

    asset_dir = ASSETS_DIR / make_document_id(path)
    image_budget = [opts.max_images_per_doc]

    try:
        doc = mupdf.open(str(path))
    except Exception as e:
        doc_out.error = f"PyMuPDF could not open the PDF: {e}"
        return

    try:
        doc_out.page_count = doc.page_count
        try:
            doc_out.raw_metadata.update({k: v for k, v in (doc.metadata or {}).items() if v})
        except Exception:
            pass
        try:
            toc = doc.get_toc() or []
            if toc:
                doc_out.raw_metadata["outline"] = [
                    {"level": lvl, "title": title, "page": pno} for lvl, title, pno in toc[:200]]
        except Exception:
            pass

        order = 0
        current_section: Optional[str] = None

        for page_index in range(min(doc.page_count, opts.max_pages)):
            page_no = page_index + 1
            page = doc.load_page(page_index)

            # 1) native/digital text layer -------------------------------------
            try:
                native = page.get_text("text") or ""
            except Exception:
                native = ""

            if len(native.strip()) >= opts.native_text_min_chars:
                doc_out.extraction_methods.append("native_pdf_text")
                blocks = []
                try:
                    blocks = page.get_text("blocks") or []
                except Exception:
                    pass
                if blocks:
                    blocks.sort(key=lambda b: (round(b[1], 1), round(b[0], 1)))
                    for b in blocks:
                        text = (b[4] or "").strip()
                        if not text:
                            continue
                        etype, heading = _classify_block(text, b, page)
                        if etype in ("heading", "title"):
                            current_section = text.strip()[:120]
                        doc_out.elements.append(DocElement(
                            element_type="heading" if etype in ("heading", "title") else etype,
                            text=text, page_number=page_no, section=current_section,
                            heading=heading, order=order,
                            bbox=[float(b[0]), float(b[1]), float(b[2]), float(b[3])],
                            extraction_method="native_pdf_text",
                            source_location=f"page={page_no}",
                        ))
                        order += 1
                else:
                    doc_out.elements.append(DocElement(
                        element_type="text", text=native.strip(), page_number=page_no,
                        section=current_section, order=order,
                        extraction_method="native_pdf_text",
                        source_location=f"page={page_no}"))
                    order += 1
            else:
                # 2) scanned page -> OCR ---------------------------------------
                doc_out.scanned_pages.append(page_no)
                if not opts.ocr:
                    doc_out.warnings.append(f"page {page_no} is scanned and OCR is disabled")
                    continue
                try:
                    image = _page_image(page, opts.ocr_dpi)
                except Exception as e:
                    doc_out.warnings.append(f"page {page_no}: could not render for OCR ({e})")
                    continue

                order = _ocr_page_into(doc_out, image, page_no, order, opts,
                                       section=current_section)

            # 3) tables --------------------------------------------------------
            if opts.tables:
                tbls = _extract_pdf_tables(page, page_no, order)
                if tbls:
                    doc_out.extraction_methods.append("pymupdf_table")
                    for t in tbls:
                        t.section = current_section
                    doc_out.elements.extend(tbls)
                    order += len(tbls)

            # 4) embedded images ----------------------------------------------
            if opts.images and image_budget[0] > 0:
                assets, figs = _extract_pdf_images(doc, page, page_no, asset_dir, opts, image_budget)
                if assets:
                    doc_out.extraction_methods.append("pymupdf_image")
                doc_out.assets.extend(assets)
                doc_out.elements.extend(figs)

            # 5) vision model for pages that still need multimodal reading -----
            if opts.vision and _page_needs_vision(doc_out, page_no):
                try:
                    image = _page_image(page, min(opts.ocr_dpi, 150))
                    desc = describe_with_vision(image, page_no)
                    if desc:
                        doc_out.extraction_methods.append("vision_model")
                        doc_out.elements.append(DocElement(
                            element_type="figure", text=desc, page_number=page_no,
                            section=current_section, order=order,
                            extraction_method="vision_model", is_inference=True,
                            source_location=f"page={page_no}",
                            extra={"note": "AI inference from the page image, not document text"}))
                        order += 1
                except Exception as e:
                    doc_out.warnings.append(f"page {page_no}: vision analysis failed ({e})")
    finally:
        try:
            doc.close()
        except Exception:
            pass


def _classify_block(text: str, block, page) -> Tuple[str, Optional[str]]:
    """Very cheap heading detection for digital PDFs: short lines that are
    numbered, all-caps or title-cased and sit alone in their block."""
    stripped = text.strip()
    if len(stripped) <= 120 and "\n" not in stripped:
        if re.match(r"^\d+(?:\.\d+)*[.)]?\s+\S", stripped) or (
                stripped.isupper() and len(stripped) > 3) or stripped.endswith(":"):
            return "heading", stripped[:120]
    return "text", None


def _page_needs_vision(doc_out: ExtractedDocument, page_no: int) -> bool:
    """A page needs the vision model when OCR produced very little text but the
    page carries figures - a drawing sheet, a photo page, a stamped form."""
    page_elements = [e for e in doc_out.elements if e.page_number == page_no]
    text_chars = sum(len(e.text) for e in page_elements
                     if e.element_type in ("text", "heading") and not e.is_inference)
    has_figure = any(e.element_type == "figure" for e in page_elements)
    return text_chars < 200 and (has_figure or page_no in doc_out.scanned_pages)


def _ocr_page_into(doc_out: ExtractedDocument, image, page_no: int, order: int,
                   opts: ExtractionOptions, section: Optional[str] = None) -> int:
    """Run OCR + layout analysis on one rendered page and append elements."""
    if opts.layout:
        regions = ocr_engines.analyse_layout(image)
    else:
        regions = []

    if regions:
        engine_name = ("paddleocr" if ocr_engines.paddle_structure_available()
                       else "layout+ocr")
        doc_out.extraction_methods.append(
            "paddleocr_structure" if ocr_engines.paddle_structure_available()
            else "heuristic_layout")
        current = section
        for r in regions:
            if not r.text.strip():
                continue
            etype = {"title": "heading", "text": "text", "list": "list", "table": "table",
                     "figure": "figure", "header": "header", "footer": "footer",
                     "reference": "text", "equation": "text"}.get(r.region_type, "text")
            if etype == "heading":
                current = r.text.strip()[:120]
            doc_out.elements.append(DocElement(
                element_type=etype, text=r.text.strip(), page_number=page_no, section=current,
                order=order, bbox=r.bbox, extraction_method=engine_name,
                ocr_confidence=round(r.confidence, 4) if r.confidence else None,
                ocr_engine=engine_name, source_location=f"page={page_no}",
                extra={"layout_region": r.region_type}))
            order += 1
        return order

    result = ocr_engines.run_ocr(image, lang=opts.ocr_language)
    if result.ok:
        doc_out.extraction_methods.append(f"ocr_{result.engine}")
        doc_out.elements.append(DocElement(
            element_type="text", text=result.text, page_number=page_no, section=section,
            order=order, extraction_method=f"ocr_{result.engine}",
            ocr_confidence=round(result.confidence, 4), ocr_engine=result.engine,
            source_location=f"page={page_no}"))
        return order + 1

    doc_out.warnings.append(
        f"page {page_no}: OCR produced no text ({result.error or 'empty result'})")
    return order


def _extract_pdf_pypdf_fallback(path: Path, doc_out: ExtractedDocument) -> None:
    try:
        from pypdf import PdfReader
    except ImportError:
        doc_out.error = ("no PDF engine available - install PyMuPDF "
                         "(pip install PyMuPDF) or pypdf")
        return
    try:
        reader = PdfReader(str(path))
        doc_out.page_count = len(reader.pages)
        doc_out.warnings.append("PyMuPDF not installed - using pypdf; no tables, "
                                "images or layout analysis on this document")
        doc_out.extraction_methods.append("pypdf_text")
        try:
            meta = reader.metadata or {}
            doc_out.raw_metadata.update({str(k).lstrip("/"): str(v) for k, v in meta.items() if v})
        except Exception:
            pass
        for i, page in enumerate(reader.pages):
            text = (page.extract_text() or "").strip()
            if not text:
                doc_out.scanned_pages.append(i + 1)
                continue
            doc_out.elements.append(DocElement(
                element_type="text", text=text, page_number=i + 1, order=i,
                extraction_method="pypdf_text", source_location=f"page={i + 1}"))
    except Exception as e:
        doc_out.error = f"pypdf failed: {e}"


# ----------------------------------------------------------------------- DOCX

def extract_docx(path: Path, opts: ExtractionOptions, doc_out: ExtractedDocument) -> None:
    try:
        from docx import Document
    except ImportError:
        doc_out.error = "python-docx not installed"
        return
    try:
        d = Document(str(path))
    except Exception as e:
        doc_out.error = f"could not open DOCX: {e}"
        return

    doc_out.extraction_methods.append("python_docx")
    doc_out.page_count = 0
    try:
        props = d.core_properties
        for key in ("title", "author", "subject", "category", "comments", "revision",
                    "last_modified_by", "keywords"):
            val = getattr(props, key, None)
            if val:
                doc_out.raw_metadata[key] = str(val)
        for key in ("created", "modified"):
            val = getattr(props, key, None)
            if val:
                doc_out.raw_metadata[key] = str(val)
    except Exception:
        pass

    order = 0
    section: Optional[str] = None
    for para in d.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        style = (para.style.name if para.style is not None else "") or ""
        is_heading = style.lower().startswith("heading") or style.lower() in ("title", "subtitle")
        if is_heading:
            section = text[:120]
        doc_out.elements.append(DocElement(
            element_type="heading" if is_heading else "text", text=text,
            section=section, heading=text[:120] if is_heading else None, order=order,
            extraction_method="python_docx", source_location=f"paragraph={order + 1}",
            extra={"style": style}))
        order += 1

    if opts.tables:
        for t_idx, table in enumerate(d.tables):
            rows = []
            for row in table.rows:
                rows.append([(cell.text or "").replace("\n", " ").strip() for cell in row.cells])
            if not any(any(c for c in r) for r in rows):
                continue
            doc_out.elements.append(DocElement(
                element_type="table",
                text="\n".join(" | ".join(r) for r in rows), section=section, order=order,
                extraction_method="python_docx_table",
                source_location=f"table={t_idx + 1}",
                extra={"rows": len(rows), "columns": max(len(r) for r in rows),
                       "grid": rows[:60]}))
            order += 1

    if opts.images:
        asset_dir = ASSETS_DIR / make_document_id(path)
        count = 0
        try:
            for rel in d.part.rels.values():
                if "image" not in rel.reltype or count >= opts.max_images_per_doc:
                    continue
                blob = rel.target_part.blob
                if len(blob) < opts.min_image_bytes:
                    continue
                ext = Path(rel.target_part.partname).suffix or ".png"
                target = asset_dir / f"docx_img{count + 1}{ext}"
                if opts.save_assets:
                    asset_dir.mkdir(parents=True, exist_ok=True)
                    target.write_bytes(blob)
                doc_out.assets.append(DocAsset(asset_type="image", path=str(target),
                                               extraction_method="python_docx_image"))
                count += 1
        except Exception as e:
            doc_out.warnings.append(f"DOCX image extraction skipped: {e}")


# ----------------------------------------------------------------------- XLSX

def extract_xlsx(path: Path, opts: ExtractionOptions, doc_out: ExtractedDocument) -> None:
    try:
        import openpyxl
    except ImportError:
        doc_out.error = "openpyxl not installed"
        return
    try:
        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    except Exception as e:
        doc_out.error = f"could not open workbook: {e}"
        return

    doc_out.extraction_methods.append("openpyxl")
    order = 0
    try:
        props = wb.properties
        for key in ("title", "creator", "lastModifiedBy", "category", "subject", "revision"):
            val = getattr(props, key, None)
            if val:
                doc_out.raw_metadata[key] = str(val)
    except Exception:
        pass

    for sheet in wb.worksheets:
        rows: List[List[str]] = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if c is None else str(c).strip() for c in row]
            if any(cells):
                rows.append(cells)
            if len(rows) > MAX_TABLE_ROWS:
                rows.pop()
                doc_out.warnings.append(f"sheet '{sheet.title}' truncated at {MAX_TABLE_ROWS} rows - "
                                        f"the rest of the sheet was NOT read")
                break
        if not rows:
            continue
        doc_out.elements.append(DocElement(
            element_type="table",
            text="\n".join(" | ".join(r) for r in rows), section=sheet.title,
            heading=sheet.title, order=order, extraction_method="openpyxl",
            source_location=f"sheet={sheet.title}",
            extra={"sheet": sheet.title, "rows": len(rows),
                   "columns": max(len(r) for r in rows), "grid": rows[:60]}))
        order += 1
    try:
        wb.close()
    except Exception:
        pass


# ----------------------------------------------------------------------- PPTX

def extract_pptx(path: Path, opts: ExtractionOptions, doc_out: ExtractedDocument) -> None:
    try:
        from pptx import Presentation
    except ImportError:
        doc_out.error = "python-pptx not installed"
        return
    try:
        prs = Presentation(str(path))
    except Exception as e:
        doc_out.error = f"could not open presentation: {e}"
        return

    doc_out.extraction_methods.append("python_pptx")
    doc_out.page_count = len(prs.slides)
    order = 0
    asset_dir = ASSETS_DIR / make_document_id(path)
    image_count = 0

    for s_idx, slide in enumerate(prs.slides, start=1):
        title = ""
        try:
            if slide.shapes.title is not None:
                title = (slide.shapes.title.text or "").strip()
        except Exception:
            pass
        if title:
            doc_out.elements.append(DocElement(
                element_type="heading", text=title, page_number=s_idx, section=title,
                heading=title, order=order, extraction_method="python_pptx",
                source_location=f"slide={s_idx}"))
            order += 1

        for shape in slide.shapes:
            try:
                if shape.has_text_frame:
                    text = "\n".join(p.text for p in shape.text_frame.paragraphs).strip()
                    if text and text != title:
                        doc_out.elements.append(DocElement(
                            element_type="text", text=text, page_number=s_idx,
                            section=title or None, order=order,
                            extraction_method="python_pptx",
                            source_location=f"slide={s_idx}"))
                        order += 1
                if getattr(shape, "has_table", False):
                    rows = [[(c.text or "").strip() for c in row.cells] for row in shape.table.rows]
                    doc_out.elements.append(DocElement(
                        element_type="table",
                        text="\n".join(" | ".join(r) for r in rows), page_number=s_idx,
                        section=title or None, order=order,
                        extraction_method="python_pptx_table",
                        source_location=f"slide={s_idx}",
                        extra={"rows": len(rows), "grid": rows[:60]}))
                    order += 1
                if opts.images and shape.shape_type is not None and hasattr(shape, "image"):
                    blob = shape.image.blob
                    if len(blob) >= opts.min_image_bytes and image_count < opts.max_images_per_doc:
                        target = asset_dir / f"slide{s_idx}_img{image_count + 1}.{shape.image.ext}"
                        if opts.save_assets:
                            asset_dir.mkdir(parents=True, exist_ok=True)
                            target.write_bytes(blob)
                        doc_out.assets.append(DocAsset(
                            asset_type="image", path=str(target), page_number=s_idx,
                            extraction_method="python_pptx_image"))
                        image_count += 1
            except Exception:
                continue

        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip()
                if notes:
                    doc_out.elements.append(DocElement(
                        element_type="text", text=notes, page_number=s_idx,
                        section=title or None, order=order,
                        extraction_method="python_pptx_notes",
                        source_location=f"slide={s_idx}&notes"))
                    order += 1
        except Exception:
            pass


# ------------------------------------------------------------------ CSV / TEXT

def extract_csv(path: Path, opts: ExtractionOptions, doc_out: ExtractedDocument) -> None:
    rows: List[List[str]] = []
    truncated = False
    try:
        import pandas as pd
        df = pd.read_csv(path, dtype=str, keep_default_na=False, on_bad_lines="skip",
                         nrows=MAX_TABLE_ROWS + 1, encoding_errors="replace")
        truncated = len(df) > MAX_TABLE_ROWS
        rows = [list(df.columns)] + df.head(MAX_TABLE_ROWS).astype(str).values.tolist()
        doc_out.extraction_methods.append("pandas_csv")
    except Exception:
        import csv as _csv
        rows = []
        try:
            with open(path, "r", encoding="utf-8", errors="replace", newline="") as fh:
                for i, row in enumerate(_csv.reader(fh)):
                    if i > MAX_TABLE_ROWS:
                        truncated = True
                        break
                    rows.append([str(c).strip() for c in row])
            doc_out.extraction_methods.append("csv_module")
        except Exception as e:
            doc_out.error = f"could not read CSV: {e}"
            return
    if truncated:
        doc_out.warnings.append(f"CSV truncated at {MAX_TABLE_ROWS} data rows - the rest of the file was NOT read")
    if not rows:
        doc_out.warnings.append("CSV is empty")
        return
    doc_out.elements.append(DocElement(
        element_type="table", text="\n".join(" | ".join(r) for r in rows),
        order=0, extraction_method=doc_out.extraction_methods[-1],
        source_location="table=1",
        extra={"rows": len(rows), "columns": max(len(r) for r in rows), "grid": rows[:60]}))


def extract_text(path: Path, opts: ExtractionOptions, doc_out: ExtractedDocument) -> None:
    try:
        content = path.read_text(encoding="utf-8", errors="replace")
    except OSError as e:
        doc_out.error = f"could not read file: {e}"
        return
    doc_out.extraction_methods.append("plain_text")

    # Split on markdown-style headings so section provenance survives.
    order, section = 0, None
    buffer: List[str] = []

    def flush():
        nonlocal order, buffer
        text = "\n".join(buffer).strip()
        buffer = []
        if text:
            doc_out.elements.append(DocElement(
                element_type="text", text=text, section=section, order=order,
                extraction_method="plain_text", source_location=f"block={order + 1}"))
            order += 1

    for line in content.splitlines():
        m = re.match(r"^(#{1,6})\s+(.*)$", line.strip())
        if m:
            flush()
            section = m.group(2).strip()[:120]
            doc_out.elements.append(DocElement(
                element_type="heading", text=section, section=section, heading=section,
                order=order, extraction_method="plain_text",
                source_location=f"heading={section[:40]}",
                extra={"level": len(m.group(1))}))
            order += 1
        else:
            buffer.append(line)
    flush()


# ---------------------------------------------------------------------- IMAGE

def extract_image(path: Path, opts: ExtractionOptions, doc_out: ExtractedDocument) -> None:
    try:
        from PIL import Image
        image = Image.open(path)
        image.load()
    except Exception as e:
        doc_out.error = f"could not open image: {e}"
        return

    doc_out.page_count = 1
    doc_out.raw_metadata.update({"width": image.width, "height": image.height,
                                 "mode": image.mode, "format": image.format or ""})
    if opts.ocr:
        _ocr_page_into(doc_out, image, 1, 0, opts)
    if opts.vision and _page_needs_vision(doc_out, 1):
        desc = describe_with_vision(image, 1)
        if desc:
            doc_out.extraction_methods.append("vision_model")
            doc_out.elements.append(DocElement(
                element_type="figure", text=desc, page_number=1, order=999,
                extraction_method="vision_model", is_inference=True,
                source_location="image=1",
                extra={"note": "AI inference from the image, not document text"}))
    doc_out.assets.append(DocAsset(asset_type="image", path=str(path), page_number=1,
                                   extraction_method="source_file"))


# -------------------------------------------------------------- UNSTRUCTURED

def extract_with_unstructured(path: Path, opts: ExtractionOptions,
                              doc_out: ExtractedDocument) -> None:
    """Local partitioning for formats the dedicated libraries above do not
    cover (HTML, EML, RTF, ODT, XML...). Forced into fully local mode: the
    'fast' strategy uses no hosted model and no remote inference, and
    telemetry/analytics env vars are pinned off before import."""
    os.environ.setdefault("SCARF_NO_ANALYTICS", "true")
    os.environ.setdefault("DO_NOT_TRACK", "true")
    os.environ.setdefault("UNSTRUCTURED_ANALYTICS", "false")
    try:
        from unstructured.partition.auto import partition
    except ImportError:
        doc_out.error = (f"no local handler for '{path.suffix}' "
                         f"(install 'unstructured' to partition this format)")
        return
    try:
        parts = partition(filename=str(path), strategy="fast")
    except Exception as e:
        doc_out.error = f"unstructured partitioning failed: {e}"
        return

    doc_out.extraction_methods.append("unstructured_fast")
    section = None
    for i, el in enumerate(parts):
        text = (getattr(el, "text", "") or "").strip()
        if not text:
            continue
        category = (getattr(el, "category", "") or el.__class__.__name__).lower()
        if "title" in category or "header" in category:
            etype, section = "heading", text[:120]
        elif "table" in category:
            etype = "table"
        elif "list" in category:
            etype = "list"
        else:
            etype = "text"
        meta = getattr(el, "metadata", None)
        page_no = getattr(meta, "page_number", None) if meta else None
        doc_out.elements.append(DocElement(
            element_type=etype, text=text, page_number=page_no, section=section, order=i,
            extraction_method="unstructured_fast",
            source_location=f"page={page_no}" if page_no else f"element={i + 1}"))


# --------------------------------------------------------------- VISION HOOK

VISION_PROMPT = (
    "This is a page from an industrial engineering document. Describe only what is "
    "visibly present: the kind of page (drawing, form, table, photograph), the visible "
    "layout, any equipment or components shown, and any labels you can read with "
    "confidence. Do not guess values you cannot read, and do not add engineering "
    "conclusions."
)


def vision_available() -> bool:
    try:
        from core.model_router import model_router
        from core.ollama_client import ollama_client
        if not ollama_client.is_available():
            return False
        model = model_router.route("vision")
        return any(m == model or m.startswith(model.split(":")[0])
                   for m in ollama_client.list_models())
    except Exception:
        return False


def describe_with_vision(image, page_no: int, prompt: str = "") -> str:
    """Ask the local vision model about a page image. Returns '' if the model
    is not available - never raises, never falls back to a text model."""
    try:
        import base64
        from core.model_router import model_router
        from core.ollama_client import ollama_client
        if not ollama_client.is_available():
            return ""
        buf = io.BytesIO()
        image.convert("RGB").save(buf, format="JPEG", quality=82)
        b64 = base64.b64encode(buf.getvalue()).decode("utf-8")
        model = model_router.route("vision")
        out = ollama_client.chat_vision(model, prompt or VISION_PROMPT, b64,
                                        purpose=f"Industrial page {page_no}")
        if not out or out.strip().startswith("["):
            return ""
        return f"[Vision model reading of page {page_no} - AI inference]\n{out.strip()}"
    except Exception as e:
        logger.warning(f"vision description failed: {e}")
        return ""


# ------------------------------------------------------------------ dispatcher

_EXTRACTORS = {
    "pdf": extract_pdf,
    "docx": extract_docx,
    "xlsx": extract_xlsx,
    "pptx": extract_pptx,
    "csv": extract_csv,
    "text": extract_text,
    "image": extract_image,
    "unstructured": extract_with_unstructured,
}


def extract_document(path: str | Path, options: Optional[ExtractionOptions] = None) -> ExtractedDocument:
    """Full local extraction of one document. Never raises: failures are
    reported on the returned object so a batch ingest can continue."""
    path = Path(path).expanduser()
    opts = options or ExtractionOptions()
    started = time.perf_counter()

    doc_out = ExtractedDocument(source_path=str(path), filename=path.name, file_type="unknown")
    if not path.is_file():
        doc_out.error = "file does not exist"
        return doc_out

    doc_out.file_type = detect_file_type(path)
    doc_out.file_hash = file_hash(path)
    doc_out.raw_metadata.setdefault("file_size_bytes", path.stat().st_size)
    doc_out.raw_metadata.setdefault("file_modified", time.strftime(
        "%Y-%m-%dT%H:%M:%S", time.localtime(path.stat().st_mtime)))

    extractor = _EXTRACTORS.get(doc_out.file_type)
    if extractor is None:
        if doc_out.file_type == "legacy_office":
            doc_out.error = ("legacy Office format (.doc/.xls/.ppt) - convert to the "
                             "modern format first; ZEUS does not run Office to read it")
        else:
            doc_out.error = f"unsupported file type '{path.suffix or doc_out.file_type}'"
        return doc_out

    try:
        extractor(path, opts, doc_out)
    except Exception as e:  # defensive - one bad file must not abort a batch
        logger.exception(f"extraction failed for {path}")
        doc_out.error = f"{type(e).__name__}: {e}"

    if doc_out.page_count == 0:
        pages = {e.page_number for e in doc_out.elements if e.page_number}
        doc_out.page_count = max(pages) if pages else (1 if doc_out.elements else 0)

    doc_out.extraction_methods = sorted(set(doc_out.extraction_methods))
    doc_out.duration_s = round(time.perf_counter() - started, 3)
    return doc_out


def pipeline_status() -> dict:
    """What the pipeline can do right now on this machine."""
    mupdf = _load_pymupdf()
    def has(mod: str) -> bool:
        import importlib.util
        try:
            return importlib.util.find_spec(mod) is not None
        except (ImportError, ValueError):
            return False
    return {
        "pdf_engine": ("PyMuPDF" if mupdf else ("pypdf (limited)" if has("pypdf") else "NONE")),
        "pdf_tables": bool(mupdf),
        "docx": has("docx"), "xlsx": has("openpyxl"), "pptx": has("pptx"),
        "csv": has("pandas") or True, "images": has("PIL"),
        "opencv_preprocessing": has("cv2"),
        "unstructured": has("unstructured"),
        "ocr": ocr_engines.engine_report(),
        "vision_model_ready": vision_available(),
        "supported_extensions": sorted(SUPPORTED_EXTENSIONS | UNSTRUCTURED_EXTENSIONS),
    }
