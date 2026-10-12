"""
Workbook builder: IndustrialAnalysis -> a real, multi-sheet .xlsx.

Structure, rather than one flat dump:

    Summary      the management view - headline, KPI band, actions, the ask
    <data>       one sheet per table block: filters, frozen header, totals
    Analysis     findings, a live severity profile and a chart
    Sources      where every number came from, and what is DEMO

Three rules that separate this from "export a table":

* **Real formulas.** A total is `=SUM(D8:D19)`, not a number Python worked out
  and pasted. The Analysis sheet counts severities with `COUNTIF` against its
  own findings and reaches across sheets with `COUNTA('Open alerts'!A8:A12)`,
  so if someone edits a row the summary follows. A reader can click a cell and
  see where the figure comes from.
* **Typed cells.** Numbers are written as numbers with a number format, dates
  as dates. A spreadsheet of strings looks the same and sorts wrong.
* **It is laid out like a person would lay it out.** Full-width title bands,
  section bands, banded rows, a highlighted total, wide columns for prose,
  gridlines off, print setup done, coloured sheet tabs.

The presentation grammar (title band, section bands, banded data rows, a
tinted total row, an assumptions block) was calibrated against a professional
reference workbook. None of its content, wording or business data is
reproduced - and where the reference hard-codes its calculated values, this
writes formulas instead.
"""

from __future__ import annotations

from typing import Any, Dict, List, Optional, Tuple

from core.branding import APP_TAGLINE as BRAND_TAGLINE
from core.logger import get_logger
from industrial.analysis import BAD, GOOD, IndustrialAnalysis, TableBlock, WARN

logger = get_logger("industrial.workbook")

# ------------------------------------------------------------- palette ----
INK = "1F2933"
MUTED_INK = "5A6B7B"
TITLE_BG = "10263F"          # full-width title band
TITLE_INK = "FFFFFF"
SUBTITLE_BG = "1C4E80"       # the strip under the title
HEADER_BG = "10263F"         # table header row
HEADER_INK = "FFFFFF"
SECTION_BG = "DCE7F1"        # section band inside a sheet
SECTION_INK = "10263F"
ACCENT = "0E7490"
BAND = "F2F6FA"              # banded data row
TOTAL_BG = "D8F0E2"
TOTAL_INK = "10603A"
NOTE_BG = "EAF2FB"
GOOD_BG, BAD_BG, WARN_BG = "E4F5E9", "FCE8E6", "FDF3DE"
GOOD_INK, BAD_INK, WARN_INK = "1B7A3D", "B42318", "8A5A00"
GRID = "C9D6E2"

TAB_SUMMARY = "10263F"
TAB_DATA = "0E7490"
TAB_ANALYSIS = "8A5A00"
TAB_SOURCES = "5A6B7B"

MAX_SHEET_NAME = 31


def _safe_sheet_name(name: str, used: set) -> str:
    cleaned = "".join(c for c in name if c not in "[]:*?/\\")[:MAX_SHEET_NAME].strip()
    cleaned = cleaned or "Sheet"
    candidate, index = cleaned, 2
    while candidate.lower() in used:
        suffix = f" {index}"
        candidate = cleaned[:MAX_SHEET_NAME - len(suffix)] + suffix
        index += 1
    used.add(candidate.lower())
    return candidate


def _looks_numeric(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _thin_border():
    from openpyxl.styles import Border, Side
    side = Side(style="thin", color=GRID)
    return Border(left=side, right=side, top=side, bottom=side)


class WorkbookBuilder:
    def __init__(self, analysis: IndustrialAnalysis):
        from openpyxl import Workbook
        self.analysis = analysis
        self.book = Workbook()
        self.book.remove(self.book.active)
        self._used_names: set = set()
        self._data_sheets: Dict[str, Tuple[str, TableBlock, int, int]] = {}

    # ------------------------------------------------------- primitives --

    @staticmethod
    def _merge(sheet, row: int, width: int):
        if width > 1:
            sheet.merge_cells(start_row=row, start_column=1, end_row=row,
                              end_column=width)

    def _band(self, sheet, row: int, text: str, width: int, fill: str, ink: str,
              size: float = 11, bold: bool = True, height: Optional[float] = None,
              italic: bool = False, align: str = "left") -> int:
        """A full-width coloured band. This is what makes a sheet read as a
        designed page rather than a CSV that happens to be in Excel."""
        from openpyxl.styles import Alignment, Font, PatternFill
        cell = sheet.cell(row=row, column=1, value=text)
        cell.font = Font(bold=bold, size=size, color=ink, italic=italic)
        cell.alignment = Alignment(horizontal=align, vertical="center",
                                   wrap_text=False, indent=1 if align == "left" else 0)
        for column in range(1, max(2, width) + 1):
            sheet.cell(row=row, column=column).fill = PatternFill("solid", fgColor=fill)
        self._merge(sheet, row, max(2, width))
        if height:
            sheet.row_dimensions[row].height = height
        return row + 1

    def _style_header(self, cell):
        from openpyxl.styles import Alignment, Font, PatternFill
        cell.fill = PatternFill("solid", fgColor=HEADER_BG)
        cell.font = Font(bold=True, color=HEADER_INK, size=10.5)
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        cell.border = _thin_border()

    def _header_row(self, sheet, row: int, columns: List[str]) -> int:
        for index, name in enumerate(columns, start=1):
            self._style_header(sheet.cell(row=row, column=index, value=name))
        sheet.row_dimensions[row].height = 26
        return row + 1

    def _title_block(self, sheet, title: str, subtitle: str = "", width: int = 8,
                     row: int = 1) -> int:
        """Title band, subtitle strip, provenance line, DEMO warning - in that
        order, with a blank row before the table starts."""
        row = self._band(sheet, row, title, width, TITLE_BG, TITLE_INK, size=15,
                         height=28)
        if subtitle:
            row = self._band(sheet, row, subtitle, width, SUBTITLE_BG, "E6F0F9",
                             size=10.5, bold=False, height=18)
        row = self._band(sheet, row, f"ZEUS  ·  {BRAND_TAGLINE}", width,
                         "FFFFFF", MUTED_INK, size=9, bold=True, height=15)
        row = self._band(sheet, row,
                         f"Generated {self.analysis.generated_at}"
                         f"  ·  local and offline", width, "FFFFFF", MUTED_INK,
                         size=9, bold=False, height=15)
        if self.analysis.contains_demo_data:
            row = self._band(sheet, row, self.analysis.demo_banner(), width,
                             WARN_BG, WARN_INK, size=9.5, height=16)
        return row + 1

    @staticmethod
    def _autosize(sheet, columns: List[str], rows: List[List[Any]],
                  start_column: int = 1):
        from openpyxl.utils import get_column_letter
        for index, name in enumerate(columns):
            header_word = max([len(w) for w in str(name).split()] or [4])
            longest = max(header_word + 4, min(len(str(name)) + 2, 22))
            for row in rows[:300]:
                if index < len(row) and row[index] is not None:
                    longest = max(longest, min(len(str(row[index])), 72))
            letter = get_column_letter(start_column + index)
            sheet.column_dimensions[letter].width = min(max(longest + 2, 11), 60)

    @staticmethod
    def _finish_sheet(sheet, tab_colour: str, landscape: bool = True,
                      freeze: Optional[str] = None, one_page: bool = False):
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.tabColor = tab_colour
        sheet.page_setup.orientation = "landscape" if landscape else "portrait"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1 if one_page else 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_options.horizontalCentered = True
        if freeze:
            sheet.freeze_panes = freeze

    # ------------------------------------------------------ summary ------

    def summary_sheet(self):
        """The management view: the recommendation, a KPI band that reads like
        the cards on the deck, the actions, and the decision requested."""
        from openpyxl.styles import Alignment, Font, PatternFill
        sheet = self.book.create_sheet(_safe_sheet_name("Summary", self._used_names))
        width = 6
        row = self._title_block(sheet, self.analysis.title, self.analysis.subtitle,
                                width=width, row=1)

        def prose(label: str, text: str, span: int, height: float, bold=False,
                  size=11, fill: Optional[str] = None):
            nonlocal row
            row = self._band(sheet, row, label.upper(), width, SECTION_BG,
                             SECTION_INK, size=9.5, height=16)
            cell = sheet.cell(row=row, column=1, value=text)
            cell.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
            cell.font = Font(size=size, bold=bold, color=INK)
            if fill:
                for column in range(1, width + 1):
                    sheet.cell(row=row, column=column).fill = PatternFill(
                        "solid", fgColor=fill)
            sheet.merge_cells(start_row=row, start_column=1, end_row=row + span - 1,
                              end_column=width)
            for offset in range(span):
                sheet.row_dimensions[row + offset].height = height
            row += span + 1

        if self.analysis.headline:
            prose("Recommendation", self.analysis.headline, span=2, height=20,
                  bold=True, size=12, fill=NOTE_BG)
        if self.analysis.executive_summary:
            prose("Executive summary", self.analysis.executive_summary, span=4,
                  height=17)

        # ---- KPI table. "Change" is only offered when a KPI actually has one:
        # an always-empty column is clutter, not completeness.
        row = self._band(sheet, row, "KEY NUMBERS", width, SECTION_BG, SECTION_INK,
                         size=9.5, height=16)
        show_change = any(k.delta for k in self.analysis.kpis)
        kpi_headers = (["Measure", "Value", "Change", "What it means"] if show_change
                       else ["Measure", "Value", "What it means"])
        row = self._header_row(sheet, row, kpi_headers)
        tone_fill = {GOOD: (GOOD_BG, GOOD_INK), BAD: (BAD_BG, BAD_INK),
                     WARN: (WARN_BG, WARN_INK)}
        border = _thin_border()
        for index, kpi in enumerate(self.analysis.kpis):
            banded = index % 2 == 1
            measure = sheet.cell(row=row, column=1, value=kpi.label)
            measure.font = Font(bold=True, size=10.5)
            measure.alignment = Alignment(vertical="center", indent=1)
            value_cell = sheet.cell(row=row, column=2, value=kpi.value)
            value_cell.font = Font(bold=True, size=12)
            value_cell.alignment = Alignment(horizontal="center", vertical="center")
            background, ink = tone_fill.get(kpi.tone, (None, None))
            if background:
                value_cell.fill = PatternFill("solid", fgColor=background)
                value_cell.font = Font(bold=True, size=12, color=ink)
            note_column = 3
            if show_change:
                change = sheet.cell(row=row, column=3, value=kpi.delta)
                change.alignment = Alignment(horizontal="center", vertical="center")
                note_column = 4
            note = sheet.cell(row=row, column=note_column, value=kpi.note)
            note.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
            note.font = Font(size=10, color=MUTED_INK)
            for column in range(1, len(kpi_headers) + 1):
                cell = sheet.cell(row=row, column=column)
                cell.border = border
                if banded and cell.fill.fgColor.rgb in (None, "00000000"):
                    cell.fill = PatternFill("solid", fgColor=BAND)
            sheet.row_dimensions[row].height = 22
            row += 1
        row += 1

        # ---- recommended actions
        if self.analysis.recommendations:
            row = self._band(sheet, row, "RECOMMENDED ACTIONS", width, SECTION_BG,
                             SECTION_INK, size=9.5, height=16)
            headers = ["Action", "Priority", "Why it matters", "Impact", "Effort",
                       "Owner"]
            row = self._header_row(sheet, row, headers)
            for position, item in enumerate(
                    sorted(self.analysis.recommendations, key=lambda r: r.priority),
                    start=1):
                banded = position % 2 == 0
                action = sheet.cell(row=row, column=1, value=f"{position}.  {item.title}")
                action.font = Font(bold=True, size=10.5)
                action.alignment = Alignment(wrap_text=True, vertical="top", indent=1)
                priority = sheet.cell(row=row, column=2, value=item.priority)
                priority.alignment = Alignment(horizontal="center", vertical="top")
                for column, value, colour in ((3, item.rationale, MUTED_INK),
                                              (4, item.impact, GOOD_INK),
                                              (5, item.effort, INK),
                                              (6, item.owner, INK)):
                    cell = sheet.cell(row=row, column=column, value=value)
                    cell.alignment = Alignment(wrap_text=True, vertical="top",
                                               indent=1)
                    cell.font = Font(size=10, color=colour)
                for column in range(1, len(headers) + 1):
                    cell = sheet.cell(row=row, column=column)
                    cell.border = border
                    if banded:
                        cell.fill = PatternFill("solid", fgColor=BAND)
                sheet.row_dimensions[row].height = 32
                row += 1
            row += 1

        if self.analysis.decision_request:
            prose("Decision requested", self.analysis.decision_request, span=2,
                  height=19, bold=True, fill=NOTE_BG)

        # ---- how to read this workbook: the reference's assumptions block,
        # doing an honest job here.
        row = self._band(sheet, row, "HOW TO READ THIS WORKBOOK", width, NOTE_BG,
                         SECTION_INK, size=10, height=17)
        notes = [
            "Every data sheet carries the full records; totals are live "
            "=SUM() formulas, not pasted numbers.",
            "The Analysis sheet counts findings with =COUNTIF and reaches "
            "across sheets with =COUNTA, so edits flow through.",
            "Cells shaded red are at or beyond a threshold; amber is a warning; "
            "green is within tolerance.",
            "The Sources sheet lists where each figure came from.",
        ]
        if self.analysis.contains_demo_data:
            notes.append("Rows marked DEMO are generated sample records, not "
                         "measurements from a real plant.")
        for note in notes:
            cell = sheet.cell(row=row, column=1, value=f"•  {note}")
            cell.font = Font(size=9.5, color=MUTED_INK)
            cell.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
            self._merge(sheet, row, width)
            sheet.row_dimensions[row].height = 15
            row += 1

        for letter, column_width in (("A", 44), ("B", 15), ("C", 46), ("D", 28),
                                     ("E", 22), ("F", 16)):
            sheet.column_dimensions[letter].width = column_width
        # The management view is a one-page read; it should print as one.
        self._finish_sheet(sheet, TAB_SUMMARY, one_page=True)
        return sheet

    # --------------------------------------------------------- data ------

    def data_sheet(self, block: TableBlock):
        """One table block: filters, frozen header, banded rows, typed cells,
        a real total row, and conditional formatting that means something."""
        from openpyxl.formatting.rule import CellIsRule, FormulaRule
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        sheet = self.book.create_sheet(_safe_sheet_name(block.title, self._used_names))
        header_row = self._title_block(sheet, block.title, block.note,
                                       width=len(block.columns))
        header_row = self._header_row(sheet, header_row, block.columns) - 1

        first_data_row = header_row + 1
        border = _thin_border()
        for offset, row_values in enumerate(block.rows):
            excel_row = first_data_row + offset
            banded = offset % 2 == 1
            for column_index, value in enumerate(row_values, start=1):
                cell = sheet.cell(row=excel_row, column=column_index)
                cell.value = value
                column_name = block.columns[column_index - 1]
                number_format = block.formats.get(column_name)
                if number_format:
                    cell.number_format = number_format
                elif _looks_numeric(value):
                    cell.number_format = "#,##0.##"
                cell.font = Font(size=10, bold=column_index == 1,
                                 color=INK if column_index == 1 else "37474F")
                cell.alignment = Alignment(
                    vertical="center", indent=1,
                    wrap_text=isinstance(value, str) and len(value) > 45,
                    horizontal="right" if _looks_numeric(value) else "left")
                cell.border = border
                if banded:
                    cell.fill = PatternFill("solid", fgColor=BAND)
            sheet.row_dimensions[excel_row].height = 18
        last_data_row = first_data_row + len(block.rows) - 1

        # ---- totals as real formulas, not pasted numbers
        if block.total_columns and block.rows:
            total_row = last_data_row + 1
            label = sheet.cell(row=total_row, column=1, value="TOTAL")
            label.font = Font(bold=True, size=10.5, color=TOTAL_INK)
            label.alignment = Alignment(vertical="center", indent=1)
            for name in block.total_columns:
                if name not in block.columns:
                    continue
                column_index = block.columns.index(name) + 1
                letter = get_column_letter(column_index)
                cell = sheet.cell(row=total_row, column=column_index)
                cell.value = f"=SUM({letter}{first_data_row}:{letter}{last_data_row})"
                cell.font = Font(bold=True, size=10.5, color=TOTAL_INK)
                cell.alignment = Alignment(horizontal="right", vertical="center",
                                           indent=1)
                cell.number_format = block.formats.get(name, "#,##0.##")
            for column_index in range(1, len(block.columns) + 1):
                cell = sheet.cell(row=total_row, column=column_index)
                cell.fill = PatternFill("solid", fgColor=TOTAL_BG)
                cell.border = border
            sheet.row_dimensions[total_row].height = 20

        if block.rows:
            last_letter = get_column_letter(len(block.columns))
            sheet.auto_filter.ref = f"A{header_row}:{last_letter}{last_data_row}"
            sheet.freeze_panes = f"A{first_data_row}"

        # ---- conditional formatting that means something
        for column_name, rule in (block.highlight or {}).items():
            if column_name not in block.columns or not block.rows:
                continue
            column_index = block.columns.index(column_name) + 1
            letter = get_column_letter(column_index)
            span = f"{letter}{first_data_row}:{letter}{last_data_row}"
            if rule == "severity":
                for value, background, ink in (("critical", BAD_BG, BAD_INK),
                                               ("high", BAD_BG, BAD_INK),
                                               ("medium", WARN_BG, WARN_INK)):
                    sheet.conditional_formatting.add(span, CellIsRule(
                        operator="equal", formula=[f'"{value}"'],
                        fill=PatternFill("solid", fgColor=background),
                        font=Font(bold=True, color=ink)))
            elif rule == "below_minimum" and "Minimum" in block.columns:
                minimum_letter = get_column_letter(block.columns.index("Minimum") + 1)
                sheet.conditional_formatting.add(span, FormulaRule(
                    formula=[f"AND({letter}{first_data_row}<>\"\","
                             f"{letter}{first_data_row}<={minimum_letter}{first_data_row})"],
                    fill=PatternFill("solid", fgColor=BAD_BG),
                    font=Font(bold=True, color=BAD_INK)))
            elif rule == "status":
                for value, background, ink in (("operational", GOOD_BG, GOOD_INK),
                                               ("maintenance", WARN_BG, WARN_INK),
                                               ("shutdown", BAD_BG, BAD_INK)):
                    sheet.conditional_formatting.add(span, CellIsRule(
                        operator="equal", formula=[f'"{value}"'],
                        fill=PatternFill("solid", fgColor=background),
                        font=Font(bold=True, color=ink)))

        self._autosize(sheet, block.columns, block.rows)
        self._finish_sheet(sheet, TAB_DATA)
        self._data_sheets[block.title] = (sheet.title, block, header_row, last_data_row)
        return sheet

    # ----------------------------------------------------- analysis ------

    def analysis_sheet(self):
        """Findings, plus counts computed by formula ACROSS the data sheets so
        the summary tracks edits instead of freezing at generation time."""
        from openpyxl.chart import BarChart, Reference
        from openpyxl.chart.label import DataLabelList
        from openpyxl.styles import Alignment, Font, PatternFill

        sheet = self.book.create_sheet(_safe_sheet_name("Analysis", self._used_names))
        width = 5
        row = self._title_block(sheet, "Analysis",
                                "Findings, a live severity profile and the record "
                                "counts behind them", width=width)
        border = _thin_border()

        if self.analysis.findings:
            row = self._band(sheet, row, "FINDINGS", width, SECTION_BG, SECTION_INK,
                             size=9.5, height=16)
            row = self._header_row(sheet, row, ["Severity", "Finding", "Detail",
                                                "Magnitude", "Evidence"])
            first_finding_row = row
            for index, finding in enumerate(self.analysis.findings):
                banded = index % 2 == 1
                severity = sheet.cell(row=row, column=1,
                                      value=finding.severity.title())
                severity.alignment = Alignment(horizontal="center", vertical="center")
                if finding.severity.lower() in ("critical", "high"):
                    severity.fill = PatternFill("solid", fgColor=BAD_BG)
                    severity.font = Font(bold=True, color=BAD_INK, size=10)
                elif finding.severity.lower() == "medium":
                    severity.fill = PatternFill("solid", fgColor=WARN_BG)
                    severity.font = Font(bold=True, color=WARN_INK, size=10)
                else:
                    severity.font = Font(bold=True, color=MUTED_INK, size=10)
                title_cell = sheet.cell(row=row, column=2, value=finding.title)
                title_cell.font = Font(bold=True, size=10.5)
                title_cell.alignment = Alignment(wrap_text=True, vertical="center",
                                                 indent=1)
                detail = sheet.cell(row=row, column=3, value=finding.detail)
                detail.alignment = Alignment(wrap_text=True, vertical="center",
                                             indent=1)
                detail.font = Font(size=10, color="37474F")
                magnitude = sheet.cell(row=row, column=4, value=finding.magnitude)
                magnitude.font = Font(bold=True, size=10, color=ACCENT)
                magnitude.alignment = Alignment(horizontal="center", vertical="center")
                evidence = sheet.cell(row=row, column=5,
                                      value=finding.evidence.replace("\n", " "))
                evidence.font = Font(size=9, color=MUTED_INK)
                evidence.alignment = Alignment(wrap_text=True, vertical="center",
                                               indent=1)
                for column in range(1, width + 1):
                    cell = sheet.cell(row=row, column=column)
                    cell.border = border
                    if banded and column not in (1,):
                        cell.fill = PatternFill("solid", fgColor=BAND)
                sheet.row_dimensions[row].height = 30
                row += 1
            last_finding_row = row - 1
            row += 1

            # ---- live COUNTIF over the findings just written
            row = self._band(sheet, row, "SEVERITY PROFILE  —  live formulas, not "
                             "pasted counts", width, SECTION_BG, SECTION_INK,
                             size=9.5, height=16)
            row = self._header_row(sheet, row, ["Severity", "Count", "Share"])
            profile_first = row
            for index, severity in enumerate(("Critical", "High", "Medium", "Low")):
                label = sheet.cell(row=row, column=1, value=severity)
                label.font = Font(bold=True, size=10)
                label.alignment = Alignment(vertical="center", indent=1)
                count = sheet.cell(row=row, column=2, value=(
                    f'=COUNTIF(A{first_finding_row}:A{last_finding_row},"{severity}")'))
                count.alignment = Alignment(horizontal="center", vertical="center")
                count.font = Font(bold=True, size=10)
                share = sheet.cell(row=row, column=3, value=(
                    f"=IFERROR(B{row}/COUNTA(A{first_finding_row}:A{last_finding_row}),0)"))
                share.number_format = "0%"
                share.alignment = Alignment(horizontal="center", vertical="center")
                for column in range(1, 4):
                    cell = sheet.cell(row=row, column=column)
                    cell.border = border
                    if index % 2 == 1:
                        cell.fill = PatternFill("solid", fgColor=BAND)
                row += 1
            profile_last = row - 1
            total_label = sheet.cell(row=row, column=1, value="Total")
            total_label.font = Font(bold=True, color=TOTAL_INK)
            total_label.alignment = Alignment(vertical="center", indent=1)
            total_cell = sheet.cell(row=row, column=2,
                                    value=f"=SUM(B{profile_first}:B{profile_last})")
            total_cell.font = Font(bold=True, color=TOTAL_INK)
            total_cell.alignment = Alignment(horizontal="center", vertical="center")
            for column in range(1, 4):
                cell = sheet.cell(row=row, column=column)
                cell.fill = PatternFill("solid", fgColor=TOTAL_BG)
                cell.border = border
            row += 2

            chart = BarChart()
            chart.type = "col"
            chart.title = "Findings by severity"
            chart.height, chart.width = 7.4, 13
            chart.y_axis.title = "Findings"
            chart.gapWidth = 60
            data = Reference(sheet, min_col=2, min_row=profile_first - 1,
                             max_row=profile_last)
            categories = Reference(sheet, min_col=1, min_row=profile_first,
                                   max_row=profile_last)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(categories)
            chart.legend = None
            chart.dLbls = DataLabelList()
            chart.dLbls.showVal = True
            sheet.add_chart(chart, f"E{profile_first - 2}")

        # ---- cross-sheet counts against the data sheets
        if self._data_sheets:
            row = self._band(sheet, row, "RECORD COUNTS  —  live, across sheets",
                             width, SECTION_BG, SECTION_INK, size=9.5, height=16)
            row = self._header_row(sheet, row, ["Dataset", "Rows", "Sheet"])
            for index, (title, (sheet_name, block, header_row, last_row)) in enumerate(
                    self._data_sheets.items()):
                name_cell = sheet.cell(row=row, column=1, value=title)
                name_cell.font = Font(bold=True, size=10)
                name_cell.alignment = Alignment(vertical="center", indent=1)
                first = header_row + 1
                count = sheet.cell(row=row, column=2)
                count.value = (f"=COUNTA('{sheet_name}'!A{first}:A{last_row})"
                               if last_row >= first else 0)
                count.alignment = Alignment(horizontal="center", vertical="center")
                count.font = Font(bold=True, size=10)
                reference = sheet.cell(row=row, column=3, value=sheet_name)
                reference.font = Font(size=9.5, color=MUTED_INK)
                reference.alignment = Alignment(vertical="center", indent=1)
                for column in range(1, 4):
                    cell = sheet.cell(row=row, column=column)
                    cell.border = border
                    if index % 2 == 1:
                        cell.fill = PatternFill("solid", fgColor=BAND)
                row += 1

        for letter, column_width in (("A", 16), ("B", 40), ("C", 56), ("D", 18),
                                     ("E", 30)):
            sheet.column_dimensions[letter].width = column_width
        self._finish_sheet(sheet, TAB_ANALYSIS)
        return sheet

    # ------------------------------------------------------ sources ------

    def sources_sheet(self):
        from openpyxl.styles import Alignment, Font, PatternFill
        sheet = self.book.create_sheet(_safe_sheet_name("Sources", self._used_names))
        width = 2
        row = self._title_block(sheet, "Sources and assumptions",
                                "Where every figure in this workbook came from",
                                width=width)
        row = self._header_row(sheet, row, ["Item", "Detail"])
        border = _thin_border()

        entries: List[Tuple[str, str]] = [
            ("Generated", self.analysis.generated_at),
            ("Generated by", f"ZEUS — {BRAND_TAGLINE} (local, offline)"),
            ("Narrative", "written by the local model, constrained to these figures"
             if self.analysis.narrative_included
             else "not used — figures and text are computed from records"),
        ]
        entries += [("Data source", item) for item in self.analysis.source_data]
        entries += [("Citation", item) for item in self.analysis.citations]
        if self.analysis.contains_demo_data:
            entries.append(("WARNING", self.analysis.demo_banner()))
        entries += [
            ("Network", "none — nothing left this machine"),
            ("Verification", "this file was reopened and parsed after writing; a "
                             "file that fails that check is reported as an error"),
        ]

        for index, (item, detail) in enumerate(entries):
            key = sheet.cell(row=row, column=1, value=item)
            key.font = Font(bold=True, size=10,
                            color=BAD_INK if item == "WARNING" else INK)
            key.alignment = Alignment(vertical="center", indent=1)
            cell = sheet.cell(row=row, column=2, value=detail)
            cell.alignment = Alignment(wrap_text=True, vertical="center", indent=1)
            cell.font = Font(size=10, color="37474F")
            for column in (1, 2):
                target = sheet.cell(row=row, column=column)
                target.border = border
                if item == "WARNING":
                    target.fill = PatternFill("solid", fgColor=WARN_BG)
                elif index % 2 == 1:
                    target.fill = PatternFill("solid", fgColor=BAND)
            sheet.row_dimensions[row].height = 20
            row += 1

        sheet.column_dimensions["A"].width = 22
        sheet.column_dimensions["B"].width = 96
        self._finish_sheet(sheet, TAB_SOURCES, landscape=False, one_page=True)
        return sheet

    # ------------------------------------------------------------ build --

    def build(self, max_data_sheets: int = 6) -> "WorkbookBuilder":
        self.summary_sheet()
        for block in self.analysis.tables[:max_data_sheets]:
            if block.rows:
                self.data_sheet(block)
        if self.analysis.financials and self.analysis.financials.rows:
            self.data_sheet(self.analysis.financials)
        self.analysis_sheet()
        self.sources_sheet()
        # Summary first is not automatic once sheets are appended out of order.
        order = ["Summary"] + [name for name in self.book.sheetnames if name != "Summary"]
        self.book._sheets = [self.book[name] for name in order]
        self.book.active = 0
        return self

    def save(self, path) -> str:
        self.book.save(str(path))
        return str(path)


def build_workbook(analysis: IndustrialAnalysis, path) -> Dict[str, Any]:
    from core import office
    builder = WorkbookBuilder(analysis).build()
    builder.save(path)
    return office.finish(path, "xlsx", extra={
        "sheet_names": builder.book.sheetnames,
        "analysis": analysis.summary_dict(),
        "demo_data": analysis.contains_demo_data,
        "narrative_included": analysis.narrative_included,
    })
