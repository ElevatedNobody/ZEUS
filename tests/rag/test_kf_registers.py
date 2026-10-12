"""
Registers - tables whose rows are records (leases, employees, filings,
contracts) - answered by the engine, row by row, with the data checked.

Synthetic workbooks are built per test run; dates are relative to today so
'the next 90 days' is always testable. Nothing here uses real company data.
"""

import datetime as dt
from pathlib import Path

import pytest

from knowledge_fabric import registers, store

TODAY = dt.date.today()


def _d(days: int) -> str:
    return (TODAY + dt.timedelta(days=days)).isoformat()


def _workbook(path: Path, sheets: dict) -> Path:
    from openpyxl import Workbook
    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        ws = book.create_sheet(name)
        for row in rows:
            ws.append(row)
    book.save(path)
    return path


LEASES = [["Lease_ID", "Property_Name", "Lessor", "Start_Date", "End_Date", "Tenure_Months", "Monthly_Rent_INR",
           "Security_Deposit_INR", "Status", "City", "State"],
          ["LS-T-001", "Pune Shed 1", "Alpha Realty", _d(-700), _d(30), 24, 100000, 600000, "Active", "Pune", "Maharashtra"],
          ["LS-T-002", "Pune Shed 2", "Alpha Realty", _d(-900), _d(60), 32, 200000, 1200000, "Active", "Pune", "Maharashtra"],
          ["LS-T-003", "Chennai Depot", "Beta Estates", _d(-1500), _d(-400), 36, 150000, 450000, "Active", "Chennai", "Tamil Nadu"],
          ["LS-T-004", "Noida Office", "Gamma Infra", _d(-300), _d(700), 33, 300000, 900000, "Expiring_90d", "Noida", "Uttar Pradesh"],
          ["LS-T-005", "Indore Store", "Beta Estates", _d(-200), _d(900), 36, 50000, 150000, "Terminated", "Indore", "Madhya Pradesh"]]
TAX = [["Assessment_ID", "Property_Code", "FY", "Tax_Amount", "Due_Date", "Paid_On", "Status"],
       ["PT-T-1", "PRP-PUN-001", "2025-26", 100000, _d(-60), _d(-70), "Paid"],
       ["PT-T-2", "PRP-PUN-002", "2025-26", 250000, _d(-60), "", "Unpaid"],
       ["PT-T-3", "PRP-CHE-003", "2025-26", 40000, _d(-60), "", "Paid"],
       ["PT-T-4", "PRP-NOI-004", "2025-26", 75000, _d(-60), _d(5), "Paid"]]
STAFF = [["Employee_ID", "Full_Name", "Department", "Grade", "DOJ", "CTC_Annual_INR", "PAN_Masked", "Aadhaar_Masked",
          "Bank_IFSC", "Status"],
         ["EMP-T-01", "Test Person 1", "Finance", "L2", "2019-02-01", 900000, "ABCDE****1F", "XXXX-XXXX-0001", "HDFC0000001", "Active"],
         ["EMP-T-02", "Test Person 2", "Finance", "L3", "2020-03-01", 1200000, "ABCDE****2F", "XXXX-XXXX-0002", "HDFC0000002", "Active"],
         ["EMP-T-03", "Test Person 3", "Quality", "L1", "2021-04-01", 600000, "ABCDE****3F", "XXXX-XXXX-0003", "HDFC0000003", "Exited"],
         ["EMP-T-04", "Test Person 4", "Quality", "L2", "2022-05-01", 750000, "ABCDE****4F", "XXXX-XXXX-0004", "HDFC0000004", "Active"]]


@pytest.fixture(scope="module")
def reg_db(tmp_path_factory):
    from knowledge_fabric import ingest
    folder = tmp_path_factory.mktemp("registers")
    db = folder / "kf_registers.db"
    files = [_workbook(folder / "Lease_Register_Master.xlsx", {"Lease_Register": LEASES}),
             _workbook(folder / "Tax_Compliance_Master.xlsx", {"Property_Tax": TAX}),
             _workbook(folder / "Employee_HR_Master.xlsx", {"Employee_Master": STAFF})]
    (folder / "gst_sop.md").write_text("# GST COMPLIANCE SOP\n**Document ID:** SOP-TAX-GST-01\n\n## Responsibilities\n"
                                       "Tax Manager reviews and submits GSTR-3B by the 20th.\n", encoding="utf-8")
    (folder / "prompts.md").write_text("# SLIDE DECK: Use cases\n## Slide 1 - Questions\n"
                                       "- “Who submits GSTR-3B?” - “When is GSTR-3B due?” - "
                                       "“List the GST forms”\n", encoding="utf-8")
    store.reset_cache()
    with store.using_database(db):
        for f in files + [folder / "gst_sop.md", folder / "prompts.md"]:
            assert ingest.ingest_document(f, origin="test_fixture", embed=False).status == "ingested"
    return db


@pytest.fixture
def reg(reg_db):
    with store.using_database(reg_db):
        yield reg_db


def ctx(role="engineer", departments=()):
    from tests.conftest import kf_ctx
    return kf_ctx(role, departments=departments)


def ask(question, role="admin", departments=(), today=TODAY):
    return registers.answer(question, ctx(role, departments), today=today)


# ------------------------------------------------------------------ ingest ---

def test_a_register_is_recognised_and_its_rows_become_records(reg):
    with store.connection() as conn:
        names = {r["name"]: r for r in store.rows(conn.execute("SELECT name, key_column, n_records FROM kf_registers"))}
        assert names["Lease_Register"]["key_column"] == "Lease_ID" and names["Lease_Register"]["n_records"] == 5
        nodes = conn.execute("SELECT COUNT(*) FROM kf_nodes WHERE kind = 'record'").fetchone()[0]
        assert nodes == 5 + 4 + 4


def test_register_rows_are_not_turned_into_dated_facts(reg):
    """'LS-T-001 rent 100000' is not a figure for the year the workbook
    mentions most - only a real period column (FY) dates a value."""
    with store.connection() as conn:
        lease_facts = conn.execute("SELECT COUNT(*) FROM kf_facts f JOIN kf_documents d ON d.document_id = f.document_id "
                                   "WHERE d.filename = 'Lease_Register_Master.xlsx'").fetchone()[0]
        tax_periods = {r[0] for r in conn.execute("SELECT DISTINCT f.period FROM kf_facts f JOIN kf_documents d "
                                                   "ON d.document_id = f.document_id WHERE d.filename = 'Tax_Compliance_Master.xlsx'")}
        doc = dict(conn.execute("SELECT document_type, year FROM kf_documents WHERE filename = 'Lease_Register_Master.xlsx'").fetchone())
    assert lease_facts == 0
    assert tax_periods and all("2025" in (p or "") for p in tax_periods)
    assert doc == {"document_type": "register", "year": None}


def test_a_register_of_personal_data_is_restricted_to_hr(reg):
    with store.connection() as conn:
        doc = dict(conn.execute("SELECT confidentiality, department, department_restricted FROM kf_documents "
                                "WHERE filename = 'Employee_HR_Master.xlsx'").fetchone())
    assert doc == {"confidentiality": "confidential", "department": "Human Resources", "department_restricted": 1}
    assert ask("How many employees are in Finance?", role="engineer") is None      # not HR: not even the register
    assert ask("How many employees are in Finance?", role="engineer", departments=("Human Resources",)) is not None
    assert ask("How many employees are in Finance?", role="admin") is not None


def test_identifiers_are_never_shown_unless_asked_for(reg):
    found = ask("How many employees are in Finance?")
    assert not {"PAN_Masked", "Aadhaar_Masked", "Bank_IFSC", "CTC_Annual_INR"} & set(found.shown_columns)
    asked = ask("What is the CTC of EMP-T-02?")
    assert asked.field_asked == "CTC_Annual_INR" and asked.records[0]["data"]["CTC_Annual_INR"] == "1200000"


# ------------------------------------------------------------------- query ---

def test_the_next_90_days_uses_the_end_date(reg):
    found = ask("Which leases expire in the next 90 days?")
    assert found.mode == "list" and {r["record_key"] for r in found.records} == {"LS-T-001", "LS-T-002"}
    assert any("End_Date" in c for c in found.conditions)


def test_counts_and_totals_are_calculated_from_the_rows(reg):
    found = ask("What is the total monthly rent of active leases?")
    assert {r["record_key"] for r in found.records} == {"LS-T-001", "LS-T-002", "LS-T-003"}
    total = next(c for c in found.calculations if c["name"].endswith("total Monthly_Rent_INR"))
    assert total["value"] == 450000 and len(total["inputs"]) == 3
    count = ask("How many leases are in Pune?")
    assert count.mode == "count" and count.calculations[0]["value"] == 2
    top = ask("Which lease has the highest security deposit?")
    assert top.mode == "max" and top.records[0]["record_key"] == "LS-T-002"


def test_group_by_counts_each_value(reg):
    found = ask("How many leases by city?")
    groups = {c["name"].split("= ")[-1]: c["value"] for c in found.calculations if "City =" in c["name"]}
    assert groups == {"Pune": 2, "Chennai": 1, "Noida": 1, "Indore": 1}


def test_unpaid_and_before_a_month(reg):
    unpaid = ask("Which property tax items are unpaid?")
    assert [r["record_key"] for r in unpaid.records] == ["PT-T-2"]
    month = (TODAY + dt.timedelta(days=45)).strftime("%B %Y")
    first = (TODAY + dt.timedelta(days=45)).replace(day=1)
    before = ask(f"Which leases end before {month}?")
    expected = {k for k, end in (("LS-T-001", 30), ("LS-T-002", 60), ("LS-T-003", -400))
                if TODAY + dt.timedelta(days=end) < first}
    assert {r["record_key"] for r in before.records} == expected


def test_a_record_named_by_its_code_is_looked_up(reg):
    found = ask("What is the security deposit for LS-T-004?")
    assert found.mode == "lookup" and found.field_asked == "Security_Deposit_INR"
    assert found.records[0]["data"]["Security_Deposit_INR"] == "900000"


def test_a_question_that_names_no_register_is_left_to_passage_search(reg):
    assert ask("Who submits GSTR-3B?") is None
    assert ask("What is the notice period for L3?") is None      # 'L3' alone does not pick the staff register


# ------------------------------------------------------------ data quality ---

def test_rows_that_contradict_the_calendar_are_reported(reg):
    checks = {(f["record_key"], f["check"]) for f in ask("Which leases expire in the next 90 days?").findings}
    assert ("LS-T-003", "status_contradicts_date") in checks            # Active, ended 400 days ago
    assert ("LS-T-004", "flag_contradicts_date") in checks              # 'Expiring_90d', ends in ~2 years
    tax = {(f["record_key"], f["check"]) for f in ask("Which property tax items are unpaid?").findings}
    assert ("PT-T-3", "done_without_date") in tax                        # Paid, no Paid_On
    assert ("PT-T-4", "future_completion") in tax                        # paid "in 5 days"


# ------------------------------------------------------------- the service ---

def test_the_answer_is_the_records_with_valid_citations(reg):
    from knowledge_fabric import service
    obj = service.analyze("What is the total monthly rent of active leases?", ctx=ctx("admin"),
                          use_model=False, use_memory=False, use_legacy=False, make_artifacts=False)
    assert obj.records["register"]["name"] == "Lease_Register"
    assert "450,000" in obj.answer and "| LS-T-001 |" in obj.answer
    assert obj.validation["ok"], obj.validation
    assert obj.evidence[0]["kind"] == "record" and obj.evidence[0]["heading"] == "LS-T-001"
    assert any(c["type"] == "record" for c in obj.contradictions)


def test_the_excel_export_recomputes_the_total(reg, tmp_path):
    from openpyxl import load_workbook
    from knowledge_fabric import renderers, service
    obj = service.analyze("What is the total monthly rent of active leases?", ctx=ctx("admin"),
                          use_model=False, use_memory=False, use_legacy=False, make_artifacts=False)
    with renderers.writing_to(tmp_path):
        made = renderers.create(obj, "xlsx")
    assert made["ok"] and made["verified"], made
    ws = load_workbook(made["path"])["Records"]
    rows = {r[0].value: r for r in ws.iter_rows() if r[0].value}
    assert rows["Total Monthly_Rent_INR"][1].value.startswith("=SUM(") and rows["Total Monthly_Rent_INR"][3].value == 450000
    assert rows["Records matching"][3].value == 3


def test_passage_search_skips_rows_and_question_lists(reg):
    from knowledge_fabric import retrieval
    from knowledge_fabric.retrieval import _only_questions
    result = retrieval.retrieve("Who submits GSTR-3B?", ctx("admin"), {"primary": "FACT_LOOKUP", "intents": []},
                                top_k=5, use_legacy=False)
    assert all(e.kind != "record" for e in result.evidence)
    assert "Tax Manager" in result.evidence[0].text
    assert _only_questions("- “Who submits GSTR-3B?” - “When is it due?” - “List the forms”")
    assert not _only_questions("Resignation notice: 30 days (L1-L2), 60 days (L3+). Exit interview follows.")


# --------------------------------------------------------- classification ---

@pytest.mark.parametrize("text, name, expected", [
    ("# GST COMPLIANCE SOP\n**Document ID:** SOP-TAX-GST-01", "gst_sop.md", "sop"),
    ("SLIDE DECK: HR map\nSlide 1 - Title\nSlide 2 - Policies incl. EHS", "HR_Policy_Map.md", "presentation"),
    ("INDUSTRIAL LEASE AGREEMENT\nLessor: X\nLessee: Y\nGoverning law", "lease_agreement.md", "contract"),
    ("Title chain: sale deed registered, encumbrance certificate", "property_title_summary.md", "legal_record"),
    ("EHS SAFETY BRIEF\nGolden rules. Report near-miss within shift.", "ehs_safety_brief.md", "safety_record"),
])
def test_document_types_from_the_text(text, name, expected):
    from knowledge_fabric import authority
    assert authority.classify_type(text, name) == expected


@pytest.mark.parametrize("title, expected", [
    ("Employee HR Master", "Human Resources"), ("Tax Compliance Master", "Finance"),
    ("Vendor Contracts Register", "Procurement"), ("EHS Safety Brief", "HSE"),
    ("Lease Register Master", "Legal"), ("Quarterly throughput summary", None),
])
def test_departments_are_inferred_only_on_a_clear_lead(title, expected):
    from knowledge_fabric.ingest import infer_department
    assert infer_department(title, "", "") == expected


def test_specific_identifiers_and_file_metadata():
    from knowledge_fabric.retrieval import _FRONT_MATTER, _identifiers
    ids = _identifiers("What is the notice period for L3 and U5 in FY2019, 2020 and Q3?")
    assert [p.pattern for p in ids] and all(p.search("60 days (L3+)") or p.search("Unit U5") for p in ids)
    assert not any(p.search("period 2020, FY2019 and Q3") for p in ids)       # years and quarters are not identifiers
    assert _FRONT_MATTER.match("--- title: Emergency Response domain: emergency_response default_claim_type: fact ---")
    assert not _FRONT_MATTER.match("--- a horizontal rule, then text: the plant restarted at 06:00")


def test_the_word_report_carries_the_records(reg, tmp_path):
    from docx import Document
    from knowledge_fabric import renderers, service
    obj = service.analyze("What is the total monthly rent of active leases?", ctx=ctx("admin"),
                          use_model=False, use_memory=False, use_legacy=False, make_artifacts=False)
    with renderers.writing_to(tmp_path):
        made = renderers.create(obj, "docx")
    assert made["ok"] and made["verified"], made
    doc = Document(made["path"])
    assert any(p.text.startswith("Records - Lease_Register") for p in doc.paragraphs)
    cells = {c.text for t in doc.tables for row in t.rows for c in row.cells}
    assert {"LS-T-001", "LS-T-002", "LS-T-003"} <= cells
