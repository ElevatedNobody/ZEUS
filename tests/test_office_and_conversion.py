import sys
import shutil
import json
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import core.builtin_tools as bt
from core.decision_engine import DecisionEngine
from core.security_manager import security_manager

security_manager.set_confirm_callback(lambda msg: True)
engine = DecisionEngine()

DOCS_DIR = Path.home() / "Documents" / "ZEUS_Documents"
CONVERT_DIR = Path("/tmp/medha_test_convert")


def _cleanup():
    shutil.rmtree(DOCS_DIR, ignore_errors=True)
    shutil.rmtree(CONVERT_DIR, ignore_errors=True)
    CONVERT_DIR.mkdir(parents=True, exist_ok=True)


# --- presentation generation ---

def test_create_presentation_makes_outline_call_plus_one_per_slide():
    """The actual reliability fix: NOT one giant call for the whole deck -
    a small outline call, then one small call per slide."""
    _cleanup()
    outline = json.dumps(["Intro", "Point One"])
    slide1 = json.dumps({"bullets": ["detail a", "detail b"], "notes": "note one"})
    with patch("core.ollama_client.OllamaClient.is_available", return_value=True), \
         patch("core.ollama_client.OllamaClient.chat", side_effect=[outline, slide1]) as mock_chat:
        r = bt.create_presentation("test topic", slide_count=2)
        assert mock_chat.call_count == 2  # 1 outline + 1 content slide (title slide needs no call)
        assert "saved" in r
        from pptx import Presentation
        prs = Presentation(r["saved"])
        assert len(prs.slides) == 2
    _cleanup()


def test_one_bad_slide_does_not_kill_the_whole_presentation():
    """Regression test for the actual bug: a single malformed slide must
    not fail the entire deck the way the old giant single-call design did."""
    _cleanup()
    outline = json.dumps(["Solar System Overview", "The Sun", "Inner Planets", "Outer Planets"])
    sun = json.dumps({"bullets": ["Center of the solar system"], "notes": "note"})
    inner = json.dumps({"bullets": ["Mercury, Venus, Earth, Mars"], "notes": "note"})
    seq = [outline, sun, inner, "garbage", "still garbage"]  # Outer Planets fails twice
    with patch("core.ollama_client.OllamaClient.is_available", return_value=True), \
         patch("core.ollama_client.OllamaClient.chat", side_effect=seq):
        r = bt.create_presentation("the solar system", slide_count=4)
        assert "saved" in r
        assert r["slides"] == 3
        assert r["slides_failed"] == ["Outer Planets"]
        from pptx import Presentation
        prs = Presentation(r["saved"])
        # Assert on what is actually ON the slides rather than on the title
        # placeholder: the deck is now rendered with ZEUS's own layout, so
        # there is no python-pptx title placeholder to read. Checking every
        # shape is the stronger form of the same claim.
        rendered = " ".join(shape.text_frame.text for slide in prs.slides
                            for shape in slide.shapes if shape.has_text_frame)
        assert "Outer Planets" not in rendered
        assert "The Sun" in rendered
        assert len(prs.slides) == 3
    _cleanup()


def test_create_presentation_fails_clearly_on_bad_outline():
    with patch("core.ollama_client.OllamaClient.is_available", return_value=True), \
         patch("core.ollama_client.OllamaClient.chat", return_value="not valid json at all"):
        r = bt.create_presentation("test topic")
        assert "error" in r


def test_create_presentation_reports_ollama_offline():
    with patch("core.ollama_client.OllamaClient.is_available", return_value=False):
        r = bt.create_presentation("test topic")
        assert "error" in r
        assert "ollama" in r["error"].lower()


def test_generic_deck_uses_the_same_renderer_as_the_industrial_one():
    """The bug this guards: there were TWO presentation generators. A deck
    about the plant went through the designed renderer; a deck about any
    other subject went through python-pptx's default 4:3 template with
    title-and-bullets placeholders. One renderer, one visual language."""
    _cleanup()
    outline = json.dumps(["Water Reuse", "Sources of Reusable Water"])
    slide = json.dumps({"key_point": "Most reusable water is already on site.",
                        "bullets": ["Roof runoff is the cleanest stream",
                                    "Greywater needs treatment before reuse"],
                        "notes": "Start with the cheapest source."})
    with patch("core.ollama_client.OllamaClient.is_available", return_value=True), \
         patch("core.ollama_client.OllamaClient.chat", side_effect=[outline, slide]):
        r = bt.create_presentation("water reuse", slide_count=2)
        assert "saved" in r
        from pptx import Presentation
        from pptx.util import Emu
        prs = Presentation(r["saved"])
        width = round(Emu(prs.slide_width).inches, 2)
        assert width > 12, f"deck is {width}in wide - still the 4:3 default template"

        # The designed renderer draws its own shapes; the default template
        # would give exactly two placeholders per slide and nothing else.
        content_slide = prs.slides[1]
        assert len(content_slide.shapes) > 3, \
            "the content slide has only placeholders - the old template is still in use"
    _cleanup()


def test_placeholder_text_never_reaches_a_slide():
    """A real deck shipped with '<point 1> <point 2> <point 3>' on a slide,
    because the prompt's own example was echoed back by the model and written
    out unchecked."""
    assert bt._is_placeholder("<point 1>")
    assert bt._is_placeholder("[your text here]")
    assert bt._is_placeholder("point 2")
    assert bt._is_placeholder("TBD")
    assert bt._is_placeholder("")
    assert not bt._is_placeholder("Roof runoff is the cleanest stream")
    assert not bt._is_placeholder("5%"), "a short real figure is not a placeholder"

    assert bt._clean_points(["<point 1>", "Real content here", "TODO"]) \
        == ["Real content here"]


def test_a_slide_of_only_placeholders_is_retried_then_dropped():
    """If the model can only produce scaffolding, that slide is reported as
    failed - never written out as if it were content."""
    _cleanup()
    outline = json.dumps(["Deck Title", "Good Slide", "Bad Slide"])
    good = json.dumps({"key_point": "This one is real.",
                       "bullets": ["A specific, concrete point"], "notes": "n"})
    junk = json.dumps({"bullets": ["<point 1>", "<point 2>"], "notes": "<note>"})
    with patch("core.ollama_client.OllamaClient.is_available", return_value=True), \
         patch("core.ollama_client.OllamaClient.chat",
               side_effect=[outline, good, junk, junk]):
        r = bt.create_presentation("test topic", slide_count=3)
        assert r.get("slides_failed") == ["Bad Slide"]
        from pptx import Presentation
        text = " ".join(shape.text_frame.text for slide in Presentation(r["saved"]).slides
                        for shape in slide.shapes if shape.has_text_frame)
        assert "<point" not in text, "placeholder scaffolding was written to the deck"
    _cleanup()


def test_presentation_command_routes_correctly():
    direct = engine._match_direct_command("make a powerpoint on the solar system")
    assert direct == ("create_presentation", {"topic": "the solar system"})
    direct2 = engine._match_direct_command("create a presentation about climate change")
    assert direct2 == ("create_presentation", {"topic": "climate change"})


# --- spreadsheet generation ---

def test_create_spreadsheet_makes_headers_call_plus_row_batches():
    """The actual reliability fix: NOT one giant call for all rows - a
    small headers call, then small batched row calls."""
    _cleanup()
    from itertools import chain, repeat
    headers = json.dumps(["Name", "Score"])
    batch = json.dumps([{"Name": "Alice", "Score": 90}, {"Name": "Bob", "Score": 85}])
    with patch("core.ollama_client.OllamaClient.is_available", return_value=True), \
         patch("core.ollama_client.OllamaClient.chat", side_effect=chain([headers], repeat(batch))):
        r = bt.create_spreadsheet("test scores")
        assert "saved" in r
        from openpyxl import load_workbook
        wb = load_workbook(r["saved"])
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        assert rows[0] == ("Name", "Score")
        assert rows[1] == ("Alice", 90)
    _cleanup()


def test_one_bad_row_batch_does_not_kill_the_whole_spreadsheet():
    """Regression test for the actual bug found while fixing this: a
    failing row batch must not fail the entire spreadsheet."""
    _cleanup()
    headers = json.dumps(["Planet", "Distance (AU)", "Moons"])
    batch1 = json.dumps([
        {"Planet": "Mercury", "Distance (AU)": 0.39, "Moons": 0},
        {"Planet": "Venus", "Distance (AU)": 0.72, "Moons": 0},
        {"Planet": "Earth", "Distance (AU)": 1.0, "Moons": 1},
        {"Planet": "Mars", "Distance (AU)": 1.52, "Moons": 2},
        {"Planet": "Jupiter", "Distance (AU)": 5.2, "Moons": 95},
    ])
    seq = [headers, batch1, "garbage", "still garbage"]
    with patch("core.ollama_client.OllamaClient.is_available", return_value=True), \
         patch("core.ollama_client.OllamaClient.chat", side_effect=seq):
        r = bt.create_spreadsheet("planets")
        assert "saved" in r
        assert r["rows"] == 5  # got the first batch even though the second failed
        from openpyxl import load_workbook
        wb = load_workbook(r["saved"])
        ws = wb.active
        assert ws.max_row == 6  # header + 5 rows
    _cleanup()


def test_create_spreadsheet_fails_clearly_when_no_headers_generated():
    with patch("core.ollama_client.OllamaClient.is_available", return_value=True), \
         patch("core.ollama_client.OllamaClient.chat", return_value="not valid json"):
        r = bt.create_spreadsheet("test topic")
        assert "error" in r


def test_spreadsheet_row_generation_is_bounded_not_infinite():
    """Regression test for a real bug found while building this: if the
    model consistently returns fewer rows than requested per batch (but
    never fails outright), the old loop condition would call forever."""
    _cleanup()
    from itertools import chain, repeat
    headers = json.dumps(["Name", "Score"])
    tiny_batch = json.dumps([{"Name": "X", "Score": 1}])  # always 1 row, never reaches target
    with patch("core.ollama_client.OllamaClient.is_available", return_value=True), \
         patch("core.ollama_client.OllamaClient.chat",
               side_effect=chain([headers], repeat(tiny_batch))) as mock_chat:
        r = bt.create_spreadsheet("test scores")
        assert "saved" in r
        assert mock_chat.call_count <= 5  # 1 headers + capped batch attempts, not unbounded
    _cleanup()


def test_spreadsheet_command_routes_correctly():
    direct = engine._match_direct_command("make an excel of top 5 countries by gdp")
    assert direct == ("create_spreadsheet", {"topic": "top 5 countries by gdp"})


# --- enhanced generation quality ---

def test_write_document_produces_real_section_headings():
    _cleanup()
    content = (
        "## Introduction\nThis introductory section has enough words to pass validation "
        "checks and sets up the topic for the reader who is following along here today.\n\n"
        "## Main Analysis\nThis body section develops the main points in detail with "
        "supporting examples that elaborate on the core ideas presented earlier in depth.\n\n"
        "## Conclusion\nThis concluding section wraps up the discussion and summarizes the "
        "key takeaways for the reader before the document reaches its natural end here."
    )
    with patch("core.ollama_client.OllamaClient.is_available", return_value=True), \
         patch("core.ollama_client.OllamaClient.chat", return_value=content):
        r = bt.write_document("test topic")
        assert "saved" in r
        if Path(r["saved"]).suffix == ".docx":
            from docx import Document
            doc = Document(r["saved"])
            heading_texts = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
            assert "Introduction" in heading_texts
            assert "Conclusion" in heading_texts
    _cleanup()


def test_create_presentation_has_theme_and_speaker_notes():
    _cleanup()
    outline = json.dumps(["Title Slide", "Point One"])
    # Real sentences, not one-letter filler: single characters are now
    # rejected as placeholder scaffolding before they can reach a slide.
    slide1 = json.dumps({"key_point": "Point one matters because it sets the frame.",
                         "bullets": ["Collection area determines yield",
                                     "Storage sizing follows demand"],
                         "notes": "Explain point one in more depth."})
    with patch("core.ollama_client.OllamaClient.is_available", return_value=True), \
         patch("core.ollama_client.OllamaClient.chat", side_effect=[outline, slide1]):
        r = bt.create_presentation("test topic", slide_count=2)
        assert "saved" in r
        from pptx import Presentation
        prs = Presentation(r["saved"])
        for slide in prs.slides:
            assert slide.background.fill.type is not None
        # the content slide (not the title slide) carries real speaker notes
        assert prs.slides[1].notes_slide.notes_text_frame.text
    _cleanup()


def test_create_spreadsheet_has_freeze_panes_and_filter():
    _cleanup()
    from itertools import chain, repeat
    headers = json.dumps(["Item", "Cost"])
    rows = json.dumps([{"Item": "A", "Cost": 100}, {"Item": "B", "Cost": 200}])
    with patch("core.ollama_client.OllamaClient.is_available", return_value=True), \
         patch("core.ollama_client.OllamaClient.chat", side_effect=chain([headers], repeat(rows))):
        r = bt.create_spreadsheet("test budget")
        assert "saved" in r
        from openpyxl import load_workbook
        wb = load_workbook(r["saved"])
        ws = wb.active
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref is not None
        assert ws["B2"].number_format == "#,##0.00"  # "Cost" recognized as a money column
    _cleanup()


def test_create_spreadsheet_adds_chart_when_data_supports_it():
    _cleanup()
    from itertools import chain, repeat
    headers = json.dumps(["Category", "Amount"])
    rows = json.dumps([{"Category": "Rent", "Amount": 1200}, {"Category": "Food", "Amount": 400}])
    with patch("core.ollama_client.OllamaClient.is_available", return_value=True), \
         patch("core.ollama_client.OllamaClient.chat", side_effect=chain([headers], repeat(rows))):
        r = bt.create_spreadsheet("monthly budget")
        from openpyxl import load_workbook
        wb = load_workbook(r["saved"])
        ws = wb.active
        # a bar chart, plus a share pie when there are few categories
        assert len(ws._charts) >= 1
    _cleanup()


# --- file conversion ---

def test_convert_image_png_to_jpg():
    _cleanup()
    from PIL import Image
    src = CONVERT_DIR / "test.png"
    dst = CONVERT_DIR / "test.jpg"
    Image.new("RGB", (20, 20), color="blue").save(src)

    r = bt.convert_file(str(src), str(dst))
    assert "converted" in r
    assert dst.exists()
    verify = Image.open(dst)
    assert verify.format == "JPEG"
    _cleanup()


def test_convert_xlsx_to_csv_and_back():
    _cleanup()
    import pandas as pd
    src = CONVERT_DIR / "data.xlsx"
    dst = CONVERT_DIR / "data.csv"
    pd.DataFrame({"A": [1, 2], "B": [3, 4]}).to_excel(src, index=False)

    r = bt.convert_file(str(src), str(dst))
    assert "converted" in r
    verify = pd.read_csv(dst)
    assert verify.to_dict("records") == [{"A": 1, "B": 3}, {"A": 2, "B": 4}]
    _cleanup()


def test_convert_docx_to_txt():
    _cleanup()
    from docx import Document
    src = CONVERT_DIR / "doc.docx"
    dst = CONVERT_DIR / "doc.txt"
    doc = Document()
    doc.add_paragraph("Hello from a real docx paragraph.")
    doc.save(str(src))

    r = bt.convert_file(str(src), str(dst))
    assert "converted" in r
    assert "Hello from a real docx paragraph." in dst.read_text()
    _cleanup()


def test_convert_unsupported_pair_reports_clearly_not_silently():
    _cleanup()
    src = CONVERT_DIR / "test.png"
    from PIL import Image
    Image.new("RGB", (5, 5)).save(src)
    r = bt.convert_file(str(src), str(CONVERT_DIR / "test.mp4"))
    assert "error" in r
    assert "no converter" in r["error"].lower()
    _cleanup()


def test_convert_nonexistent_source_reports_clearly():
    r = bt.convert_file("/tmp/does_not_exist_xyz.png", "/tmp/out.jpg")
    assert "error" in r


def test_convert_command_routes_correctly():
    direct = engine._match_direct_command("convert report.docx to report.pdf")
    assert direct == ("convert_file", {"src": "report.docx", "dst": "report.pdf"})


if __name__ == "__main__":
    import inspect
    fails = 0
    for name, fn in list(globals().items()):
        if name.startswith("test_") and inspect.isfunction(fn):
            try:
                fn()
                print(f"PASS {name}")
            except AssertionError as e:
                fails += 1
                print(f"FAIL {name}: {e}")
            except Exception as e:
                fails += 1
                print(f"ERROR {name}: {e}")
    _cleanup()
    sys.exit(1 if fails else 0)
