"""ONE VERIFIED ANALYSIS -> MULTIPLE OUTPUT FORMATS: every file carries the same
analysis, passes its validator, and states no number the analysis does not hold."""

import pytest

from knowledge_fabric import renderers
from tests.conftest import kf_ask

LOSSES = "Why did we make losses between 2019 and 2021 and what can we do?"


@pytest.fixture
def losses(kf):
    return kf_ask(LOSSES)


@pytest.mark.parametrize("fmt", ["xlsx", "docx", "pptx", "pdf", "brief", "action_plan", "md"])
def test_every_format_is_verified(losses, fmt):
    result = renderers.create(losses, fmt)
    assert result["ok"] and result["verified"], result
    assert result["consistency"]["analysis_id_present"] and not result["consistency"]["unsupported_numbers"]


def test_all_formats_come_from_the_same_analysis(losses):
    ids = {renderers.create(losses, fmt)["analysis_id"] for fmt in ("xlsx", "pptx", "docx")}
    assert ids == {losses.analysis_id}


def test_the_workbook_recomputes_and_agrees_with_the_engine(losses):
    from openpyxl import load_workbook
    path = renderers.create(losses, "xlsx")["path"]
    book = load_workbook(path)
    assert {"Summary", "Series", "Checks", "Facts", "Calculations", "Evidence", "Recommendations",
            "Implementation Plan", "KPIs", "Risks", "Provenance"} <= set(book.sheetnames)
    checks = book["Checks"]
    rows = list(checks.iter_rows(min_row=2, values_only=True))
    assert rows and all(str(r[2]).startswith("=") for r in rows)          # live formulas
    # the engine column holds the analysis values the formulas must reproduce
    series = book["Series"]
    column = [r[1] for r in series.iter_rows(min_row=2, values_only=True) if r[1] is not None]
    total = next(r for r in rows if r[1] == "total")
    assert total[3] == pytest.approx(sum(column))


def test_document_text_cannot_become_a_spreadsheet_formula():
    assert renderers._cell_text("=HYPERLINK(\"http://x\")").startswith(" =")
    assert renderers._cell_text("+cmd|' /C calc'!A0").startswith(" +")
    assert renderers._cell_text("Revenue rose") == "Revenue rose"


def test_a_number_the_analysis_does_not_hold_is_caught(losses):
    from docx import Document
    path = renderers.create(losses, "docx")["path"]
    doc = Document(path)
    doc.add_paragraph("Savings of 987654 are guaranteed.")
    doc.save(path)
    check = renderers.consistency(losses, __import__("pathlib").Path(path))
    assert not check["ok"] and "987654" in check["unsupported_numbers"]


def test_the_deck_fits_its_slides(losses):
    result = renderers.create(losses, "pptx")
    assert result["file_check"]["checks"]["nothing_overflows"] and result["file_check"]["checks"]["readable_type"]


def test_files_do_not_carry_the_operators_name(losses):
    from docx import Document
    from openpyxl import load_workbook
    from pptx import Presentation
    assert Document(renderers.create(losses, "docx")["path"]).core_properties.author == "ZEUS"
    assert load_workbook(renderers.create(losses, "xlsx")["path"]).properties.creator == "ZEUS"
    assert Presentation(renderers.create(losses, "pptx")["path"]).core_properties.author == "ZEUS"


def test_a_follow_up_request_renders_the_saved_analysis(kf):
    first = kf_ask(LOSSES + " Create an Excel file.", make_artifacts=True)
    assert first.artifacts and first.artifacts[0]["verified"]
    follow = kf_ask("Now create a management presentation from this analysis", make_artifacts=True)
    assert follow.analysis_id == first.analysis_id
    assert any(a["format"] == "pptx" and a["verified"] for a in follow.artifacts)


def test_no_evidence_means_no_file(kf):
    obj = kf_ask("What was revenue in 1995?")
    result = renderers.create(obj, "docx")
    assert not result["ok"] and "no evidence" in result["error"]
