"""
The orchestration layer: classify -> plan -> route -> execute -> validate ->
auto-fix -> deliver.

These tests check the properties that make the layer worth having, not just
that the modules import:

  * the pipeline WRAPS the existing engine - every previous response field is
    still there, and every route still resolves to the same tool
  * deterministic work does not touch a model
  * the router picks the smallest capable model and steps down under memory
    pressure
  * validators fail on real defects and, just as importantly, do NOT veto a
    tool that succeeded in a shape they did not expect
  * a repair changes something specific and is bounded
  * one sentence naming three deliverables produces three verified files from
    one analysis
"""

import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from memory import database  # noqa: E402

DB_PATH = Path(tempfile.gettempdir()) / "medha_test_orchestration.db"

from core.orchestration import activity, agents as agent_registry, autofix  # noqa: E402
from core.orchestration import task_graph, validators  # noqa: E402
from core.orchestration.classifier import classify, detect_deliverables  # noqa: E402
from core.orchestration.contracts import (  # noqa: E402
    Complexity, MedhaError, ModelChoice, OrchestrationResult, PlanStep, Stage,
    StageRecord, Status, TaskKind,
)
from core.orchestration.router import (  # noqa: E402
    DETERMINISTIC_KINDS, CapabilityRouter, estimate_ram_gb,
)
from core.tool_registry import registry  # noqa: E402
from industrial import plant_data as pd  # noqa: E402

OUT = Path(tempfile.mkdtemp(prefix="medha_orch_"))
_engine = None


def setup_module(module):
    database.DB_PATH = DB_PATH
    if DB_PATH.exists():
        DB_PATH.unlink()
    database.init_db()
    pd.init_plant_schema(force=True)
    pd.seed_demo_data(force=True)
    # Plugins register tools too (the calculator, for one). The engine does
    # this at startup, so the test environment must match it or the roster
    # check below would flag a tool that genuinely exists at runtime.
    try:
        from plugins.plugin_loader import discover_plugins
        discover_plugins()
    except Exception:
        pass


def engine():
    global _engine
    if _engine is None:
        from core.decision_engine import DecisionEngine
        _engine = DecisionEngine()
    return _engine


# ------------------------------------------------------------ contracts --

def test_engine_response_keeps_every_field_it_always_had():
    """The whole layer is additive. If this breaks, the UI breaks."""
    result = OrchestrationResult(response="hello", intent="COMMAND",
                                 agent="operations", refresh=["tasks"])
    payload = result.to_engine_response()
    for key in ("response", "intent", "refresh", "agent"):
        assert key in payload, f"'{key}' disappeared from the engine response"
    assert payload["response"] == "hello"
    assert payload["intent"] == "COMMAND"
    assert isinstance(payload["pipeline"], dict), "pipeline detail should ride along"


def test_errors_are_structured_not_strings():
    error = MedhaError(module="pptx_renderer", error_type="layout_overflow",
                        detail="card 4 hangs off the bottom",
                        recovery="retry_with_fewer_items")
    data = error.as_dict()
    assert data["status"] == "failed"
    assert data["module"] == "pptx_renderer"
    assert data["error_type"] == "layout_overflow"
    assert data["recovery"] == "retry_with_fewer_items"


# ----------------------------------------------------------- classifier --

def test_classifier_recognises_the_format_that_was_asked_for():
    assert classify("make me a powerpoint about pumps").kind is TaskKind.POWERPOINT
    assert classify("build an excel workbook of the stock").kind is TaskKind.EXCEL
    assert classify("fix this python function").kind is TaskKind.CODE
    assert classify("show the maintenance history for P-204").kind is TaskKind.INDUSTRIAL_DATA


def test_classifier_spots_several_deliverables_in_one_sentence():
    request = ("read this PDF, analyse it, update Word, create Excel and "
               "make a PPT")
    found = detect_deliverables(request.lower())
    assert TaskKind.WORD in found
    assert TaskKind.EXCEL in found
    assert TaskKind.POWERPOINT in found
    assert classify(request).kind is TaskKind.MULTI_STEP


def test_classifier_is_cheap_for_the_obvious_cases():
    """A greeting must not be classified as hard work needing a 7B model."""
    assert classify("hello").complexity is Complexity.SIMPLE
    assert classify("copy report.docx to the archive folder").complexity \
        is Complexity.TRIVIAL


def test_deep_analysis_is_recognised_as_hard():
    result = classify("analyse the root cause of the repeat bearing failures and "
                      "recommend a strategy for the board")
    assert result.complexity is Complexity.HARD


# --------------------------------------------------------------- router --

def test_deterministic_work_gets_no_model():
    router = CapabilityRouter()
    for kind in (TaskKind.FILE_OPS, TaskKind.CONVERSION, TaskKind.INDUSTRIAL_DATA,
                 TaskKind.AUTOMATION, TaskKind.OCR):
        choice = router.route_for(kind, Complexity.SIMPLE)
        assert choice.deterministic, f"{kind.value} should not need a model"
        assert choice.model == ""
        assert choice.reason, "a routing decision must say why"


def test_ram_estimate_reads_the_model_tag():
    small = estimate_ram_gb("qwen2.5:0.5b-instruct-q4_K_M")
    medium = estimate_ram_gb("qwen2.5:3b-instruct-q4_K_M")
    large = estimate_ram_gb("mistral:7b-instruct-q4_K_M")
    assert small < medium < large
    assert large < 16, "a 7B q4 model must not be estimated above the machine's RAM"


def test_router_steps_down_when_memory_is_tight(monkeypatch):
    """The 16 GB rule: if the preferred model will not fit, take a smaller one
    rather than swapping."""
    router = CapabilityRouter()
    monkeypatch.setattr("core.orchestration.router.available_ram_gb", lambda: 2.0)
    monkeypatch.setattr(router, "_resolve_slot",
                        lambda slot: {"reasoning": "mistral:7b-instruct-q4_K_M",
                                      "default": "qwen2.5:3b-instruct-q4_K_M",
                                      "fast": "qwen2.5:0.5b-instruct-q4_K_M",
                                      "json_structured": "qwen2.5:3b-instruct-q4_K_M",
                                      "classify": "qwen2.5:0.5b-instruct-q4_K_M",
                                      }.get(slot))
    choice = router.route_for(TaskKind.CHAT, Complexity.HARD)
    assert "7b" not in choice.model.lower(), \
        f"a 7B model was chosen with 2 GB free: {choice.model}"
    assert "stepped down" in choice.reason


def test_router_honours_the_slot_the_agent_asked_for(monkeypatch):
    router = CapabilityRouter()
    monkeypatch.setattr("core.orchestration.router.available_ram_gb", lambda: 32.0)
    monkeypatch.setattr(router, "_resolve_slot",
                        lambda slot: f"model-for-{slot}")
    choice = router.route_for(TaskKind.CHAT, Complexity.HARD,
                              prefer_slot="json_structured")
    assert choice.model == "model-for-json_structured"


def test_router_names_the_missing_model_instead_of_substituting(monkeypatch):
    router = CapabilityRouter()
    monkeypatch.setattr(router, "_resolve_slot", lambda slot: None)
    choice = router.route_for(TaskKind.VISION, Complexity.SIMPLE)
    assert "ollama pull" in choice.reason, \
        "a missing model must be named, not silently swapped for a text model"


def test_every_deterministic_kind_is_genuinely_tool_backed():
    """If a kind claims to need no model, tools must exist to do the work."""
    assert TaskKind.INDUSTRIAL_DATA in DETERMINISTIC_KINDS
    assert registry.get("plant_status") is not None
    assert TaskKind.CONVERSION in DETERMINISTIC_KINDS
    assert registry.get("convert_file") is not None


# ----------------------------------------------------------- validators --

def _real_files():
    from industrial import analysis, deck_builder, report_builder, workbook_builder
    built = analysis.build_operations_analysis(narrative=False)
    deck = OUT / "v.pptx"
    book = OUT / "v.xlsx"
    document = OUT / "v.docx"
    deck_builder.build_deck(built, deck)
    workbook_builder.build_workbook(built, book)
    report_builder.build_report(built, document)
    return deck, book, document


def test_validators_pass_real_generated_files():
    deck, book, document = _real_files()
    assert validators.validate("pptx", str(deck)).ok
    assert validators.validate("xlsx", str(book)).ok
    assert validators.validate("docx", str(document)).ok


def test_validator_rejects_a_missing_file():
    report = validators.validate("docx", str(OUT / "nope.docx"))
    assert not report.ok
    assert report.error.error_type == "file_missing"
    assert report.error.recovery == "regenerate"


def test_validator_rejects_an_empty_file():
    empty = OUT / "empty.xlsx"
    empty.write_bytes(b"")
    report = validators.validate("xlsx", str(empty))
    assert not report.ok
    assert report.error.error_type == "file_empty"


def test_validator_rejects_a_file_that_will_not_reopen():
    corrupt = OUT / "corrupt.pptx"
    corrupt.write_bytes(b"this is not a presentation")
    report = validators.validate("pptx", str(corrupt))
    assert not report.ok
    assert report.error.error_type == "file_corrupt"


def test_pptx_validator_catches_content_running_off_the_slide():
    """The defect a file-opens check cannot see."""
    from pptx import Presentation
    from pptx.util import Inches
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(6.5), Inches(4), Inches(3))
    box.text_frame.text = "this hangs off the bottom"
    path = OUT / "overflow.pptx"
    deck.save(str(path))

    report = validators.validate("pptx", str(path))
    assert not report.ok
    assert report.error.error_type == "layout_overflow"
    assert report.error.recovery == "retry_with_fewer_items"


def test_code_validator_actually_parses():
    good = validators.validate("code", "def add(a, b):\n    return a + b\n")
    assert good.ok
    bad = validators.validate("code", "def add(a, b)\n    return a + b\n")
    assert not bad.ok
    assert bad.error.error_type == "syntax_error"
    assert "line" in bad.error.detail


def test_code_validator_tolerates_a_markdown_fence():
    fenced = "```python\ndef f():\n    return 1\n```"
    assert validators.validate("code", fenced).ok


def test_a_validator_never_vetoes_a_tool_it_simply_does_not_understand():
    """The bug this guards: the file validator called a successful
    create_folder a failure because the tool reports its result under
    'created' rather than 'path'."""
    report = validators.validate("file", {"created": "/tmp", "note": "ok"})
    assert report.ok

    unknown_shape = validators.validate("file", {"count": 3, "message": "done"})
    assert unknown_shape.ok, "absence of a path is not evidence of failure"
    assert unknown_shape.checks["path_returned"] is False


def test_rag_validator_accepts_an_answer_in_the_tools_own_shape():
    report = validators.validate("rag", {"term": "FCC",
                                         "definition": "fluid catalytic cracking"})
    assert report.ok
    assert validators.validate("rag", {}).ok is False


def test_bundle_validator_fails_if_any_file_in_the_pack_is_bad():
    deck, book, document = _real_files()
    good = validators.validate("bundle", {"a": str(deck), "b": str(book),
                                          "c": str(document)})
    assert good.ok

    broken = OUT / "broken.docx"
    broken.write_bytes(b"nope")
    bad = validators.validate("bundle", {"a": str(deck), "b": str(broken)})
    assert not bad.ok, "a pack with one corrupt file is not a finished pack"


def test_unknown_validator_kind_is_reported_honestly():
    report = validators.validate("something_made_up", "x")
    assert report.ok
    assert report.checks["validator_exists"] is False
    assert "not checked" in report.detail


# -------------------------------------------------------------- autofix --

def test_repair_changes_something_specific():
    step = PlanStep(id="s1", tool="executive_presentation", args={"subject": "P-204"})
    error = MedhaError(module="pptx_renderer", error_type="layout_overflow",
                        recovery="retry_with_fewer_items")
    repair = autofix.plan_repair(error, step)
    assert repair is not None
    assert repair.args != step.args, "a repair that changes nothing is not a repair"
    assert repair.args["max_items"] < 6


def test_repair_feeds_the_syntax_error_back_to_the_coder():
    step = PlanStep(id="s1", tool="write_code", args={})
    error = MedhaError(module="code_agent", error_type="syntax_error",
                        detail="line 3: invalid syntax")
    repair = autofix.plan_repair(error, step)
    assert "line 3" in repair.args["fix_hint"]


def test_repairs_are_bounded():
    step = PlanStep(id="s1", tool="executive_presentation", args={})
    step.attempts = autofix.MAX_ATTEMPTS
    error = MedhaError(module="pptx_renderer", error_type="layout_overflow")
    assert autofix.plan_repair(error, step) is None, "repair loop is not bounded"


def test_an_unknown_failure_is_not_retried_blindly():
    step = PlanStep(id="s1", tool="whatever", args={})
    error = MedhaError(module="x", error_type="something_nobody_planned_for")
    assert autofix.plan_repair(error, step) is None


# ----------------------------------------------------------- task graph --

def test_one_sentence_three_deliverables_becomes_three_steps():
    request = "create an executive report and an excel workbook and a powerpoint for P-204"
    plan = task_graph.build_plan(request, classify(request))
    tools = [step.tool for step in plan]
    assert len(plan) == 3, f"expected three steps, got {tools}"
    assert "executive_report" in tools
    assert "executive_workbook" in tools
    assert "executive_presentation" in tools
    for step in plan:
        assert step.args.get("subject") == "P-204", "the tag should reach every step"


def test_a_named_source_file_is_read_once_and_shared():
    request = "read data.xlsx and make a powerpoint and a word report from it"
    plan = task_graph.build_plan(request, classify(request))
    assert plan, "no plan was produced"
    assert plan[0].tool == "analyze_spreadsheet"
    dependents = [step for step in plan[1:] if plan[0].id in step.depends_on]
    assert dependents, "the deliverables should depend on the single read step"


def test_a_source_filename_is_not_mistaken_for_a_deliverable():
    """The bug this guards: the '.xlsx' in the INPUT path was read as an order
    for a workbook, so the plan grew a phantom fourth step that then failed."""
    request = "read /tmp/data.xlsx and make a powerpoint and a word report from it"
    found = detect_deliverables(request.lower())
    assert TaskKind.EXCEL not in found, \
        "the source file's extension was treated as a requested output"
    assert TaskKind.POWERPOINT in found and TaskKind.WORD in found


def test_a_multi_deliverable_request_beats_a_single_matched_route():
    """A route can only see the first action in the sentence. When several
    deliverables are named, the multi-step plan must win or the request stops
    after the read."""
    request = "read /tmp/data.xlsx and make a powerpoint and a word report from it"
    plan = task_graph.build_plan(request, classify(request),
                                 direct=("read_file", {"path": "/tmp/data.xlsx"}))
    assert len(plan) > 1, f"collapsed to a single step: {[s.tool for s in plan]}"
    tools = [step.tool for step in plan]
    assert any("presentation" in tool for tool in tools)
    assert any("report" in tool for tool in tools)


def test_the_validator_follows_the_evidence_not_the_tool_name():
    """`analyze_spreadsheet` contains the word 'spreadsheet' but writes no
    file; validating its result as a workbook made a working tool look broken.
    """
    from core.orchestration.agents import validator_for
    assert validator_for("data_analysis", "analyze_spreadsheet") == "data"
    assert validator_for("report", "executive_presentation") == "pptx"
    assert validator_for("report", "executive_workbook") == "xlsx"
    assert validator_for("report", "executive_briefing_pack") == "bundle"


def test_path_and_tag_extraction():
    assert task_graph.extract_path('summarise "C:/data/report file.xlsx" please') \
        == "C:/data/report file.xlsx"
    assert task_graph.extract_path("analyse ./exports/q1.csv") == "./exports/q1.csv"
    assert task_graph.extract_subject("what about Pump P-204 today") == "P-204"
    assert task_graph.extract_subject("what about the plant today") == ""


def test_the_planner_never_admits_a_tool_that_does_not_exist(monkeypatch):
    monkeypatch.setattr("core.planner.decompose", lambda goal: [
        {"tool": "make_powerpoint_magically", "args": {}, "description": "invented"},
        {"tool": "plant_status", "args": {}, "description": "real"},
    ])
    plan = task_graph.model_plan("do something")
    assert [step.tool for step in plan] == ["plant_status"]


# --------------------------------------------------------------- agents --

def test_every_tool_an_agent_claims_actually_exists():
    missing = []
    for specialist in agent_registry.SPECIALISTS.values():
        for tool in specialist.tools:
            if registry.get(tool) is None:
                missing.append(f"{specialist.id}:{tool}")
    assert not missing, f"agents claim tools that do not exist: {missing}"


def test_agents_share_one_registry_rather_than_owning_copies():
    """No tool may be claimed by two specialists - that is how duplicate
    implementations start."""
    seen = {}
    clashes = []
    for specialist in agent_registry.SPECIALISTS.values():
        for tool in specialist.tools:
            if tool in seen:
                clashes.append(f"{tool}: {seen[tool]} and {specialist.id}")
            seen[tool] = specialist.id
    assert not clashes, f"tools claimed twice: {clashes}"


def test_specialists_appear_in_the_original_agent_registry():
    from core.agents import AGENT_REGISTRY
    for agent_id in ("operations", "maintenance", "safety", "inventory", "report"):
        assert agent_id in AGENT_REGISTRY, f"{agent_id} is not visible to the UI roster"


def test_deterministic_agents_declare_no_model():
    for specialist in agent_registry.SPECIALISTS.values():
        if specialist.deterministic:
            assert specialist.model_slot == "", \
                f"{specialist.id} is deterministic but asks for a model"


# --------------------------------------------------- end to end pipeline --

def test_a_plain_command_still_resolves_to_the_same_tool():
    result = engine().process("show the plant status")
    assert result["intent"] == "COMMAND"
    assert "plant_status" in result["pipeline"]["tools"]
    assert result["pipeline"]["model"] == "", \
        "a database query should not load a language model"


def test_the_pipeline_records_every_stage():
    result = engine().process("show active alerts")
    stages = [entry["stage"] for entry in result["pipeline"]["stages"]]
    for expected in ("classify", "memory", "plan", "route", "execute", "deliver"):
        assert expected in stages, f"stage '{expected}' was not recorded"


def test_one_request_three_verified_files():
    """The flagship chain from the brief, end to end."""
    result = engine().process(
        "create an executive report and an excel workbook and a powerpoint for P-204")
    pipeline = result["pipeline"]
    assert result["intent"] == "MULTI_STEP_TASK"
    assert pipeline["succeeded"] is True, result["response"]

    artifacts = pipeline["artifacts"]
    assert len(artifacts) == 3, f"expected three files, got {artifacts}"
    extensions = sorted(Path(path).suffix for path in artifacts)
    assert extensions == [".docx", ".pptx", ".xlsx"]
    for path in artifacts:
        assert Path(path).exists() and Path(path).stat().st_size > 0
    assert pipeline["validated"] is True, "files were delivered unverified"


def test_a_failing_tool_reports_its_own_message_not_a_generic_one():
    """The bug this guards: a validator's generic complaint replacing the
    tool's specific, useful error."""
    result = engine().process("read the file /definitely/not/here.txt")
    assert "definitely" in result["response"] or "not" in result["response"].lower()
    assert "no_path_returned" not in result["response"], \
        "a validator's slug leaked into the user-facing message"


def test_the_activity_board_shows_task_agent_model_tools_status():
    engine().process("show the equipment inventory")
    recent = activity.recent(1)
    assert recent, "the run was not recorded"
    entry = recent[0]
    for field in ("task", "agent", "model", "tools", "status"):
        assert field in entry, f"the UI cannot show '{field}'"
    assert entry["tools"], "no tool was recorded for a tool-backed request"


def test_the_run_is_written_to_the_audit_log():
    from core.security_manager import AUDIT_LOG
    engine().process("show the plant status")
    assert AUDIT_LOG.exists()
    tail = AUDIT_LOG.read_text(encoding="utf-8", errors="ignore")[-4000:]
    assert "orchestrator" in tail, "the pipeline did not write an audit line"


def test_orchestration_tools_are_reachable_and_honest():
    for name in ("pipeline_status", "list_specialists", "model_routing",
                 "validation_report", "explain_last_run"):
        assert registry.get(name) is not None, f"{name} is not registered"
    routing = registry.execute("model_routing")
    assert routing.success
    assert "handled_without_a_model" in routing.output
    assert routing.output["available_ram_gb"] > 0


def test_the_layer_adds_no_network_calls():
    """Offline is not a setting here; it is a property of the code."""
    import re
    banned = re.compile(r"\b(requests\.|urllib|openai|anthropic|googleapis|"
                        r"api\.openai|boto3)\b")
    folder = Path(__file__).resolve().parent.parent / "core" / "orchestration"
    for module in folder.glob("*.py"):
        text = module.read_text(encoding="utf-8")
        found = banned.search(text)
        assert not found, f"{module.name} references {found.group(0)}"
