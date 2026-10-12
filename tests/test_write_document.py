"""
Comprehensive tests for the 'write assignment/essay/report' feature.
This is the feature that was reported broken multiple times, so it gets
disproportionate test coverage - every failure mode we could think of.
"""

import sys
import shutil
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import core.builtin_tools as bt
from core.decision_engine import DecisionEngine
from core.security_manager import security_manager

security_manager.set_confirm_callback(lambda msg: True)
engine = DecisionEngine()

TEST_DOCS_DIR = Path.home() / "Documents" / "ZEUS_Documents"


def _cleanup():
    shutil.rmtree(TEST_DOCS_DIR, ignore_errors=True)


# --- routing: every phrasing variant that real users type ---

PHRASING_CASES = [
    "write assignment on photosynthesis",
    "write an assignment on photosynthesis",
    "write my assignment on photosynthesis",
    "write me an assignment about photosynthesis",
    "write me an assignment on photosynthesis",
    "can you write my assignment on photosynthesis",
    "can you write an assignment about photosynthesis",
    "could you write me an assignment on photosynthesis",
    "please write my assignment on photosynthesis",
    "write an essay on photosynthesis",
    "write an essay about photosynthesis",
    "write my essay on photosynthesis",
    "write a report on photosynthesis",
    "write a report about photosynthesis",
    "write my report on photosynthesis",
    "write an article on photosynthesis",
    "write a paper on photosynthesis",
    "write a document about photosynthesis",
    "write assignment for photosynthesis",
]


def test_all_realistic_phrasings_route_to_write_document():
    failures = []
    for phrasing in PHRASING_CASES:
        direct = engine._match_direct_command(phrasing)
        if direct is None or direct[0] != "write_document":
            failures.append(phrasing)
    assert not failures, f"These phrasings did NOT route to write_document: {failures}"


def test_topic_extraction_is_clean_not_polluted_with_connectors():
    cases = {
        "write an assignment on photosynthesis": "photosynthesis",
        "write my assignment about the french revolution": "the french revolution",
        "write a report for climate change": "climate change",
    }
    for phrasing, expected_topic in cases.items():
        direct = engine._match_direct_command(phrasing)
        assert direct[1]["topic"] == expected_topic, f"{phrasing!r} -> {direct[1]}"


def test_non_write_requests_do_not_false_positive():
    # Make sure broadening the parser didn't cause it to over-match.
    non_matches = ["write a python script to sort a list", "what should I write in my resume"]
    for text in non_matches:
        direct = engine._match_direct_command(text)
        if direct:
            assert direct[0] != "write_document", f"False positive on: {text!r}"


# --- execution: the actual bug that caused repeated failures ---

def test_stale_ollama_cache_does_not_permanently_break_write_document():
    """This is the actual root-cause regression test: is_available() used
    to cache its result forever, so one bad check at startup could make
    every future write_document call fail even with Ollama running fine."""
    from core.ollama_client import ollama_client
    import time
    ollama_client._available = False
    ollama_client._available_checked_at = time.time() - 100  # force stale

    with patch("httpx.get") as mock_get, \
         patch("core.ollama_client.OllamaClient.chat",
               return_value="This is an introductory paragraph with enough words to pass validation checks. It sets up the topic clearly and establishes context for the reader who is following along.\n\nThis is the body paragraph that develops the main points in more detail with supporting examples. It elaborates on the core ideas presented in the introduction above.\n\nThis is the concluding paragraph that wraps up the discussion and summarizes the key takeaways for the reader."):
        mock_get.return_value.status_code = 200
        _cleanup()
        r = engine.process("write my assignment on gravity", {"history": []})
        assert "error" not in r["response"].lower() or "saved" in r["response"].lower(), r["response"]
        assert "saved" in r["response"]


def test_write_document_success_produces_valid_docx():
    _cleanup()
    with patch("core.ollama_client.OllamaClient.is_available", return_value=True), \
         patch("core.ollama_client.OllamaClient.chat",
               return_value="Gravity is a fundamental force that governs the motion of objects with mass throughout the universe. It was first formally described by Isaac Newton in his law of universal gravitation.\n\nEinstein later refined our understanding of gravity through general relativity, describing it as the curvature of spacetime. This explained phenomena that Newtonian physics could not fully account for, such as the precession of Mercury.\n\nIn conclusion, gravity remains one of the four fundamental forces and continues to be an active area of physics research today."):
        r = bt.write_document("gravity")
        assert "saved" in r
        path = Path(r["saved"])
        assert path.exists()
        if r["format"] == "docx":
            from docx import Document
            doc = Document(str(path))
            assert len(doc.paragraphs) >= 3
    _cleanup()


def test_write_document_reports_ollama_offline_clearly():
    with patch("core.ollama_client.OllamaClient.is_available", return_value=False):
        r = bt.write_document("some topic")
        assert "error" in r
        assert "ollama" in r["error"].lower()
        assert "pull" in r["error"].lower()  # tells user the fix


def test_write_document_handles_model_not_found_response():
    with patch("core.ollama_client.OllamaClient.is_available", return_value=True), \
         patch("core.ollama_client.OllamaClient.chat", return_value="[Model 'qwen2.5:7b' not found. Try: ollama pull qwen2.5:7b]"):
        r = bt.write_document("some topic")
        assert "error" in r


def test_write_document_handles_empty_model_response():
    with patch("core.ollama_client.OllamaClient.is_available", return_value=True), \
         patch("core.ollama_client.OllamaClient.chat", return_value=""):
        r = bt.write_document("some topic")
        assert "error" in r


def test_write_document_rejects_suspiciously_short_generation():
    with patch("core.ollama_client.OllamaClient.is_available", return_value=True), \
         patch("core.ollama_client.OllamaClient.chat", return_value="Too short."):
        r = bt.write_document("some topic")
        assert "error" in r
        assert "little" in r["error"].lower() or "content" in r["error"].lower()


def test_write_document_strips_markdown_code_fences_if_model_adds_them():
    _cleanup()
    fenced_response = (
        "```\nThis is an introductory paragraph with enough words to pass content validation "
        "thresholds and set up the topic clearly for the reader who is following along.\n\n"
        "This is the body paragraph developing the main points in detail with supporting "
        "examples that elaborate on the core ideas presented in the introduction above.\n\n"
        "This is the concluding paragraph that wraps up the discussion and summarizes the "
        "key takeaways for the reader before the document ends.\n```"
    )
    with patch("core.ollama_client.OllamaClient.is_available", return_value=True), \
         patch("core.ollama_client.OllamaClient.chat", return_value=fenced_response):
        r = bt.write_document("test topic")
        assert "saved" in r
        path = Path(r["saved"])
        if path.suffix == ".txt":
            text = path.read_text(encoding="utf-8")
        else:
            from docx import Document
            doc = Document(str(path))
            text = "\n".join(p.text for p in doc.paragraphs)
        assert "```" not in text, f"Code fence leaked into saved document: {text!r}"
    _cleanup()


def test_write_document_strips_llm_preamble():
    _cleanup()
    preamble_response = (
        "Here is your assignment:\nThis is an introductory paragraph with enough words to "
        "pass content validation thresholds and set up the topic clearly for the reader.\n\n"
        "This is the body paragraph developing the main points in detail with supporting "
        "examples that elaborate on the core ideas presented in the introduction above.\n\n"
        "This is the concluding paragraph that wraps up the discussion and summarizes the "
        "key takeaways for the reader before the document ends here."
    )
    with patch("core.ollama_client.OllamaClient.is_available", return_value=True), \
         patch("core.ollama_client.OllamaClient.chat", return_value=preamble_response):
        r = bt.write_document("test topic")
        assert "saved" in r
        path = Path(r["saved"])
        if path.suffix == ".txt":
            text = path.read_text(encoding="utf-8")
        else:
            from docx import Document
            doc = Document(str(path))
            text = "\n".join(p.text for p in doc.paragraphs)
        assert "here is your" not in text.lower(), f"LLM preamble leaked into document: {text!r}"
    _cleanup()


def test_write_document_handles_special_characters_in_topic_for_filename():
    _cleanup()
    with patch("core.ollama_client.OllamaClient.is_available", return_value=True), \
         patch("core.ollama_client.OllamaClient.chat",
               return_value="This is an introductory paragraph with enough words to pass the minimum content validation threshold easily. It sets up the topic clearly and establishes context for the reader who is following along here.\n\nThis is the body paragraph that develops the main points in more detail with supporting examples throughout. It elaborates further on the core ideas presented in the introduction above in more depth.\n\nThis is the concluding paragraph that wraps up the discussion and summarizes the key takeaways for the reader now."):
        r = bt.write_document("AI/ML: what's next?! (2026)")
        assert "saved" in r
        # Filename must not contain illegal filesystem characters
        filename = Path(r["saved"]).name
        for bad_char in ["/", ":", "?", "!", "*"]:
            assert bad_char not in filename, f"Unsafe filename character leaked through: {filename!r}"
    _cleanup()


def test_write_document_handles_long_topic_gracefully():
    _cleanup()
    long_topic = "the complete history and socioeconomic impact of " + "renewable energy " * 20
    with patch("core.ollama_client.OllamaClient.is_available", return_value=True), \
         patch("core.ollama_client.OllamaClient.chat",
               return_value="This is an introductory paragraph with enough words to pass the minimum content validation threshold easily. It sets up the topic clearly and establishes context for the reader who is following along here.\n\nThis is the body paragraph that develops the main points in more detail with supporting examples throughout. It elaborates further on the core ideas presented in the introduction above in more depth.\n\nThis is the concluding paragraph that wraps up the discussion and summarizes the key takeaways for the reader now."):
        r = bt.write_document(long_topic)
        assert "saved" in r
        filename = Path(r["saved"]).stem
        assert len(filename) <= 60  # truncated to something filesystem-safe
    _cleanup()


def test_write_document_does_not_overwrite_existing_file_with_same_topic():
    _cleanup()
    with patch("core.ollama_client.OllamaClient.is_available", return_value=True), \
         patch("core.ollama_client.OllamaClient.chat",
               return_value="This is the first version of the introductory paragraph with sufficient length to pass content validation. It introduces the topic in a clear and structured way for the reader to follow easily.\n\nThis is the first version of the body paragraph containing the main supporting arguments and details. It builds on the introduction with concrete examples relevant to the topic at hand.\n\nThis is the first version of the concluding paragraph that ties everything together for the reader."):
        r1 = bt.write_document("duplicate topic test")
        path1 = Path(r1["saved"])
        original_content = path1.read_bytes()

    with patch("core.ollama_client.OllamaClient.is_available", return_value=True), \
         patch("core.ollama_client.OllamaClient.chat",
               return_value="This is the second version of the introductory paragraph with sufficient length to pass content validation. It introduces the topic in a different way for the reader to follow easily this time around.\n\nThis is the second version of the body paragraph containing different supporting arguments and details. It builds on the introduction with different concrete examples relevant to the topic at hand.\n\nThis is the second version of the concluding paragraph that ties everything together differently for the reader."):
        r2 = bt.write_document("duplicate topic test")
        path2 = Path(r2["saved"])

    assert path1.exists()
    assert path2.exists()
    # Should not have clobbered the first file
    assert path1.read_bytes() == original_content or path1 != path2
    _cleanup()


def test_write_document_falls_back_to_txt_without_python_docx():
    _cleanup()
    import builtins
    real_import = builtins.__import__

    def fake_import(name, *args, **kwargs):
        if name == "docx":
            raise ImportError("simulated missing python-docx")
        return real_import(name, *args, **kwargs)

    with patch("core.ollama_client.OllamaClient.is_available", return_value=True), \
         patch("core.ollama_client.OllamaClient.chat",
               return_value="This is an introductory paragraph with enough words to pass the minimum content validation threshold easily. It sets up the topic clearly and establishes context for the reader who is following along here.\n\nThis is the body paragraph that develops the main points in more detail with supporting examples throughout. It elaborates further on the core ideas presented in the introduction above in more depth.\n\nThis is the concluding paragraph that wraps up the discussion and summarizes the key takeaways for the reader now."), \
         patch("builtins.__import__", side_effect=fake_import):
        r = bt.write_document("docx missing test")
        assert "saved" in r
        assert r["format"] == "txt"
        assert Path(r["saved"]).exists()
    _cleanup()


def test_write_document_uses_single_attempt_not_stacked_retries():
    """Regression test for the actual crash: retries=2 + 120s timeout could
    take up to 360s, exceeding the tool's own timeout and orphaning the
    background thread mid-execution. write_document must call chat() with
    retries=0 so one slow generation can't silently triple its own wait."""
    _cleanup()
    with patch("core.ollama_client.OllamaClient.is_available", return_value=True), \
         patch("core.ollama_client.OllamaClient.chat",
               return_value="Intro paragraph with enough words to pass validation checks nicely. "
                            "Body paragraph with more detail and supporting examples throughout. "
                            "Conclusion paragraph wrapping up the discussion for the reader now.") as mock_chat:
        bt.write_document("test topic")
        _, kwargs = mock_chat.call_args
        assert kwargs.get("retries") == 0, f"Expected retries=0, got {kwargs.get('retries')}"
        assert kwargs.get("timeout", 0) >= 150, "Timeout should be generous for long-form generation"
    _cleanup()


if __name__ == "__main__":
    import inspect
    fails = 0
    for name, fn in list(globals().items()):
        if name.startswith("test_") and inspect.isfunction(fn):
            try:
                fn()
                print(f"PASS {name}")
            except AssertionError as e:
                fails += 1
                print(f"FAIL {name}: {e}")
            except Exception as e:
                fails += 1
                print(f"ERROR {name}: {e}")
    _cleanup()
    sys.exit(1 if fails else 0)
