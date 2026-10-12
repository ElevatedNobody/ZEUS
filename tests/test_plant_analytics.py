"""
Data-grounded deliverables: plant database -> IndustrialAnalysis -> files.

The thing being defended here is that generated documents contain TRUE
numbers. A deck that looks professional and states a pressure that no
instrument ever recorded is worse than no deck, because it will be believed.

So these tests check three things, in order of how much damage the failure
would do:

  1. Figures in an analysis trace to rows in the database.
  2. A request with no plant data behind it produces no analysis, rather than
     an invented one.
  3. Demonstration data is always labelled as such.

The rendering tests then check the two faults that a structural assertion
cannot see and only appear once the file is opened: text overflowing its box,
and a statistic that is arithmetically fine but meaningless.
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from industrial import demo_refinery as demo


@pytest.fixture(scope="module")
def seeded():
    """The demonstration refinery, loaded once for the module.

    Seeding is idempotent, so this is safe whether or not the developer has
    already loaded it.
    """
    import importlib.util
    spec = importlib.util.spec_from_file_location(
        "seed_demo_refinery", ROOT / "scripts" / "seed_demo_refinery.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    module.seed()
    from industrial import plant_analytics
    return plant_analytics


# ================================================== the data is coherent ====

def test_the_dataset_agrees_with_itself():
    """Runs the shipped consistency check. If the seed data contradicts
    itself, every downstream assertion is meaningless."""
    import importlib.util
    spec = importlib.util.spec_from_file_location(
        "seed_demo_refinery", ROOT / "scripts" / "seed_demo_refinery.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    assert module.verify() == 0


def test_the_dataset_is_marked_as_demonstration_data():
    assert "SAMPLE" in demo.COMPANY["note"].upper()
    assert demo.DEMO == "DEMO"


def test_no_real_organisation_is_impersonated():
    """The dataset must not carry a real company's name, since documents
    generated from it look official."""
    forbidden = ("reliance", "indian oil", "bharat petroleum", "hindustan petroleum",
                 "mrpl", "mangalore refinery", "shell", "bp plc", "exxon", "chevron")
    blob = " ".join([demo.COMPANY["name"], demo.COMPANY["site"],
                     demo.COMPANY["location"]]).lower()
    for name in forbidden:
        assert name not in blob, f"the demo company resembles {name}"


# ============================================== figures trace to the data ====

def test_an_equipment_dossier_reports_the_real_status(seeded):
    from industrial import plant_data

    analysis = seeded.equipment_dossier("P-101")
    assert analysis is not None
    row = plant_data.resolve_equipment("P-101")

    status = next(k for k in analysis.kpis if k.label == "Status")
    assert status.value.lower() == str(row["status"]).replace("_", " ").lower()

    interventions = next(k for k in analysis.kpis if k.label == "Interventions")
    assert int(interventions.value) == len(plant_data.maintenance_history("P-101",
                                                                          limit=50))


def test_the_maintenance_table_matches_the_database_row_for_row(seeded):
    from industrial import plant_data

    analysis = seeded.equipment_dossier("P-101")
    table = analysis.table("Maintenance history")
    records = plant_data.maintenance_history("P-101", limit=50)
    assert table is not None
    assert len(table.rows) == len(records)

    dates_in_table = {str(row[0]) for row in table.rows}
    dates_in_database = {str(r.get("performed_on")) for r in records}
    assert dates_in_table == dates_in_database, "the table invented or lost a row"


def test_the_trend_values_come_from_the_readings_table(seeded):
    analysis = seeded.equipment_dossier("P-101")
    trend = next((t for t in analysis.tables if "trend" in t.title.lower()), None)
    assert trend is not None, "P-101 has readings but no trend table was built"

    charted = [float(row[1]) for row in trend.rows]
    stored = {round(value, 2) for _, value in
              seeded._readings("P-101")["vibration_velocity"]}
    for value in charted:
        assert round(value, 2) in stored, f"{value} is not a recorded reading"


def test_an_unknown_asset_produces_nothing(seeded):
    """Not an empty document - nothing. A dossier for equipment that does not
    exist could only be fabricated."""
    assert seeded.equipment_dossier("ZZ-999") is None
    assert seeded.equipment_dossier("") is None


def test_a_non_plant_question_gets_no_analysis(seeded):
    for question in ("what is the capital of France",
                     "write me a poem about the sea",
                     "how do I sort a list in python"):
        assert seeded.analysis_for(question) is None, (
            f"{question!r} produced a plant analysis it has no data for")


def test_an_equipment_tag_beats_a_topic_word(seeded):
    """'maintenance report for P-101' is about P-101, not about maintenance in
    general. Getting this backwards returns fleet-wide figures under a title
    the reader will take as being about one pump."""
    analysis = seeded.analysis_for("create a maintenance report for P-101")
    assert analysis is not None
    assert analysis.title.startswith("P-101")


def test_topic_requests_reach_their_builders(seeded):
    expected = {
        "asset reliability overview": "Asset Reliability Overview",
        "maintenance performance review": "Maintenance Performance Review",
        "safety and incident report": "Safety Performance Report",
        "production throughput": "Production Performance",
        "spares and stock levels": "Spares and Inventory Status",
    }
    for request, title in expected.items():
        analysis = seeded.analysis_for(request)
        assert analysis is not None, f"{request!r} produced nothing"
        assert analysis.title == title


# ================================================== demo data is labelled ====

def test_every_seeded_analysis_is_flagged_as_demo(seeded):
    for request in ("P-101 maintenance report", "asset reliability overview",
                    "maintenance performance review", "safety report",
                    "inventory status", "refinery overview"):
        analysis = seeded.analysis_for(request)
        if analysis is None:
            continue
        assert analysis.contains_demo_data, (
            f"{request!r} was built from DEMO rows but is not flagged")
        assert "DEMO" in analysis.demo_banner()


def test_findings_carry_their_evidence(seeded):
    """A finding without a source cannot be checked by the person reading it."""
    analysis = seeded.equipment_dossier("P-101")
    assert analysis.findings
    for finding in analysis.findings:
        assert finding.evidence, f"{finding.title!r} has no evidence"


def test_recommendations_come_from_the_record_not_from_nowhere(seeded):
    """Every recommendation must be traceable to an inspection, an open work
    order or a stock level - never generic maintenance advice."""
    from industrial import plant_data

    analysis = seeded.equipment_dossier("P-101")
    inspections = plant_data.inspection_history("P-101", limit=50)
    recommendations = {r.title for r in analysis.recommendations}
    assert recommendations

    sources = {str(i.get("recommendation", ""))[:140] for i in inspections}
    sources |= {f"Close out work order {m.get('work_order', '')} and verify "
                f"before return to service"
                for m in plant_data.maintenance_history("P-101", limit=50)}
    sources |= {f"Replenish {s.get('part_number')} to minimum stock"
                for s in seeded._spares_for("P-101")}
    # A next-due date that has passed is also a record-derived reason.
    sources |= {f"Schedule the follow-up to {m.get('work_order', 'the last job')}, "
                f"due {m.get('next_due')}"
                for m in plant_data.maintenance_history("P-101", limit=50)}
    for recommendation in recommendations:
        assert recommendation in sources, (
            f"{recommendation!r} does not trace to a record")


# ====================================================== rendered output ====

def test_a_deck_is_produced_and_opens(seeded, tmp_path):
    from pptx import Presentation

    from industrial.deck_builder import build_deck

    analysis = seeded.equipment_dossier("P-101")
    path = tmp_path / "p101.pptx"
    result = build_deck(analysis, path)

    assert path.is_file() and path.stat().st_size > 20_000
    assert "error" not in result
    deck = Presentation(str(path))
    assert len(deck.slides) >= 6


def test_the_deck_charts_the_condition_trend(seeded, tmp_path):
    """The vibration trend is the argument of a P-101 report. A row limit
    tuned for bar charts used to reject it silently, and the deck went out
    with no chart at all."""
    from pptx import Presentation

    from industrial.deck_builder import build_deck

    path = tmp_path / "p101.pptx"
    build_deck(seeded.equipment_dossier("P-101"), path)

    charts = [shape.chart for slide in Presentation(str(path)).slides
              for shape in slide.shapes if shape.has_chart]
    assert charts, "the deck contains no chart"
    points = len(charts[0].plots[0].categories)
    assert points >= 20, f"the trend was reduced to {points} points"


def test_a_trend_chart_is_not_summed(seeded):
    """Adding up a series of measurements produces a number that does not
    exist. The read-out for a trend has to be latest/change/peak, never
    total and share."""
    from industrial.analysis import TableBlock
    from industrial.deck_builder import DeckBuilder

    trend = TableBlock(title="Vibration velocity trend",
                       columns=["Date", "Vibration"],
                       rows=[[f"2026-01-{d:02d}", 2.0 + d * 0.1] for d in range(1, 15)])
    categories = TableBlock(title="Interventions by type",
                            columns=["Type", "Count"],
                            rows=[["Preventive", 8], ["Corrective", 3]])
    assert DeckBuilder._is_trend(trend) is True
    assert DeckBuilder._is_trend(categories) is False


def test_prose_is_trimmed_to_the_box_it_is_drawn_in():
    """Text that overflows its card prints on top of whatever is below it.
    This was visible only by rendering the deck - no structural assertion
    catches it."""
    from industrial.deck_builder import _chars_that_fit, _fit_prose

    long_text = ("This is a sentence that goes on. " * 40).strip()
    budget = _chars_that_fit(6.89, 1.76, 13.5, 1.5)
    fitted = _fit_prose(long_text, 6.89, 1.76, 13.5, 1.5)

    assert len(fitted) <= budget
    assert len(fitted) > budget * 0.5, "trimmed so hard the box is half empty"
    assert fitted.endswith((".", "…")), f"cut mid-word: {fitted[-40:]!r}"


def test_short_prose_is_left_alone():
    from industrial.deck_builder import _fit_prose
    text = "P-101 is under maintenance."
    assert _fit_prose(text, 6.89, 1.76, 13.5, 1.5) == text


def test_a_workbook_is_produced_and_opens(seeded, tmp_path):
    from openpyxl import load_workbook

    from industrial.workbook_builder import build_workbook

    analysis = seeded.equipment_dossier("P-101")
    path = tmp_path / "p101.xlsx"
    build_workbook(analysis, path)

    assert path.is_file() and path.stat().st_size > 8_000
    workbook = load_workbook(str(path))
    assert len(workbook.sheetnames) >= 3
    assert any("maintenance" in name.lower() for name in workbook.sheetnames)


def test_the_workbook_carries_the_demo_banner(seeded, tmp_path):
    from openpyxl import load_workbook

    from industrial.workbook_builder import build_workbook

    path = tmp_path / "p101.xlsx"
    build_workbook(seeded.equipment_dossier("P-101"), path)

    text = " ".join(
        str(cell.value) for sheet in load_workbook(str(path)).worksheets
        for row in sheet.iter_rows(max_row=12) for cell in row if cell.value)
    assert "DEMO" in text.upper(), "a workbook of sample data is not labelled"


def test_every_topic_builds_all_three_formats(seeded, tmp_path):
    """The whole point of the IndustrialAnalysis seam: one analysis renders to
    deck, workbook and report without any of them needing special cases."""
    from industrial.deck_builder import build_deck
    from industrial.workbook_builder import build_workbook

    for request in ("asset reliability overview", "maintenance performance review",
                    "safety report", "production throughput", "inventory status",
                    "refinery overview"):
        analysis = seeded.analysis_for(request)
        assert analysis is not None, f"{request!r} produced nothing"
        stem = request.replace(" ", "_")[:24]
        deck = tmp_path / f"{stem}.pptx"
        book = tmp_path / f"{stem}.xlsx"
        build_deck(analysis, deck)
        build_workbook(analysis, book)
        assert deck.stat().st_size > 15_000, f"{request}: deck is suspiciously small"
        assert book.stat().st_size > 5_000, f"{request}: workbook is suspiciously small"
