"""
Phase 17 end-to-end validation.

Each test drives a real user sentence through the real decision engine, the
real tool registry and the real generators, then verifies the artefact on disk.
Nothing is stubbed except the confirmation prompt (a person clicking OK) and
the documents directory (so tests do not write into ~/Documents).

Ollama is NOT running in this environment, which is the point of several of
these: the data-driven industrial paths must work anyway, and the model-driven
ones must fail honestly instead of pretending.
"""

import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from memory import database  # noqa: E402

DB_PATH = Path(tempfile.gettempdir()) / "medha_test_e2e.db"

from core import office  # noqa: E402
from core.decision_engine import DecisionEngine  # noqa: E402
from core.security_manager import security_manager  # noqa: E402
from core.tool_registry import registry  # noqa: E402
from industrial import plant_data as pd  # noqa: E402

OUT = Path(tempfile.mkdtemp(prefix="medha_e2e_docs_"))
_engine = DecisionEngine()


def setup_function(function):
    database.DB_PATH = DB_PATH
    database.set_active_user(None)
    security_manager.set_confirm_callback(lambda message: True)
    office.documents_dir = lambda subfolder="": OUT
    import industrial.plant_docs as docs
    docs.office.documents_dir = lambda subfolder="": OUT


def setup_module(module):
    database.DB_PATH = DB_PATH
    if DB_PATH.exists():
        DB_PATH.unlink()
    database.init_db()
    pd.init_plant_schema(force=True)
    pd.seed_demo_data(force=True)


def _run(sentence: str):
    """USER COMMAND -> INTENT -> TOOL -> EXECUTION -> RESULT, the real path."""
    match = _engine._match_direct_command(sentence)
    assert match is not None, f"no route for: {sentence!r}"
    tool, args = match
    result = registry.execute(tool, **args)
    assert result.success, f"{tool} failed: {result.error}"
    return tool, result.output


def _real_file(output, extension):
    assert "error" not in output, output
    assert output.get("verified") is True, output
    path = Path(output["saved"])
    assert path.exists() and path.stat().st_size > 1000
    assert path.suffix.lower() == f".{extension}"
    return path


# --------------------------------------------------------------- TEST A --

def test_a_word_document_command_reaches_the_word_generator():
    """'Create a Word document about refinery maintenance.'

    Ollama is down here, so the generator cannot invent prose. What must be
    true is that the command REACHES write_document and that the failure is
    honest - this is the routing that was missing.
    """
    tool, args = _engine._match_direct_command("Create a Word document about refinery maintenance.")
    assert tool == "write_document"
    assert "refinery maintenance" in args["topic"]
    output = registry.execute(tool, **args).output
    assert "saved" in output or "Ollama" in str(output.get("error", "")), output
    if "saved" in output:
        _real_file(output, "docx")


def test_a2_data_driven_word_report_really_is_generated():
    """The industrial equivalent needs no model at all, so it must produce a
    real .docx here and now."""
    tool, output = _run("Generate a maintenance report for P-204")
    assert tool == "maintenance_report"
    path = _real_file(output, "docx")
    from docx import Document
    assert "P-204" in "\n".join(p.text for p in Document(str(path)).paragraphs)


# --------------------------------------------------------------- TEST B --

def test_b_powerpoint_command_routes_with_the_slide_count():
    tool, args = _engine._match_direct_command(
        "Make a 10-slide presentation on predictive maintenance.")
    assert tool == "create_presentation"
    assert args["slide_count"] == 10
    assert args["topic"] == "predictive maintenance"


def test_b2_operations_presentation_really_is_generated():
    tool, output = _run("Create a management presentation of the maintenance situation.")
    assert tool == "operations_presentation"
    path = _real_file(output, "pptx")
    from pptx import Presentation
    assert len(Presentation(str(path)).slides) >= 5


# --------------------------------------------------------------- TEST C --

def test_c_excel_inventory_really_is_generated():
    tool, output = _run("Make an equipment inventory spreadsheet")
    assert tool == "export_plant_data"
    path = _real_file(output, "xlsx")
    from openpyxl import load_workbook
    workbook = load_workbook(str(path))
    assert workbook.active.max_row > 1


def test_c2_excel_command_for_free_text_reaches_the_generator():
    tool, args = _engine._match_direct_command("Create an Excel sheet for inspection data.")
    assert tool == "create_spreadsheet"
    assert args["topic"] == "inspection data"


# --------------------------------------------------------------- TEST D --

def test_d_active_alerts_return_structured_results():
    tool, output = _run("Show active alerts.")
    assert tool == "list_plant_alerts"
    assert output["alert_count"] >= 1
    first = output["alerts"][0]
    for field in ("alert_id", "raised_at", "equipment_tag", "category", "severity",
                  "description", "status", "source"):
        assert field in first, f"alert is missing {field}"
    assert output["by_severity"]


# --------------------------------------------------------------- TEST E --

def test_e_maintenance_history_returns_structured_results():
    tool, output = _run("Show maintenance history for Pump P-204.")
    assert tool == "maintenance_history"
    assert output["equipment"] == "P-204"
    assert output["record_count"] >= 1
    assert output["records"][0]["equipment_tag"] == "P-204"


# --------------------------------------------------------------- TEST F --

def test_f_export_maintenance_history_to_excel():
    tool, output = _run("Export the maintenance history to Excel")
    assert tool == "export_plant_data"
    path = _real_file(output, "xlsx")
    from openpyxl import load_workbook
    workbook = load_workbook(str(path))
    assert "About" in workbook.sheetnames
    assert workbook["Maintenance"].max_row == output["rows"] + 1


# --------------------------------------------------------------- TEST G --

def test_g_maintenance_report_saved_as_word():
    tool, output = _run("Generate a maintenance report")
    assert tool == "maintenance_report"
    _real_file(output, "docx")


# --------------------------------------------------------------- TEST H --

def test_h_equipment_presentation_summarises_one_item():
    tool, output = _run("Create a presentation for equipment P-204")
    assert tool == "equipment_presentation"
    _real_file(output, "pptx")


# --------------------------------------------------------------- TEST I --

def test_i_everything_here_worked_without_a_network():
    """Nothing in these flows opened a socket to anything but loopback Ollama
    (which is not even running). Assert no web client is loaded and no HTTP
    server framework was imported by any of it."""
    forbidden = {"fastapi", "flask", "uvicorn", "django", "starlette", "tornado",
                 "boto3", "openai", "google.cloud"}
    assert not (forbidden & set(sys.modules))


def test_i2_industrial_answers_work_with_ollama_down():
    from core.ollama_client import ollama_client
    assert not ollama_client.is_available(), (
        "this test is meaningful only when Ollama is NOT running")
    for sentence in ("Show the plant status", "Show active alerts.",
                     "Show the equipment inventory.", "Show system information",
                     "Show maintenance history for P-204", "Show the safety status"):
        tool, output = _run(sentence)
        assert output, f"{sentence} returned nothing"


# ----------------------------------------------------------- honesty --

def test_generated_documents_are_verified_not_assumed():
    """core.office.finish must refuse to report success for a missing file."""
    missing = Path(tempfile.gettempdir()) / "medha_not_written.docx"
    if missing.exists():
        missing.unlink()
    result = office.finish(missing, "docx", open_after=False)
    assert "error" in result
    assert "saved" not in result


def test_an_empty_file_is_not_reported_as_generated():
    empty = Path(tempfile.mkdtemp()) / "empty.xlsx"
    empty.write_bytes(b"")
    result = office.finish(empty, "xlsx", open_after=False)
    assert "error" in result and "empty" in result["error"]


def test_a_corrupt_file_is_not_reported_as_generated():
    """The check that actually matters: the file exists and is non-empty, but
    cannot be reopened by the library that is supposed to have written it."""
    corrupt = Path(tempfile.mkdtemp()) / "corrupt.docx"
    corrupt.write_bytes(b"this is not a docx" * 100)
    result = office.finish(corrupt, "docx", open_after=False)
    assert "error" in result
    assert "reopened" in result["error"] or "verification" in result


def test_the_full_chain_is_reachable_for_every_documented_command():
    """Phase 6: the natural-language commands the brief lists must all route."""
    sentences = [
        "Show the current equipment status.",
        "Show active alerts.",
        "Show Pump P-204 maintenance history.",
        "Generate a maintenance report",
        "Make an equipment inventory spreadsheet",
        "Create a management presentation of the maintenance situation.",
        "Create a Word approval note for P-204",
        "Export today's alerts to Excel.",
        "Show today's maintenance schedule.",
        "Show the inspection history for P-204",
        "Generate a safety report.",
        "Find the SOP for pump maintenance.",
        "Show system information",
        "Show the logs",
    ]
    unrouted = [s for s in sentences if _engine._match_direct_command(s) is None]
    assert not unrouted, f"these commands do not reach a tool: {unrouted}"
